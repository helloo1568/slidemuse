#!/usr/bin/env python3
"""Real-backend regression controls for Chinese, tables, fonts and chart labels."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from build_editable_ppt import text_frame
from json_io import read_json
from measured_layout import findings
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt
from render_deck import build_review, render_selected, sha256
from render_environment import fingerprint
from render_process import rendering_options


def build_controls(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(12), Inches(6.75)

    def slide(title):
        page = prs.slides.add_slide(prs.slide_layouts[6])
        box = page.shapes.add_textbox(Inches(.4), Inches(.2), Inches(11), Inches(.6))
        box.name = "title"
        text_frame(box.text_frame, title, {"font_size": 26, "font_face": "Arial"})
        return page

    page = slide("01 中文长标题与正常正文 / Chinese control")
    box = page.shapes.add_textbox(Inches(.5), Inches(1.3), Inches(10), Inches(2))
    box.name = "spacious"
    text_frame(box.text_frame, "交通运输研究：保留来源与边界\nNormal mixed text 2026", {"font_size": 24, "font_face": "Arial"})
    page = slide("02 密集文字 / intentional overflow control")
    box = page.shapes.add_textbox(Inches(.5), Inches(1.3), Inches(2), Inches(.4))
    box.name = "overflow"
    text_frame(box.text_frame, "First\nSecond\nThird\nFourth\nFifth", {"font_size": 24})
    page = slide("03 实测文字重叠 / intentional overlap control")
    for name in ("overlap-a", "overlap-b"):
        box = page.shapes.add_textbox(Inches(.5), Inches(1.3), Inches(6), Inches(.8))
        box.name = name
        text_frame(box.text_frame, "Actual overlapping text 交通运输" if name == "overlap-a" else "Second text collides with the first", {"font_size": 24})
    page = slide("04 表格与中文 / table control")
    table = page.shapes.add_table(3, 2, Inches(.5), Inches(1.4), Inches(10), Inches(2)).table
    for row, values in enumerate((("城市", "出行量 / synthetic"), ("Alpha", "120"), ("Beta", "80"))):
        for col, value in enumerate(values):
            text_frame(table.cell(row, col).text_frame, value, {"font_size": 20})
    page = slide("05 混合字体与字号 / font control")
    box = page.shapes.add_textbox(Inches(.5), Inches(1.3), Inches(10), Inches(2))
    box.name = "mixed-fonts"
    text_frame(box.text_frame, "", {"font_size": 24})
    for value, face, size in (("Arial 24 ", "Arial", 24), ("中文 28 ", "Microsoft YaHei", 28), ("Serif 20", "Times New Roman", 20)):
        run = box.text_frame.paragraphs[0].add_run()
        run.text, run.font.name, run.font.size = value, face, Pt(size)
    page = slide("06 图表标签 / chart-label visual control")
    data = CategoryChartData()
    data.categories = ["Long category label Alpha", "中文分类 Beta"]
    data.add_series("Synthetic values", [120, 80])
    chart = page.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.5), Inches(1.4), Inches(10), Inches(4.8), data).chart
    chart.plots[0].has_data_labels = True
    prs.save(path)


def verify(output, backend="powerpoint", width=1200, timeout=120):
    output = output.resolve()
    if output.exists():
        raise ValueError("Regression output must be new; preserve prior regression evidence")
    output.mkdir(parents=True)
    deck = output / "controls.pptx"
    build_controls(deck)
    renders = output / "rendered"
    with rendering_options(timeout=timeout, retries=1, measure_text=backend == "powerpoint"):
        actual = render_selected(deck, renders, width, backend, list(range(1, 7)))
    diagnostics = []
    pages = []
    for index in range(1, 7):
        path = renders / f"{index:03d}.png"
        with Image.open(path) as image:
            image.verify()
        pages.append({"page_number": index, "path": str(path), "sha256": sha256(path)})
        measure = renders / f"{index:03d}.layout.json"
        if measure.exists():
            diagnostics.extend(findings(read_json(measure), f"s{index:02d}"))
    checks = {"six_actual_pages": len(pages) == 6}
    if actual == "powerpoint":
        checks.update(spacious_control=any(f["element_id"] == "spacious" and f["status"] == "no_obvious_risk" for f in diagnostics),
                      overflow_control=any(f["element_id"] == "overflow" and f["status"] == "risk" for f in diagnostics),
                      overlap_control=any(f.get("kind") == "text_overlap" and f["element_id"] in ("overlap-a", "overlap-b") for f in diagnostics))
    build_review([Path(p["path"]) for p in pages], output, None)
    report = {"version": "1.0", "backend": actual, "environment": fingerprint(), "deck_sha256": sha256(deck),
              "pages": pages, "checks": checks, "measured_layout": diagnostics,
              "mechanical_status": "pass" if all(checks.values()) else "fail", "visual_review_status": "pending",
              "scope": "Synthetic positive/negative regression controls; actual images require inspection. Chart labels are visual-only."}
    (output / "regression-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    parser.add_argument("--backend", choices=("powerpoint", "libreoffice"), default="powerpoint")
    parser.add_argument("--width", type=int, default=1200)
    parser.add_argument("--timeout", type=float, default=120)
    args = parser.parse_args()
    report = verify(args.output, args.backend, args.width, args.timeout)
    print(json.dumps({"status": report["mechanical_status"], "checks": report["checks"]}, ensure_ascii=True))
    if report["mechanical_status"] != "pass":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
