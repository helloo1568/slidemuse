#!/usr/bin/env python3
"""Inspect exported PPTX objects and compare them with the source scene."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
from collections import Counter
from pathlib import Path

from chart_style import style_errors
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches
from scene import asset_path, load_scene, walk
from speaker_notes import check_notes


def objects(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from objects(shape.shapes)


def kind(shape):
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if shape.has_table:
        return "table"
    if shape.has_chart:
        return "chart"
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "image"
    if shape.shape_type == MSO_SHAPE_TYPE.LINE:
        return "line"
    if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
        return "text"
    return "shape"


def color_hex(color):
    try:
        return f"#{color.rgb}"
    except (AttributeError, ValueError):
        return None


def pixel_hash(source):
    with Image.open(source) as im:
        im = ImageOps.exif_transpose(im).convert("RGBA")
        return hashlib.sha256(str(im.size).encode() + im.tobytes()).hexdigest()


def expected_geometry(element, base, sx, sy):
    """Return the unrotated OOXML box and image crop from Scene coordinates."""
    kind = element["type"]
    crop = (0, 0, 0, 0)  # left, top, right, bottom
    if kind == "group":
        boxes = [expected_geometry(child, base, sx, sy)[0] for child in element["children"]]
        left, top = min(b[0] for b in boxes), min(b[1] for b in boxes)
        right, bottom = max(b[0] + b[2] for b in boxes), max(b[1] + b[3] for b in boxes)
        return (left, top, right - left, bottom - top), crop
    if kind == "line":
        x1, y1, x2, y2 = (
            round(element[key] * scale)
            for key, scale in zip(("x1", "y1", "x2", "y2"), (sx, sy, sx, sy))
        )
        return (min(x1, x2), min(y1, y2), abs(x2 - x1), abs(y2 - y1)), crop
    x, y, w, h = (
        round(element[key] * scale)
        for key, scale in zip(("x", "y", "w", "h"), (sx, sy, sx, sy))
    )
    if kind == "image":
        with Image.open(asset_path(base, element["path"])) as image:
            iw, ih = ImageOps.exif_transpose(image).size
        fit = element.get("fit", "contain")
        if fit == "contain":
            scale = min(w / iw, h / ih)
            fitted_w, fitted_h = round(iw * scale), round(ih * scale)
            x, y = x + (w - fitted_w) // 2, y + (h - fitted_h) // 2
            w, h = fitted_w, fitted_h
        elif fit == "cover":
            if iw / ih > w / h:
                margin = (1 - (w / h) / (iw / ih)) / 2
                crop = (margin, 0, margin, 0)
            else:
                margin = (1 - (iw / ih) / (w / h)) / 2
                crop = (0, margin, 0, margin)
    return (x, y, w, h), crop


def group_transform_matches(shape, box):
    # Scene groups use absolute child coordinates and have no extra transform.
    # Inspect both coordinate spaces; checking child boxes alone misses a moved group.
    transform = shape._element.grpSpPr.xfrm
    if transform is None:
        return False
    for offset, extent in ((transform.off, transform.ext), (transform.chOff, transform.chExt)):
        if offset is None or extent is None:
            return False
        actual = (offset.x, offset.y, extent.cx, extent.cy)
        if any(abs(a - b) > 4 for a, b in zip(actual, box)):
            return False
    return not (transform.flipH or transform.flipV)


def audit(pptx: Path, scene_path: Path | None = None) -> dict:
    prs = Presentation(pptx)
    scene, warnings = load_scene(scene_path) if scene_path else (None, [])
    errors, pages = [], []
    if scene:
        errors.extend(check_notes(prs, scene["slides"], scene=True))
    if scene and len(prs.slides) != len(scene["slides"]):
        errors.append("Slide count differs from scene")
    if scene:
        canvas = scene["canvas"]
        expected_width = Inches(canvas.get("width_inches", 13.333333))
        expected_height = round(expected_width * canvas["height"] / canvas["width"])
        if (prs.slide_width, prs.slide_height) != (expected_width, expected_height):
            errors.append("Slide dimensions differ from scene")
        sx, sy = expected_width / canvas["width"], expected_height / canvas["height"]
    for index, slide in enumerate(prs.slides):
        flat = list(objects(slide.shapes))
        counts = Counter(kind(s) for s in flat)
        expected = (
            scene["slides"][index] if scene and index < len(scene["slides"]) else None
        )
        by_id = {}
        for shape in flat:
            eid = shape.name.split(" | ", 1)[0]
            if eid in by_id:
                errors.append(f"Slide {index + 1}: duplicate object name/id {eid}")
            by_id[eid] = shape
            if kind(shape) == "image":
                try:
                    with Image.open(io.BytesIO(shape.image.blob)) as im:
                        im.verify()
                except OSError:
                    errors.append(f"Slide {index + 1}/{eid}: corrupt image")
                if (
                    shape.width * shape.height
                    >= prs.slide_width * prs.slide_height * 0.8
                ):
                    warnings.append(
                        f"Slide {index + 1}/{eid}: large raster requires visual layer inspection"
                    )
        if expected:
            expected_elements = list(walk(expected["elements"]))
            expected_ids = {e["id"] for e in expected_elements}
            extra = set(by_id) - expected_ids
            if extra:
                errors.append(f"Slide {index + 1}: unexpected objects {sorted(extra)}")
            for e in expected_elements:
                label = f"Slide {index + 1}/{e['id']}"
                shape = by_id.get(e["id"])
                if shape is None:
                    errors.append(f"{label}: missing object")
                    continue
                if kind(shape) != e["type"]:
                    errors.append(f"{label}: expected {e['type']}, got {kind(shape)}")
                    continue
                wanted_box, wanted_crop = expected_geometry(e, scene_path.resolve().parent, sx, sy)
                if e["type"] not in ("group", "line"):
                    actual_box = [shape.left, shape.top, shape.width, shape.height]
                    if any(abs(a - b) > 4 for a, b in zip(wanted_box, actual_box)):
                        errors.append(f"{label}: geometry differs from scene")
                if e["type"] == "line":
                    wanted_points = [
                        round(e[k] * scale)
                        for k, scale in zip(("x1", "y1", "x2", "y2"), (sx, sy, sx, sy))
                    ]
                    if any(
                        abs(a - b) > 4
                        for a, b in zip(
                            wanted_points,
                            (shape.begin_x, shape.begin_y, shape.end_x, shape.end_y),
                        )
                    ):
                        errors.append(f"{label}: line endpoints differ")
                if (
                    e["type"] != "line"
                    and abs(shape.rotation - e.get("rotation", 0) % 360) > 0.001
                ):
                    errors.append(f"{label}: rotation differs from scene")
                if e["type"] == "image":
                    original = asset_path(scene_path.resolve().parent, e["path"])
                    if pixel_hash(original) != pixel_hash(io.BytesIO(shape.image.blob)):
                        errors.append(f"{label}: image pixels differ from asset")
                    actual_crop = (shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom)
                    # OOXML stores crops at 1/100000 resolution.
                    if any(abs(a - b) > 1.1e-5 for a, b in zip(actual_crop, wanted_crop)):
                        errors.append(f"{label}: image crop differs from scene")
                    transform = shape._element.spPr.xfrm
                    if transform is not None and (transform.flipH or transform.flipV):
                        errors.append(f"{label}: image flip differs from scene")
                if e["type"] == "text" and shape.text != e["text"]:
                    errors.append(f"{label}: text differs from scene")
                if e["type"] == "shape" and "corner_radius" in e and abs(shape.adjustments[0] - e["corner_radius"]) > 1.1e-5:
                    errors.append(f"{label}: corner radius differs from scene")
                if e["type"] == "group":
                    if not group_transform_matches(shape, wanted_box):
                        errors.append(f"{label}: group transform differs from scene")
                    child_ids = [s.name.split(" | ", 1)[0] for s in shape.shapes]
                    if child_ids != [child["id"] for child in e["children"]]:
                        errors.append(f"{label}: group membership/order differs")
                if e["type"] == "table":
                    actual = [[c.text for c in row.cells] for row in shape.table.rows]
                    if actual != e["rows"]:
                        errors.append(f"{label}: table content differs")
                if e["type"] == "chart":
                    chart = shape.chart
                    actual = [
                        {"name": s.name, "values": list(s.values)} for s in chart.series
                    ]
                    wanted = [
                        {"name": s["name"], "values": s["values"]} for s in e["series"]
                    ]
                    if (
                        actual != wanted
                        or [c.label for c in chart.plots[0].categories]
                        != e["categories"]
                    ):
                        errors.append(f"{label}: chart data differs")
                    if chart.part.chart_workbook.xlsx_part is None:
                        errors.append(
                            f"{label}: chart has no embedded editable workbook"
                        )
                    style_differs = False
                    if "data_label_color" in e:
                        style_differs |= (
                            not chart.plots[0].has_data_labels
                            or color_hex(chart.plots[0].data_labels.font.color)
                            != e["data_label_color"].upper()
                        )
                    if e["chart_type"] not in ("pie", "doughnut"):
                        value_axis = chart.value_axis
                        for field, actual_scale in (
                            ("value_axis_min", value_axis.minimum_scale),
                            ("value_axis_max", value_axis.maximum_scale),
                        ):
                            if field in e:
                                style_differs |= actual_scale != e[field]
                        if "major_gridlines" in e or "major_gridline_color" in e:
                            style_differs |= value_axis.has_major_gridlines != e.get(
                                "major_gridlines", True
                            )
                        if "major_gridline_color" in e and value_axis.has_major_gridlines:
                            style_differs |= (
                                color_hex(value_axis.major_gridlines.format.line.color)
                                != e["major_gridline_color"].upper()
                            )
                        if "tick_label_color" in e:
                            style_differs |= any(
                                color_hex(axis.tick_labels.font.color)
                                != e["tick_label_color"].upper()
                                for axis in (chart.category_axis, value_axis)
                            )
                    if style_differs:
                        errors.append(f"{label}: chart style differs from scene")
                    errors.extend(f"{label}: chart style differs from scene ({finding})"
                                  for finding in style_errors(chart, e))
                if e["type"] == "image" and e.get("role") == "reference":
                    errors.append(f"{label}: reference image packaged as slide content")
            if [s.name.split(" | ", 1)[0] for s in slide.shapes] != [
                e["id"] for e in expected["elements"]
            ]:
                errors.append(f"Slide {index + 1}: top-level stacking order differs")
            if expected.get("source_image"):
                source_hash = pixel_hash(
                    asset_path(scene_path.resolve().parent, expected["source_image"])
                )
                if any(
                    kind(s) == "image"
                    and pixel_hash(io.BytesIO(s.image.blob)) == source_hash
                    for s in flat
                ):
                    errors.append(
                        f"Slide {index + 1}: original reference pixels remain in slide media"
                    )
        native = sum(counts[k] for k in ("text", "shape", "line", "chart", "table"))
        pages.append(
            {
                "slide": index + 1,
                "objects": dict(counts),
                "native_objects": native,
                "movable_raster_objects": counts["image"],
                "elements": [
                    {
                        "id": s.name.split(" | ", 1)[0],
                        "name": s.name,
                        "type": kind(s),
                        "shape_id": s.shape_id,
                    }
                    for s in flat
                ],
            }
        )
    return {
        "pptx": str(pptx.resolve()),
        "scene_compared": bool(scene),
        "slides": pages,
        "errors": errors,
        "warnings": sorted(set(warnings)),
        "visual_review": "not performed by this script",
        "scope": "Checks declared objects, not recovery of every source pixel or visual fidelity.",
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--scene", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument(
        "--strict",
        action="store_true",
        help="Return nonzero for warnings as well as errors",
    )
    args = parser.parse_args()
    try:
        report = audit(args.pptx, args.scene)
    except (ValueError, OSError) as error:
        parser.exit(2, f"audit_editability: {error}\n")
    data = json.dumps(report, ensure_ascii=True, indent=2)
    if args.output:
        if args.output.resolve() in {
            args.pptx.resolve(),
            args.scene.resolve() if args.scene else None,
        }:
            parser.error("report output must not overwrite an input")
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(data + "\n", encoding="utf-8")
    print(data)
    raise SystemExit(
        1 if report["errors"] or (args.strict and report["warnings"]) else 0
    )


if __name__ == "__main__":
    main()
