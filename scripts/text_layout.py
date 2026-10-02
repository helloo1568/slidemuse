"""Conservative native-text layout estimates; never a substitute for rendering."""
from __future__ import annotations

import math
import unicodedata

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE

EMU_PER_PT = 12700


def width_units(char):
    """Deliberately small advances, to flag strong density risks rather than fit."""
    if char == '\t':
        return None  # Tab stops and list indents are not modeled.
    category = unicodedata.category(char)
    if category in ('Mn', 'Mc', 'Me', 'Cf', 'So', 'Sk') or unicodedata.bidirectional(char) in ('R', 'AL'):
        return None
    if char.isspace():
        return .2
    if unicodedata.east_asian_width(char) in ('W', 'F'):
        return .85
    if char in 'ilI.,:;!\'`|':
        return .15
    if char in 'MWmw@%':
        return .65
    return .3


def estimate_frame(frame, width, height):
    """Read explicit PPTX styles, including mixed runs and hard/soft breaks."""
    result = {'status': 'unverified', 'reason': '', 'method': 'conservative-character-estimate'}
    if not frame.text.strip():
        return {**result, 'status': 'empty', 'reason': 'No visible text'}
    body = frame._txBody.bodyPr
    if (body.get('vert', 'horz') != 'horz' or body.get('numCol', '1') != '1'
            or body.get('rot', '0') != '0'):
        return {**result, 'reason': 'Vertical, multi-column or internally rotated text requires rendered review'}
    if frame.auto_size != MSO_AUTO_SIZE.NONE:
        return {**result, 'reason': 'Auto-fit or inherited auto-fit requires rendered review'}
    if frame.word_wrap is None:
        return {**result, 'reason': 'Inherited wrapping requires rendered review'}
    width = (width - frame.margin_left - frame.margin_right) / EMU_PER_PT
    height = (height - frame.margin_top - frame.margin_bottom) / EMU_PER_PT
    result.update(available_width_pt=round(width, 3), available_height_pt=round(height, 3))
    if width <= 0 or height <= 0:
        return {**result, 'status': 'risk', 'reason': 'Text margins consume the available box'}
    needed, widest = 0., 0.
    lines = 0
    for paragraph in frame.paragraphs:
        # Compiler writes these explicitly. Unresolved inherited styles must not
        # silently turn into a plausible default font/spacing and a clean result.
        if paragraph.font.size is None or paragraph.line_spacing is None:
            return {**result, 'reason': 'Inherited font size or line spacing requires rendered review'}
        psize = paragraph.font.size.pt
        if psize <= 0 or not math.isfinite(psize):
            return {**result, 'reason': 'Unsupported font size'}
        if paragraph._p.xpath('./a:pPr/a:buChar | ./a:pPr/a:buAutoNum'):
            return {**result, 'reason': 'Bullets and indentation require rendered review'}
        if paragraph._p.xpath('./a:pPr[@marL or @indent]'):
            return {**result, 'reason': 'Paragraph indentation requires rendered review'}
        tokens = []
        for node in paragraph._p:
            if node.tag.endswith('}br'):
                tokens.append(('\n', psize))
            elif node.tag.endswith('}r'):
                # A run's sz overrides the explicit paragraph size.
                size_nodes = node.xpath('./a:rPr/@sz')
                size = int(size_nodes[0]) / 100 if size_nodes else psize
                if size <= 0 or not math.isfinite(size):
                    return {**result, 'reason': 'Unsupported run font size'}
                value = ''.join(node.xpath('./a:t/text()'))
                tokens.extend((char, size) for char in value)
            elif node.tag.endswith('}fld'):
                return {**result, 'reason': 'Dynamic text fields require rendered review'}
        x, line_size, pheight = 0., psize, 0.
        spacing = paragraph.line_spacing

        def line_height(size, spacing=spacing):
            return spacing.pt if hasattr(spacing, 'pt') else size * spacing

        for char, size in tokens:
            if char in ('\n', '\v'):
                pheight += line_height(line_size)
                lines += 1
                widest = max(widest, x)
                x, line_size = 0., psize
                continue
            units = width_units(char)
            if units is None:
                return {**result, 'reason': 'Tabs, complex scripts or combining marks require rendered review'}
            advance = units * size
            if frame.word_wrap and x and x + advance > width:
                # Character wrapping packs tighter than word wrapping. This is
                # intentionally conservative; no claim of real Office metrics.
                pheight += line_height(line_size)
                lines += 1
                widest = max(widest, x)
                x, line_size = 0., psize
            x += advance
            line_size = max(line_size, size)
        pheight += line_height(line_size)
        lines += 1
        widest = max(widest, x)
        needed += pheight + (paragraph.space_before or 0) / EMU_PER_PT + (paragraph.space_after or 0) / EMU_PER_PT
    risk = needed > height * 1.25 or widest > width * 1.25
    return {**result, 'status': 'risk' if risk else 'no_obvious_risk',
            'reason': 'Conservative text estimate exceeds the box by more than 25%' if risk else 'No strong risk in this estimate; actual layout is unverified',
            'estimated_height_pt': round(needed, 3), 'estimated_width_pt': round(widest, 3), 'estimated_lines': lines}


def inspect_shape(shape):
    """Frames and individual table cells; grouped shapes are walked by caller."""
    findings = []
    if shape.has_text_frame:
        findings.append({'element_id': shape.name.split(' | ', 1)[0], **estimate_frame(shape.text_frame, shape.width, shape.height)})
    elif shape.has_table:
        for r, row in enumerate(shape.table.rows):
            for c, cell in enumerate(row.cells):
                if cell.is_spanned:
                    continue
                if cell.is_merge_origin:
                    item = {'status': 'unverified', 'reason': 'Merged cells require rendered review'}
                else:
                    item = estimate_frame(cell.text_frame, shape.table.columns[c].width, row.height)
                findings.append({'element_id': shape.name.split(' | ', 1)[0], 'cell': [r + 1, c + 1], **item})
    for item in findings:
        if shape.rotation:
            item.update(status='unverified', reason='Rotated text requires rendered review')
        for parent in shape._element.iterancestors():
            if parent.tag.endswith('}grpSp'):
                transforms = parent.xpath('./p:grpSpPr/a:xfrm')
                if not transforms:
                    item.update(status='unverified', reason='Inherited group transform requires rendered review')
                    continue
                transform = transforms[0]
                scaled = any(transform.xpath(f'./a:ext/@{axis}') != transform.xpath(f'./a:chExt/@{axis}') for axis in ('cx', 'cy'))
                if scaled or any(transform.get(key, '0') not in ('0', 'false') for key in ('rot', 'flipH', 'flipV')):
                    item.update(status='unverified', reason='Transformed group text requires rendered review')
    return [item for item in findings if item['status'] != 'empty']


def inspect_deck(presentation, slide_ids=None):
    def walk(shapes):
        for shape in shapes:
            yield shape
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(shape.shapes)

    findings = []
    for index, slide in enumerate(presentation.slides):
        for shape in walk(slide.shapes):
            for item in inspect_shape(shape):
                item.update(page_number=index + 1, slide_id=slide_ids[index] if slide_ids and index < len(slide_ids) else str(index + 1))
                findings.append(item)
    return findings


def risk_warning(item):
    cell = f"/cell {item['cell']}" if 'cell' in item else ''
    return f"Slide {item['page_number']}/{item['element_id']}{cell}: possible text overflow; inspect the actual render"
