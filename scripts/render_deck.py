#!/usr/bin/env python3
"""Render every PPTX slide and create a review sheet with optional source comparison."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from PIL import Image, ImageChops, ImageDraw, ImageOps, ImageStat
from pptx import Presentation


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def render_powerpoint(deck: Path, output: Path, width: int, height: int, slides: list[int] | None = None) -> None:
    shell = shutil.which("powershell.exe") or shutil.which("powershell")
    if not shell:
        raise RuntimeError("PowerPoint renderer requires Windows PowerShell")
    subprocess.run(
        [shell, "-NoProfile", "-NonInteractive", "-File", str(Path(__file__).with_name("render_powerpoint.ps1")),
         "-Deck", str(deck), "-OutputDir", str(output), "-Width", str(width), "-Height", str(height),
         *(["-SlideNumbers", ",".join(str(n) for n in slides)] if slides is not None else [])],
        check=True, capture_output=True, text=True,
    )


def render_libreoffice(deck: Path, output: Path, width: int, height: int, count: int,
                      slides: list[int] | None = None) -> None:
    office = shutil.which("soffice") or shutil.which("libreoffice")
    raster = shutil.which("pdftoppm")
    if not office or not raster:
        raise RuntimeError("LibreOffice renderer requires soffice and pdftoppm on PATH")
    with tempfile.TemporaryDirectory() as temp:
        temp_dir = Path(temp)
        subprocess.run(
            [office, "-env:UserInstallation=file:///" + temp_dir.as_posix(), "--headless", "--convert-to", "pdf", "--outdir", str(temp_dir), str(deck)],
            check=True, capture_output=True, text=True,
        )
        pdf = temp_dir / (deck.stem + ".pdf")
        if not pdf.is_file():
            raise RuntimeError("LibreOffice did not produce a PDF")
        wanted = slides if slides is not None else list(range(1, count + 1))
        for first, last in ([(1, count)] if slides is None else [(n, n) for n in wanted]):
            subprocess.run(
                [raster, "-f", str(first), "-l", str(last), "-png", "-r", "144", str(pdf), str(temp_dir / "slide")],
                check=True, capture_output=True, text=True,
            )
        pages = {int(path.stem.rsplit("-", 1)[1]): path for path in temp_dir.glob("slide-*.png")}
        if set(pages) != set(wanted):
            raise RuntimeError(f"LibreOffice rendered {len(pages)} of {count} slides")
        for index in wanted:
            page = pages[index]
            with Image.open(page) as image:
                ImageOps.pad(image.convert("RGB"), (width, height), method=Image.Resampling.LANCZOS, color="white").save(output / f"{index:03d}.png")


def render_selected(deck: Path, output: Path, width: int, backend: str,
                    slides: list[int]) -> str:
    """Render selected current-deck indices into a fresh directory, preserving indices."""
    presentation = Presentation(deck)
    count = len(presentation.slides)
    if (not slides or len(set(slides)) != len(slides)
            or any(isinstance(n, bool) or not isinstance(n, int) or not 1 <= n <= count for n in slides)):
        raise ValueError("Selected slide indices must be unique and inside the deck")
    if backend not in ("auto", "powerpoint", "libreoffice"):
        raise ValueError("Unknown render backend")
    if output.exists() and any(output.iterdir()):
        raise ValueError("Selected render output must be empty")
    output.mkdir(parents=True, exist_ok=True)
    height = round(width * presentation.slide_height / presentation.slide_width)
    candidates = [backend] if backend != "auto" else (["powerpoint", "libreoffice"] if os.name == "nt" else ["libreoffice"])
    errors = []
    for candidate in candidates:
        try:
            if candidate == "powerpoint":
                render_powerpoint(deck, output, width, height, slides)
            else:
                render_libreoffice(deck, output, width, height, count, slides)
            for n in slides:
                with Image.open(output / f"{n:03d}.png") as page:
                    if page.size != (width, height):
                        raise ValueError("Selected renderer produced incorrect dimensions")
                    page.verify()
            return candidate
        except (OSError, ValueError, RuntimeError, subprocess.CalledProcessError) as error:
            errors.append(f"{candidate}: {error}")
            for n in slides:
                (output / f"{n:03d}.png").unlink(missing_ok=True)
    raise RuntimeError("No presentation renderer succeeded: " + "; ".join(errors))


def reference_files(spec_path: Path, count: int) -> list[Path]:
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    slides = spec.get("slides", [])
    if len(slides) != count:
        raise ValueError("Page Spec slide count differs from the PPTX")
    result = []
    for slide in slides:
        if slide.get("image_status") not in ("approved", "revision"):
            raise ValueError(f"Reference slide {slide.get('id')} has no approved or historical image")
        source = (spec_path.parent / slide["image_file"]).resolve()
        if not source.is_relative_to(spec_path.parent.resolve()):
            raise ValueError(f"Reference image escapes Page Spec directory: {slide['image_file']}")
        if not source.is_file():
            raise ValueError(f"Reference image is missing: {source}")
        result.append(source)
    return result


def build_review(images: list[Path], output: Path, references: list[Path] | None) -> list[dict]:
    with Image.open(images[0]) as first:
        thumb_w = max(1, min(640, round(960 * first.width / first.height)))
        thumb_h = max(1, min(960, round(thumb_w * first.height / first.width)))
    rows = []
    report = []
    for index, path in enumerate(images, 1):
        with Image.open(path) as opened:
            rendered = ImageOps.pad(opened.convert("RGB"), (thumb_w, thumb_h), method=Image.Resampling.LANCZOS, color="white")
            original_size = list(opened.size)
        parts = [rendered]
        item = {"slide": index, "rendered": str(path), "rendered_sha256": sha256(path), "size": original_size}
        if references:
            with Image.open(references[index - 1]) as opened:
                source = ImageOps.pad(ImageOps.exif_transpose(opened).convert("RGB"), (thumb_w, thumb_h), method=Image.Resampling.LANCZOS, color="white")
            diff = ImageChops.difference(rendered, source)
            item["reference"] = str(references[index - 1])
            item["reference_sha256"] = sha256(references[index - 1])
            item["mean_absolute_difference"] = round(sum(ImageStat.Stat(diff).mean) / 3, 2)
            parts = [source, rendered, diff]
        row = Image.new("RGB", (thumb_w * len(parts), thumb_h + 28), "#f3f4f6")
        draw = ImageDraw.Draw(row)
        draw.text((8, 7), f"Slide {index:03d}" + ("  |  source / render / difference" if references else ""), fill="#172b4d")
        for column, part in enumerate(parts):
            row.paste(part, (column * thumb_w, 28))
        rows.append(row)
        report.append(item)
    sheet = Image.new("RGB", (rows[0].width, sum(row.height for row in rows)), "white")
    y = 0
    for row in rows:
        sheet.paste(row, (0, y))
        y += row.height
    sheet.save(output / "review.png")
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("deck", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--page-spec", type=Path, help="Approved images for a visual comparison sheet")
    parser.add_argument("--backend", choices=("auto", "powerpoint", "libreoffice"), default="auto")
    parser.add_argument("--width", type=int, default=1600)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    deck = args.deck.resolve()
    output = args.output.resolve()
    if not deck.is_file() or args.width < 320:
        parser.error("deck must exist and width must be at least 320 pixels")
    presentation = Presentation(deck)
    count = len(presentation.slides)
    if not count:
        parser.error("PPTX has no slides")
    height = round(args.width * presentation.slide_height / presentation.slide_width)
    if output.exists() and (not output.is_dir() or any(output.iterdir())) and not args.overwrite:
        parser.error("output directory is not empty; pass --overwrite to replace its review files")
    output.mkdir(parents=True, exist_ok=True)
    for path in output.glob("[0-9][0-9][0-9].png"):
        path.unlink()
    backends = [args.backend] if args.backend != "auto" else (["powerpoint", "libreoffice"] if os.name == "nt" else ["libreoffice"])
    errors = []
    backend = None
    for candidate in backends:
        try:
            if candidate == "powerpoint":
                render_powerpoint(deck, output, args.width, height)
            else:
                render_libreoffice(deck, output, args.width, height, count)
            backend = candidate
            break
        except (OSError, RuntimeError, subprocess.CalledProcessError) as error:
            errors.append(f"{candidate}: {error}")
            for path in output.glob("[0-9][0-9][0-9].png"):
                path.unlink()
    if backend is None:
        raise SystemExit("No presentation renderer succeeded: " + "; ".join(errors))
    images = [output / f"{index:03d}.png" for index in range(1, count + 1)]
    if any(not path.is_file() for path in images):
        raise SystemExit("Renderer output is incomplete")
    references = reference_files(args.page_spec.resolve(), count) if args.page_spec else None
    slides = build_review(images, output, references)
    report = {"deck": str(deck), "deck_sha256": sha256(deck), "backend": backend,
              "page_spec_sha256": sha256(args.page_spec.resolve()) if args.page_spec else None,
              "slides": slides, "review_sheet": str(output / "review.png"),
              "note": "Pixel differences are diagnostic only; inspect typography, clipping, content and composition manually."}
    (output / "render-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"backend": backend, "slides": count, "review": str(output / "review.png")}, ensure_ascii=False))


if __name__ == "__main__":
    main()
