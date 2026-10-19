"""Guarded generation contracts, image review, and lossless image-deck assembly."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import runpy
import tempfile
from pathlib import Path

from audit_page_content import text_present
from content_density import validate_content, visible_text
from filelock import FileLock
from json_io import read_json
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.util import Inches

CONTRACT = "slidemuse-content.json"
REVIEW = "content-review.json"
CHECKS = ("evidence_accuracy", "readability", "density", "visual_relationships")


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(dir=path.parent, suffix=".json")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            json.dump(value, stream, ensure_ascii=False, indent=2)
        os.replace(temporary, path)
    finally:
        Path(temporary).unlink(missing_ok=True)


def inside(root, value):
    root = root.resolve()
    path = (root / value).resolve()
    if not path.is_relative_to(root):
        raise ValueError(f"Path escapes generation deck: {value}")
    return path


def prepare_spec(spec):
    """Preserve exact approved content inside every upstream prompt."""
    import copy
    spec = copy.deepcopy(spec)
    if not isinstance(spec.get("slides"), list) or not spec["slides"]:
        raise ValueError("Generation spec requires slides")
    for number, slide in enumerate(spec["slides"], 1):
        if not isinstance(slide, dict):
            raise TypeError("Generation slide must be an object")
        if "number" in slide and (type(slide["number"]) is not int or slide["number"] != number):
            raise ValueError("Generation pages must be consecutive and ordered")
        slide["number"] = number
        validate_content(slide, required=True)
        context = slide.get("local_context", {})
        slide["local_context"] = {"original_context": context, "approved_content": {
            key: slide[key] for key in ("id", "core_message", "key_points", "explanation", "boundary",
                                       "source_refs", "density", "elements", "display", "required_visible_values", "speaker_notes")
            if key in slide}, "required_visible_text": visible_text(slide)}
        slide["constraints"] = [*slide.get("constraints", []),
            ("Render every required_visible_text item exactly; keep evidence, explanation and boundaries visible. "
             "When display.coverage is supplied, approved semantic fields describe meaning carried by mapped elements; "
             "do not repeat them as extra bullets. Review semantic coverage separately. "
             "Keep speaker_notes off the image. Do not replace substantive content with decoration or invent evidence.")]
    return spec


def from_page_spec(path):
    from validate_page_spec import load_page_spec
    spec, _ = load_page_spec(path, strict=True, content=True)
    for slide in spec["slides"]:
        if any("preserve_asset" in e for e in slide["elements"]):
            raise ValueError("Protected evidence requires the Page Spec production pipeline: use slidemuse.py produce prepare; do not regenerate its pixels through generate")
    result = {"deck_name": spec["title"], "language": "Chinese", "goal": spec["title"],
            "selected_image_backend": "built-in image tool", "style": spec["style"],
            "content_version": spec["content_version"], "content_authorization": spec["content_authorization"],
            "slides": [{**s, "number": s["page_number"], "local_context": {"page_spec_sha256": sha(path)}}
                       for s in spec["slides"]]}
    references = spec["style"].get("tokens", {}).get("reference_images", [])
    if references:
        result["approved_style_reference"] = {"path": str((path.parent / references[0]).resolve()),
                                               "role": "style-only; preserve approved content"}
    return result


def save_contract(deck, spec):
    jobs_path = deck / "slide_jobs.json"
    with FileLock(str(deck / ".slide_jobs.json.lock")):
        jobs = read_json(jobs_path)
        pages = []
        for entry, slide in zip(jobs["slides"], sorted(spec["slides"], key=lambda s: s.get("number", 0))):
            if entry["status"] == "accepted":
                image = inside(deck, entry["out"])
                entry["result"]["final_image_sha256"] = sha(image)
            pages.append({"id": entry["slide_id"], "number": entry["number"],
                          "expected_text": visible_text(slide), "content": slide,
                          "job_sha256": sha(inside(deck, entry["job"]))})
        jobs["slidemuse_content_policy"] = "1.0"
        write(jobs_path, jobs)
        write(deck / CONTRACT, {"version": "1.0", "pages": pages,
                               "scope": "Declared content; sources and readability require actual review."})


def checked_images(deck, jobs, ratio=None):
    slides = jobs.get("slides", [])
    if not slides:
        raise ValueError("No generation jobs; prepare and record all pages first")
    if [s.get("number") for s in slides] != list(range(1, len(slides) + 1)):
        raise ValueError("Generation pages must be consecutive and ordered")
    images = []
    for slide in slides:
        label = slide["slide_id"]
        if label != f"slide_{slide['number']:02d}":
            raise ValueError("Generation page ID differs from page number")
        if slide.get("status") not in ("recorded", "accepted"):
            raise ValueError(f"{label}: unfinished generation job ({slide.get('status')})")
        path = inside(deck, slide["out"])
        if not path.is_file() or (slide.get("result") or {}).get("final_image_sha256") != sha(path):
            raise ValueError(f"{label}: missing image or recorded image hash changed")
        with Image.open(path) as im:
            im.verify()
        if ratio:
            with Image.open(path) as im:
                w, h = ImageOps.exif_transpose(im).size
                if abs(w - h * ratio) > max(1, ratio) + 1e-6:
                    raise ValueError(f"{label}: image aspect ratio differs from export")
        images.append(path)
    actual = {p.resolve() for p in (deck / "origin_image").iterdir()
              if p.is_file() and re.fullmatch(r"slide_\d+\.(png|jpe?g|gif|bmp)", p.name, re.IGNORECASE)}
    if actual != set(images):
        raise ValueError("Missing, extra or duplicate final slide images")
    return images


def check_contract(deck, jobs):
    contract = read_json(deck / CONTRACT)
    pages = contract.get("pages", [])
    if [p["id"] for p in pages] != [s["slide_id"] for s in jobs["slides"]]:
        raise ValueError("Content contract page order differs from jobs")
    for page, slide in zip(pages, jobs["slides"]):
        if page["job_sha256"] != sha(inside(deck, slide["job"])):
            raise ValueError(f"{page['id']}: prompt changed; refresh approved content and review")
    return contract


def review(deck, record, initialize=False):
    jobs = read_json(deck / "slide_jobs.json")
    images = checked_images(deck, jobs)
    contract = check_contract(deck, jobs)
    if initialize:
        if record.exists():
            raise ValueError("Review record already exists; retain it and choose a new path")
        write(record, {"version": "1.0", "contract_sha256": sha(deck / CONTRACT), "slides": [
            {"id": p["id"], "image_sha256": sha(image), "observed_text": "", "notes": "",
             "checks": dict.fromkeys(CHECKS, "pending")} for p, image in zip(contract["pages"], images)]})
        return {"status": "awaiting_review", "record": str(record)}
    data = read_json(record)
    if data.get("version") != "1.0" or data.get("contract_sha256") != sha(deck / CONTRACT):
        raise ValueError("Stale or unsupported content review")
    if [r.get("id") for r in data.get("slides", [])] != [p["id"] for p in contract["pages"]]:
        raise ValueError("Review must cover every page once, in order")
    findings = []
    for page, image, item in zip(contract["pages"], images, data["slides"]):
        if item.get("image_sha256") != sha(image):
            raise ValueError(f"{page['id']}: stale image review")
        text = item.get("observed_text", "")
        if not isinstance(text, str):
            raise TypeError("observed_text must be a transcription string")
        missing = [s for s in page["expected_text"] if not text_present(s, text)]
        checks = item.get("checks", {})
        passed = all(checks.get(c) == "pass" for c in CHECKS)
        if missing or not passed or not str(item.get("notes", "")).strip():
            findings.append({"id": page["id"], "missing_text": missing, "checks": checks})
    return {"status": "pass" if not findings else "incomplete_or_failed", "findings": findings,
            "scope": "Matches declared content against recorded observations; does not OCR or authenticate evidence."}


def assemble(argv, vendor_script):
    parser = argparse.ArgumentParser(description="Assemble recorded, unchanged images; retain exact source media.")
    parser.add_argument("base_dir", type=Path)
    parser.add_argument("output")
    parser.add_argument("--aspect-ratio", "--ar", choices=("16:9", "4:3"), default="16:9")
    args = parser.parse_args(argv)
    name = args.output if args.output.lower().endswith(".pptx") else args.output + ".pptx"
    if Path(name).name != name or any(c in name for c in ("/", "\\")):
        raise ValueError("Output must be a PPTX filename, not a path")
    deck = inside(args.base_dir, Path(name).stem)
    if not (deck / "slide_jobs.json").is_file():
        raise ValueError("Missing slide_jobs.json; prepare and record every page before assembly")
    with FileLock(str(deck / ".slide_jobs.json.lock")):
        jobs = read_json(deck / "slide_jobs.json")
        ratio = 16 / 9 if args.aspect_ratio == "16:9" else 4 / 3
        images = checked_images(deck, jobs, ratio)
        if jobs.get("slidemuse_content_policy") == "1.0" and not (deck / CONTRACT).is_file():
            raise ValueError("Current generation job is missing its content contract")
        if (deck / CONTRACT).exists():
            report = review(deck, deck / REVIEW)
            if report["status"] != "pass":
                raise ValueError("Content review incomplete or failed; run generate review and inspect findings")
        notes = runpy.run_path(str(vendor_script))["load_speaker_notes"](str(deck))
        before = [sha(p) for p in images]
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(10), Inches(10 / ratio)
        for i, path in enumerate(images, 1):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            picture = slide.shapes.add_picture(str(path), 0, 0, width=prs.slide_width, height=prs.slide_height)
            if hashlib.sha256(picture.image.blob).hexdigest() != before[i - 1]:
                raise ValueError("Image embedding changed source media")
            slide.notes_slide.notes_text_frame.text = notes.get(i, "")
        target = deck / name
        fd, temporary = tempfile.mkstemp(dir=deck, suffix=".pptx")
        os.close(fd)
        try:
            prs.save(temporary)
            reopened = Presentation(temporary)
            if len(reopened.slides) != len(images) or before != [sha(p) for p in images]:
                raise ValueError("Images changed during assembly")
            for i, slide in enumerate(reopened.slides, 1):
                if (hashlib.sha256(slide.shapes[0].image.blob).hexdigest() != before[i - 1]
                        or slide.notes_slide.notes_text_frame.text != notes.get(i, "")):
                    raise ValueError("Exported media/notes differ from input")
            os.replace(temporary, target)
        finally:
            Path(temporary).unlink(missing_ok=True)
    print(json.dumps({"output": str(target), "slides": len(images),
                      "content_review": "pass" if (deck / CONTRACT).exists() else "legacy_not_checked",
                      "visual_review": "Render the actual PPTX in Office before delivery."}, ensure_ascii=True))
    return 0
