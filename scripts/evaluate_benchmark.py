#!/usr/bin/env python3
"""Freeze a material batch and preserve independently reviewed delivery attempts.

This validates evidence and statistics, not reviewer identity or semantic truth.
Keep each attempt's artifacts in its own directory; changed evidence is rejected.
"""

import argparse
import hashlib
import io
import json
import math
import re
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

from audit_editability import audit as audit_editability
from evaluate_delivery import evaluate
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from scene import load_scene, walk
from validate_page_spec import load_page_spec

CATEGORIES = ("paper", "statistics", "policy", "teaching", "project", "reconstruction")
DIMENSIONS = ("content", "facts_sources", "visual_data", "editing_updates", "delivery_reproduction")
REQUIRED_RECORDS = ("content_draft", "authorizations", "run_log", "edit_check")
COMMIT = re.compile(r"[0-9a-f]{40}")
IDENTIFIER = re.compile(r"[a-z0-9][a-z0-9_-]{0,79}")
REVIEW_CONTRACT_NOTE = (
    "Content requirements only, not a factual verdict. Verify facts, formulas, values and "
    "source relationships against the original source materials. Unknown data fields are "
    "omitted and counted; the original Page Spec remains the evaluation input."
)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def object_digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                                     separators=(",", ":")).encode("utf-8")).hexdigest()


def read(path):
    value = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise TypeError(f"Expected JSON object: {path}")
    return value


def resolve(base, value):
    if not isinstance(value, str) or not value.strip():
        raise ValueError("Evidence path must be a nonempty string")
    return (base / value).resolve()


def canonical_source(value):
    """Normalize explicit aliases, including common DOI and arXiv versions."""
    if not isinstance(value, str) or not value.strip():
        raise ValueError("Missing canonical material identity")
    value = value.strip().casefold()
    value = re.sub(r"^https?://(?:www\.)?doi\.org/", "doi:", value)
    value = re.sub(r"^https?://(?:www\.)?arxiv\.org/(?:abs|pdf|html)/", "arxiv:", value)
    if value.startswith("arxiv:"):
        value = value.split("?", 1)[0].split("#", 1)[0]
        return re.sub(r"v\d+(?:\.pdf)?$|\.pdf$", "", value)
    if value.startswith(("http://", "https://")):
        url = urlsplit(value)
        # Fragments are locations within the same source, not new materials.
        return urlunsplit((url.scheme, url.netloc, url.path.rstrip("/"), url.query, ""))
    return value


def seal_write(path, value):
    """Exclusive creation prevents accidentally replacing first results."""
    value = {**value, "record_sha256": object_digest(value)}
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8") as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2)
        stream.write("\n")
    return value


def sealed_read(path):
    value = read(path)
    expected = value.pop("record_sha256", None)
    if expected != object_digest(value):
        raise ValueError(f"Record was changed or is unsealed: {path}")
    return {**value, "record_sha256": expected}


def _review_cells(value):
    """Copy scalar data or arrays of scalars; never pass through arbitrary objects."""
    if value is None or type(value) in (str, bool, int):
        return value
    if type(value) is float and math.isfinite(value):
        return value
    if isinstance(value, list):
        return [_review_cells(item) for item in value]
    raise ValueError("Review contract data requires finite scalar cells or cell arrays")


def _review_data(value):
    """Project the open Page Spec data field using explicit table/chart keys."""
    if isinstance(value, list):
        return _review_cells(value), 0
    if not isinstance(value, dict):
        raise TypeError("Review contract data must be a table/chart object or cell array")
    fields = ("chart_type", "categories", "labels", "columns", "headers", "rows",
              "values", "baseline", "unit", "units")
    projected = {key: _review_cells(value[key]) for key in fields if key in value}
    omitted = len(set(value) - {*fields, "series"})
    if "series" in value:
        if not isinstance(value["series"], list) or not value["series"]:
            raise ValueError("Review contract series must be a nonempty array")
        projected["series"] = []
        for series in value["series"]:
            if not isinstance(series, dict) or "values" not in series:
                raise ValueError("Review contract series requires explicit values")
            projected["series"].append({key: _review_cells(series[key])
                                        for key in ("name", "values") if key in series})
            omitted += len(set(series) - {"name", "values"})
    if not any(key in projected for key in ("rows", "values", "series")):
        raise ValueError("Review contract data has no supported rows, values or series")
    return projected, omitted


def project_review_contract(spec):
    """Return only declared semantic fields, excluding prompts and work judgments.

    The caller supplies a validated Page Spec. Formulas are retained verbatim in
    text; speaker_notes are declared delivery content, while notes are work logs.
    No unknown mapping is copied, including mappings nested inside data cells.
    """
    contract = {"title": spec["title"],
                "canvas": {key: spec["canvas"][key] for key in ("width", "height")},
                "slides": [], "omitted_data_fields": 0}
    for slide in spec["slides"]:
        page = {key: slide[key] for key in ("id", "page_number", "title", "core_message",
                                          "source_refs", "required_visible_values", "speaker_notes")
                if key in slide}
        page["elements"] = []
        for element in slide["elements"]:
            item = {key: element[key] for key in ("id", "kind", "role", "source_ref", "text")
                    if key in element}
            if element["kind"] in ("table", "chart"):
                item["data"], omitted = _review_data(element["data"])
                contract["omitted_data_fields"] += omitted
            page["elements"].append(item)
        contract["slides"].append(page)
    return contract


def _review_contract_record(spec_path):
    spec_path = Path(spec_path).resolve()
    before = digest(spec_path)
    try:
        spec, _ = load_page_spec(spec_path)
    except (ValueError, TypeError, KeyError, OSError):
        # Schema diagnostics may include a rejected value containing a prompt.
        raise ValueError("Review contract requires a valid readable Page Spec") from None
    projection = project_review_contract(spec)
    if digest(spec_path) != before:
        raise ValueError("Page Spec changed while preparing its review contract")
    return {"version": "1.0", "kind": "page-spec-review-contract",
            "page_spec_sha256": before, "projection_sha256": object_digest(projection),
            "contract": projection, "note": REVIEW_CONTRACT_NOTE}


def create_review_contract(spec_path, output):
    """Exclusively create a sealed projection without changing the Page Spec."""
    output = Path(output).resolve()
    if output == Path(spec_path).resolve():
        raise ValueError("Review contract must not replace its Page Spec")
    return seal_write(output, _review_contract_record(spec_path))


def verify_review_contract(contract_path, spec_path):
    """Verify sealing, source bytes and the complete current whitelist projection.

    This also rejects extra fields even if a caller recalculates the seal. It is
    integrity verification, not a source-material factual check or a signature.
    """
    value = sealed_read(Path(contract_path))
    expected = _review_contract_record(spec_path)
    actual = {key: item for key, item in value.items() if key != "record_sha256"}
    if object_digest(actual) != object_digest(expected):
        raise ValueError("Review contract differs from its current Page Spec projection")
    return value


def freeze(registry_path, output):
    registry_path, output = Path(registry_path).resolve(), Path(output).resolve()
    registry = read(registry_path)
    if registry.get("version") != "1.0" or not COMMIT.fullmatch(registry.get("candidate_commit", "")):
        raise ValueError("Registry requires version 1.0 and a full candidate commit")
    cases = registry.get("cases")
    if not isinstance(cases, list) or not cases:
        raise ValueError("Register the entire batch before evaluation")
    ids, sources = set(), set()
    for case in cases:
        if not isinstance(case, dict) or not IDENTIFIER.fullmatch(case.get("id", "")) or case["id"] in ids:
            raise ValueError("Missing, invalid or duplicate case ID")
        if case.get("category") not in CATEGORIES or case.get("split") not in ("development", "holdout"):
            raise ValueError("Unknown category or split")
        aliases = case.get("source_aliases", [])
        if not isinstance(aliases, list):
            raise TypeError("source_aliases must be an array")
        identities = {canonical_source(item) for item in [case.get("source_id"), *aliases]}
        if identities & sources:
            raise ValueError("Duplicate material or overlapping source aliases")
        ids.add(case["id"])
        sources.update(identities)
        if case["split"] == "holdout" and case.get("used_for_tuning") is not False:
            raise ValueError("Holdout requires an explicit unused-for-tuning declaration")
        if case.get("delivery_scope") not in ("image_and_editable", "reconstruction"):
            raise ValueError("Benchmark requires editable delivery; declare its actual entry point")
    rubric = resolve(registry_path.parent, registry.get("rubric"))
    if not rubric.is_file() or not rubric.read_text(encoding="utf-8").strip():
        raise ValueError("An actual predeclared rubric is required")
    registry["rubric"] = str(rubric)
    return seal_write(output, {"version": "1.0", "registry": registry,
                               "rubric_sha256": digest(rubric),
                               "frozen_at": datetime.now(timezone.utc).isoformat()})


def batch(path):
    value = sealed_read(path)
    if digest(value["registry"]["rubric"]) != value["rubric_sha256"]:
        raise ValueError("Frozen rubric changed; create a new batch, preserve old results")
    return value


def verify_image_deck(deck, spec, report):
    presentation = Presentation(deck)
    if not math.isclose(presentation.slide_width / presentation.slide_height,
                        spec["canvas"]["width"] / spec["canvas"]["height"], rel_tol=1e-6):
        raise ValueError("Image deck canvas differs from Page Spec")
    if len(presentation.slides) != len(spec["slides"]):
        raise ValueError("Image deck slide count differs from Page Spec")
    for slide, page in zip(presentation.slides, report["slides"]):
        if len(slide.shapes) != 1 or slide.shapes[0].shape_type != MSO_SHAPE_TYPE.PICTURE:
            raise ValueError("Image delivery must carry exactly one approved page image per slide")
        picture = slide.shapes[0]
        with Image.open(page["reference"]) as reference, Image.open(io.BytesIO(picture.image.blob)) as embedded:
            original = ImageOps.exif_transpose(reference).convert("RGBA")
            actual = ImageOps.exif_transpose(embedded).convert("RGBA")
            if original.size != actual.size or original.tobytes() != actual.tobytes():
                raise ValueError("Image deck picture differs from approved reference pixels")
            image_ratio = original.width / original.height
        sw, sh = presentation.slide_width, presentation.slide_height
        ratio = sw / sh
        cover_crop = ((1 - ratio / image_ratio) / 2, 0) if image_ratio > ratio else (0, (1 - image_ratio / ratio) / 2)
        cover = (abs(picture.left) <= 2 and abs(picture.top) <= 2
                 and abs(picture.width - sw) <= 2 and abs(picture.height - sh) <= 2
                 and all(math.isclose(observed, expected, abs_tol=1e-5) for observed, expected in zip(
                     (picture.crop_left, picture.crop_right, picture.crop_top, picture.crop_bottom),
                     (cover_crop[0], cover_crop[0], cover_crop[1], cover_crop[1]))))
        cw, ch = (sw, int(sw / image_ratio)) if image_ratio >= ratio else (int(sh * image_ratio), sh)
        contain = (abs(picture.left - int((sw - cw) / 2)) <= 2 and abs(picture.top - int((sh - ch) / 2)) <= 2
                   and abs(picture.width - cw) <= 2 and abs(picture.height - ch) <= 2
                   and all(value == 0 for value in (picture.crop_left, picture.crop_right, picture.crop_top, picture.crop_bottom)))
        if not cover and not contain:
            raise ValueError("Image deck must fit the complete approved page with expected contain/cover geometry")


def translated_object(before, after):
    """Recognize a uniform move, including absolute-coordinate group children."""
    deltas = {"x": [], "y": []}

    def compare(old, new):
        if old.get("type") != new.get("type") or old.get("id") != new.get("id"):
            return False
        old_rest, new_rest = dict(old), dict(new)
        if old["type"] == "group":
            children_old, children_new = old_rest.pop("children"), new_rest.pop("children")
            if len(children_old) != len(children_new) or not all(compare(a, b) for a, b in zip(children_old, children_new)):
                return False
        else:
            for axis in ("x", "y"):
                for field in (axis, axis + "1", axis + "2"):
                    if field in old_rest:
                        if field not in new_rest:
                            return False
                        deltas[axis].append(new_rest.pop(field) - old_rest.pop(field))
        return old_rest == new_rest

    if not compare(before, after) or not all(deltas.values()):
        return False
    return all(all(math.isclose(value, items[0], abs_tol=1e-7) for value in items)
               for items in deltas.values()) and any(abs(items[0]) > 1e-7 for items in deltas.values())


def validate_work_records(paths, case, bind, delivery_scene):
    authorization = read(paths["authorizations"])
    if (not isinstance(authorization.get("user_quote"), str) or not authorization["user_quote"].strip()
            or not isinstance(authorization.get("stages"), dict)
            or any(authorization["stages"].get(key) is not True for key in ("content", "style", "editable"))):
        raise ValueError("Recorded content/style/editable authorization required")
    log = read(paths["run_log"])
    events = log.get("events")
    elapsed = log.get("elapsed_seconds")
    if (log.get("complete") is not True or not isinstance(events, list) or not events
            or type(elapsed) not in (int, float) or not math.isfinite(elapsed) or elapsed < 0):
        raise ValueError("Complete run ledger and observed duration required")
    for event in events:
        if (not isinstance(event, dict) or event.get("stage") not in
                ("style", "initial_page", "repair", "program_fallback", "build", "render", "edit")
                or not isinstance(event.get("evidence"), str) or not event["evidence"].strip()
                or event.get("outcome") not in ("pass", "fail")):
            raise ValueError("Each run event needs its real stage, outcome and evidence locator")
        evidence = bind(resolve(paths["run_log"].parent, event["evidence"]))
        if not evidence.stat().st_size:
            raise ValueError("Empty invocation evidence")
    if not any(event["stage"] == "render" for event in events):
        raise ValueError("Rendering invocation missing from run ledger")
    if case["delivery_scope"] != "reconstruction" and (
            sum(event["stage"] == "style" for event in events) < 4
            or not any(event["stage"] == "initial_page" for event in events)):
        raise ValueError("Image workflow requires four style calls and initial page calls")
    checks = read(paths["edit_check"])
    operations = checks.get("operations")
    if not isinstance(operations, list) or not operations:
        raise ValueError("Actual editing exercise records required")
    kinds = set()
    for operation in operations:
        if (not isinstance(operation, dict) or operation.get("kind") not in
                ("text_edit", "object_move", "chart_data_edit", "table_data_edit")
                or operation.get("status") != "pass" or not operation.get("notes")):
            raise ValueError("Editing exercise needs a successful actual operation and explanation")
        kinds.add(operation["kind"])
        before = bind(resolve(paths["edit_check"].parent, operation.get("before_scene")))
        if digest(before) != digest(delivery_scene):
            raise ValueError("Editing exercise must start from the delivered Scene")
        after = bind(resolve(paths["edit_check"].parent, operation.get("after_scene")))
        deck = bind(resolve(paths["edit_check"].parent, operation.get("after_deck")))
        render_path = bind(resolve(paths["edit_check"].parent, operation.get("render_report")))
        scene_before, _ = load_scene(before)
        scene_after, _ = load_scene(after)
        eid, sid = operation.get("element_id"), operation.get("slide_id")
        def find(scene, slide_id=sid, element_id=eid):
            return next((element for slide in scene["slides"] if slide["id"] == slide_id
                         for element in walk(slide["elements"]) if element["id"] == element_id), None)
        old, new = find(scene_before), find(scene_after)
        if old is None or new is None or old["type"] != new["type"]:
            raise ValueError("Editing exercise must preserve an existing element identity")
        fields = {"text_edit": ("text",), "object_move": ("x", "y", "x1", "x2", "y1", "y2"),
                  "chart_data_edit": ("series",), "table_data_edit": ("rows",)}[operation["kind"]]
        expected_type = {"text_edit": "text", "chart_data_edit": "chart", "table_data_edit": "table"}
        if (operation["kind"] in expected_type and old["type"] != expected_type[operation["kind"]]) or not any(
                old.get(field) != new.get(field) for field in fields) and operation["kind"] != "object_move":
            raise ValueError("Editing exercise did not perform its declared change")
        if operation["kind"] == "object_move" and not translated_object(old, new):
            raise ValueError("Object movement must actually translate the whole object without altering its content")
        if operation["kind"] == "chart_data_edit" and (
                [series["values"] for series in old["series"]] == [series["values"] for series in new["series"]]):
            raise ValueError("Chart data exercise must change actual series values")
        audit = audit_editability(deck, after)
        if audit["errors"]:
            raise ValueError("Modified deck differs from its Scene")
        report = read(render_path)
        if (Path(report.get("deck", "")).resolve() != deck or report.get("deck_sha256") != digest(deck)
                or report.get("backend") not in ("powerpoint", "libreoffice")
                or len(report.get("slides", [])) != len(scene_after["slides"])):
            raise ValueError("Edited deck requires actual current rendering")
        rendered_paths = set()
        for number, page in enumerate(report["slides"], 1):
            rendered_path = Path(page["rendered"]).resolve()
            if page.get("slide") != number or rendered_path in rendered_paths:
                raise ValueError("Edited rendering pages must be unique and in slide order")
            rendered_paths.add(rendered_path)
            if digest(bind(page["rendered"])) != page.get("rendered_sha256"):
                raise ValueError("Edited rendering changed")
        for scene_path, scene_data in ((before, scene_before), (after, scene_after)):
            for slide in scene_data["slides"]:
                if slide.get("source_image"):
                    bind(resolve(scene_path.parent, slide["source_image"]))
                for element in walk(slide["elements"]):
                    if element["type"] == "image":
                        bind(resolve(scene_path.parent, element["path"]))
    if not {"text_edit", "object_move"} <= kinds:
        raise ValueError("Both text editing and object movement must be exercised")
    return kinds


def assess(frozen, attempt, base):
    registry = frozen["registry"]
    cases = {case["id"]: case for case in registry["cases"]}
    case = cases.get(attempt.get("case_id"))
    if case is None:
        raise ValueError("Attempt is not part of the frozen batch")
    result = {"status": "incomplete", "issues": [], "files": {}, "deliveries": {},
              "major_factual_errors": None, "reviewer_id": None}

    def bind(path):
        path = Path(path).resolve()
        result["files"][str(path)] = digest(path)
        return path

    if attempt.get("candidate_commit") != registry["candidate_commit"]:
        result["issues"].append("Candidate differs from frozen commit")
        return result
    if case["split"] == "holdout" and attempt.get("holdout_clean") is not True:
        result["issues"].append("Holdout contamination or missing declaration")
        return result
    if attempt.get("execution_status") == "aborted":
        result["issues"].append("Execution aborted; remains in denominator")
        return result
    if attempt.get("execution_status") != "delivered" or not attempt.get("producer_id"):
        result["issues"].append("No complete delivery or producer identity")
        return result
    try:
        sources = attempt.get("sources", [])
        if not isinstance(sources, list) or not sources:
            raise ValueError("Source material snapshots missing")
        for source in sources:
            path = bind(resolve(base, source))
            if not path.stat().st_size:
                raise ValueError("Empty source snapshot")
        records = attempt.get("records", {})
        record_paths = {}
        for key in REQUIRED_RECORDS:
            record_paths[key] = bind(resolve(base, records.get(key)))
            if not record_paths[key].stat().st_size:
                raise ValueError("Empty work record")
        additional = records.get("additional_evidence", [])
        if not isinstance(additional, list):
            raise TypeError("Additional evidence must be a path array")
        for path in additional:
            bind(resolve(base, path))
        deliveries = attempt.get("deliveries", {})
        required = ("editable",) if case["delivery_scope"] == "reconstruction" else ("image", "editable")
        if set(deliveries) != set(required):
            raise ValueError("Missing or unexpected delivery branch")
        delivery_scene = bind(resolve(base, deliveries["editable"].get("scene")))
        edit_kinds = validate_work_records(record_paths, case, bind, delivery_scene)
        for name in required:
            item = deliveries[name]
            spec = bind(resolve(base, item.get("page_spec")))
            render = bind(resolve(base, item.get("render_report")))
            review = bind(resolve(base, item.get("visual_review")))
            observations = item.get("observations", [])
            if not isinstance(observations, list) or not observations:
                raise ValueError("Per-page content observations missing")
            observations = [bind(resolve(base, path)) for path in observations]
            scene = bind(resolve(base, item["scene"])) if item.get("scene") else None
            if (name == "editable") != bool(scene):
                raise ValueError("Editable delivery needs Scene; image delivery must use its actual image deck")
            geometry = bind(resolve(base, item["chart_observations"])) if item.get("chart_observations") else None
            score = evaluate(spec, render, review, observations, scene, geometry)
            page_spec = read(spec)
            data_kinds = {element["kind"] for slide in page_spec["slides"] for element in slide["elements"]}
            if ("chart" in data_kinds and "chart_data_edit" not in edit_kinds) or (
                    "table" in data_kinds and "table_data_edit" not in edit_kinds):
                raise ValueError("Declared chart/table data edits were not exercised")
            if score["render_backend"] not in ("powerpoint", "libreoffice"):
                raise ValueError("Benchmark requires actual PowerPoint or LibreOffice rendering")
            report = read(render)
            if name == "image":
                verify_image_deck(score["deck"], page_spec, report)
            bind(score["deck"])
            for page in report["slides"]:
                bind(page["rendered"])
                bind(page["reference"])
            if scene:
                # Asset paths are already validated by audit_editability/Scene loader.
                scene_data, _ = load_scene(scene)
                for slide in scene_data.get("slides", []):
                    if slide.get("source_image"):
                        bind(resolve(scene.parent, slide["source_image"]))
                    for element in walk(slide["elements"]):
                        if element["type"] == "image":
                            bind(resolve(scene.parent, element["path"]))
            result["deliveries"][name] = score
    except (ValueError, TypeError, KeyError, OSError) as exc:
        result["issues"].append(str(exc))
        return result
    input_hash = object_digest({"files": result["files"], "case_id": case["id"],
                               "rubric_sha256": frozen["rubric_sha256"],
                               "candidate_commit": registry["candidate_commit"]})
    result["review_input_sha256"] = input_hash
    try:
        independent = read(bind(resolve(base, attempt.get("independent_review"))))
        if independent.get("version") != "1.0" or independent.get("input_sha256") != input_hash:
            raise ValueError("Independent review missing or bound to different inputs")
        reviewer = independent.get("reviewer_id")
        if (not isinstance(reviewer, str) or not reviewer.strip()
                or reviewer.strip().casefold() == str(attempt["producer_id"]).strip().casefold()
                or independent.get("independent") is not True):
            raise ValueError("Actual independent reviewer declaration required")
        result["reviewer_id"] = reviewer
        dimensions = independent.get("dimensions", {})
        if set(dimensions) != set(DIMENSIONS):
            raise ValueError("Five semantic review dimensions required")
        for dimension in dimensions.values():
            if (not isinstance(dimension, dict) or dimension.get("status") not in ("pass", "fail", "incomplete")
                    or not isinstance(dimension.get("notes"), str) or not dimension["notes"].strip()):
                raise ValueError("Dimension needs an actual judgment and explanation")
        errors = independent.get("major_factual_errors")
        if type(errors) is not int or errors < 0:
            raise ValueError("Major factual error count missing or invalid")
        result["major_factual_errors"] = errors
        result["independent_review"] = independent
        statuses = [item["status"] for item in result["deliveries"].values()]
        statuses.extend(item["status"] for item in dimensions.values())
        if independent.get("prior_conclusions_exposed") is not False:
            statuses.append("incomplete")
            result["issues"].append("Unexposed review declaration missing or prior conclusions were disclosed")
        result["status"] = ("fail" if errors or "fail" in statuses else
                            "incomplete" if "incomplete" in statuses else "pass")
    except (ValueError, TypeError, OSError) as exc:
        result["issues"].append(str(exc))
        if any(item["status"] == "fail" for item in result["deliveries"].values()):
            result["status"] = "fail"
    return result


def history(directory, frozen):
    results = []
    previous = None
    for number, path in enumerate(sorted(Path(directory).glob("attempt-*.json")), 1):
        record = sealed_read(path)
        if (path.name != f"attempt-{number:06d}.json" or record.get("sequence") != number
                or record.get("previous_sha256") != previous
                or record.get("batch_sha256") != frozen["record_sha256"]):
            raise ValueError("Attempt history is missing, reordered or belongs to another batch")
        results.append(record)
        previous = record["record_sha256"]
    head_path = Path(directory) / "head.json"
    if results or head_path.exists():
        head = sealed_read(head_path)
        if (head.get("count") != len(results) or head.get("tip_sha256") != previous
                or head.get("batch_sha256") != frozen["record_sha256"]):
            raise ValueError("Attempt history tail is missing or its head belongs to another batch")
    return results


def record_attempt(batch_path, attempt_path, directory):
    frozen = batch(batch_path)
    prior = history(directory, frozen)
    attempt_path = Path(attempt_path).resolve()
    attempt = read(attempt_path)
    result = assess(frozen, attempt, attempt_path.parent)
    number = len(prior) + 1
    result = seal_write(Path(directory) / f"attempt-{number:06d}.json", {
        "version": "1.0", "sequence": number, "batch_sha256": frozen["record_sha256"],
        "previous_sha256": prior[-1]["record_sha256"] if prior else None,
        "recorded_at": datetime.now(timezone.utc).isoformat(), "attempt": attempt,
        "base_directory": str(attempt_path.parent), "assessment": result,
    })
    next_head = Path(directory) / "head-next.json"
    seal_write(next_head, {"count": number, "tip_sha256": result["record_sha256"],
                           "batch_sha256": frozen["record_sha256"]})
    next_head.replace(Path(directory) / "head.json")
    return result


def summarize(batch_path, directory):
    frozen = batch(batch_path)
    cases = frozen["registry"]["cases"]
    first, final, original = {}, {}, {}
    for record in history(directory, frozen):
        case_id = record["attempt"]["case_id"]
        old = record["assessment"]
        changed = []
        for path, expected in old["files"].items():
            try:
                if digest(path) != expected:
                    changed.append(path)
            except OSError:
                changed.append(path)
        if changed:
            current = {"status": "incomplete", "major_factual_errors": None,
                       "issues": ["Recorded evidence changed or disappeared"], "changed_files": changed}
        else:
            current = assess(frozen, record["attempt"], Path(record["base_directory"]))
            # Newly created review/evidence cannot promote an already sealed attempt.
            if old["status"] != "pass" and current["status"] == "pass":
                current["status"] = old["status"]
                current["issues"].append("Later evidence requires a new attempt; original result retained")
        if case_id not in first:
            first[case_id], original[case_id] = current, old["status"]
        final[case_id] = current
    rows = [{"id": case["id"], "category": case["category"], "split": case["split"],
             "first_recorded_status": original.get(case["id"], "not_run"),
             "first": first.get(case["id"], {"status": "incomplete", "issues": ["Not run"]}),
             "final": final.get(case["id"], {"status": "incomplete", "issues": ["Not run"]})}
            for case in cases]

    def counts(group):
        passed = sum(row["first"]["status"] == "pass" for row in group)
        return {"total": len(group), "first_passed": passed,
                "first_usable_rate": passed / len(group) if group else None,
                "final_passed": sum(row["final"]["status"] == "pass" for row in group)}

    overall = counts(rows)
    categories = {key: counts([row for row in rows if row["category"] == key]) for key in CATEGORIES}
    splits = {key: counts([row for row in rows if row["split"] == key]) for key in ("development", "holdout")}
    final_errors = [row["final"].get("major_factual_errors") for row in rows]
    eligible = len(rows) >= 30 and splits["holdout"]["total"] >= 10 and all(
        value["total"] >= 5 for value in categories.values())
    complete = (eligible and overall["first_usable_rate"] >= .9
                and overall["final_passed"] == len(rows) and all(value == 0 for value in final_errors))
    return {"version": "1.0", "batch_sha256": frozen["record_sha256"], "overall": overall,
            "categories": categories, "splits": splits, "cases": rows,
            "ultimate_goal_eligible_batch": eligible, "ultimate_goal_checks_passed": complete,
            "note": "Checks rely on recorded independent judgments; the script cannot prove semantics or reviewer identity."}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    freeze_parser = sub.add_parser("freeze")
    freeze_parser.add_argument("registry", type=Path)
    freeze_parser.add_argument("output", type=Path)
    record_parser = sub.add_parser("record")
    record_parser.add_argument("batch", type=Path)
    record_parser.add_argument("attempt", type=Path)
    record_parser.add_argument("directory", type=Path)
    inspect_parser = sub.add_parser("inspect", help="Prepare review bindings without recording a delivery attempt")
    inspect_parser.add_argument("batch", type=Path)
    inspect_parser.add_argument("attempt", type=Path)
    inspect_parser.add_argument("output", type=Path)
    contract_parser = sub.add_parser("review-contract", help="Exclusively create a clean Page Spec content contract")
    contract_parser.add_argument("page_spec", type=Path)
    contract_parser.add_argument("output", type=Path)
    verify_parser = sub.add_parser("verify-review-contract", help="Verify the sealed contract against its original Page Spec")
    verify_parser.add_argument("page_spec", type=Path)
    verify_parser.add_argument("contract", type=Path)
    report_parser = sub.add_parser("report")
    report_parser.add_argument("batch", type=Path)
    report_parser.add_argument("directory", type=Path)
    report_parser.add_argument("output", type=Path)
    args = parser.parse_args()
    if args.command == "freeze":
        result = freeze(args.registry, args.output)
    elif args.command == "review-contract":
        result = create_review_contract(args.page_spec, args.output)
    elif args.command == "verify-review-contract":
        result = verify_review_contract(args.contract, args.page_spec)
    elif args.command == "record":
        result = record_attempt(args.batch, args.attempt, args.directory)
    elif args.command == "inspect":
        result = assess(batch(args.batch), read(args.attempt), args.attempt.resolve().parent)
        if args.output.resolve() in (args.batch.resolve(), args.attempt.resolve()) or str(args.output.resolve()) in result["files"]:
            parser.error("Inspection output must not replace batch, attempt or evidence")
        seal_write(args.output, result)
    else:
        result = summarize(args.batch, args.directory)
        # Also refuse to overwrite a frozen batch, record or evidence input.
        if args.output.resolve() == args.batch.resolve() or args.output.resolve().parent == args.directory.resolve():
            parser.error("Report output must be outside frozen batch/history")
        seal_write(args.output, result)
    print(json.dumps({key: result[key] for key in ("record_sha256", "page_spec_sha256", "projection_sha256", "status", "review_input_sha256", "overall", "ultimate_goal_checks_passed")
                      if key in result}, ensure_ascii=True))


if __name__ == "__main__":
    main()
