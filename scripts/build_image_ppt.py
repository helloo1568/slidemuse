#!/usr/bin/env python3
"""Build an image deck from an approved Page Spec or a legacy image directory."""
from __future__ import annotations

import argparse
import json
import math
import os
import re
import tempfile
from contextlib import nullcontext
from pathlib import Path

from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches
from speaker_notes import canonical_notes, check_notes
from validate_page_spec import load_page_spec

DEFAULT_EXTENSIONS = (".png", ".jpg", ".jpeg")
def natural_key(path: Path) -> list[object]:
    parts = re.split(r"(\d+)", path.name.lower())
    return [int(part) if part.isdigit() else part for part in parts]
def collect_images(input_dir: Path, extensions: tuple[str, ...]) -> list[Path]:
    normalized = {ext.lower() if ext.startswith(".") else f".{ext.lower()}" for ext in extensions}
    return sorted(
        (path for path in input_dir.iterdir() if path.is_file() and path.suffix.lower() in normalized),
        key=natural_key,
    )
def add_background(slide, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(color)
def image_size(image_path: Path) -> tuple[int, int]:
    with Image.open(image_path) as image:
        image = ImageOps.exif_transpose(image)
        return image.size
def prepare_image(
    image_path: Path,
    index: int,
    work_dir: Path,
    max_width: int,
    jpeg_quality: int,
    background: str,
) -> Path:
    """Return an embeddable image, optionally downscaled and re-encoded."""
    with Image.open(image_path) as image:
        image = ImageOps.exif_transpose(image)
        if max_width > 0 and image.width > max_width:
            new_height = round(image.height * max_width / image.width)
            image = image.resize((max_width, new_height), Image.Resampling.LANCZOS)
        if jpeg_quality > 0:
            suffix = ".jpg"
            if image.mode in ("RGBA", "LA") or (image.mode == "P" and "transparency" in image.info):
                rgba = image.convert("RGBA")
                canvas = Image.new("RGB", rgba.size, f"#{background}")
                canvas.paste(rgba, mask=rgba.getchannel("A"))
                image = canvas
            elif image.mode != "RGB":
                image = image.convert("RGB")
            target = work_dir / f"{index:04d}-{image_path.stem}{suffix}"
            image.save(target, "JPEG", quality=jpeg_quality, optimize=True)
        else:
            suffix = image_path.suffix.lower()
            target = work_dir / f"{index:04d}-{image_path.stem}{suffix}"
            if suffix in (".jpg", ".jpeg"):
                image.save(target, "JPEG", quality=95)
            else:
                image.save(target)
    return target
def add_fitted_picture(
    slide,
    image_path: Path,
    slide_width: int,
    slide_height: int,
    fit: str,
    name: str = "",
) -> None:
    if fit == "stretch":
        picture = slide.shapes.add_picture(str(image_path), 0, 0, width=slide_width, height=slide_height)
    else:
        image_width, image_height = image_size(image_path)
        if image_width <= 0 or image_height <= 0:
            raise ValueError(f"Invalid image dimensions: {image_path}")
        image_ratio = image_width / image_height
        slide_ratio = slide_width / slide_height
        if fit == "contain":
            if image_ratio >= slide_ratio:
                width = slide_width
                height = int(slide_width / image_ratio)
                left = 0
                top = int((slide_height - height) / 2)
            else:
                height = slide_height
                width = int(slide_height * image_ratio)
                top = 0
                left = int((slide_width - width) / 2)
            picture = slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)
        else:
            # Cover: place the picture at full slide size and center-crop the
            # overflow via crop attributes, so no pixels extend past the canvas.
            picture = slide.shapes.add_picture(
                str(image_path), 0, 0, width=slide_width, height=slide_height
            )
            if image_ratio > slide_ratio:
                crop = (1 - slide_ratio / image_ratio) / 2
                picture.crop_left = crop
                picture.crop_right = crop
            elif image_ratio < slide_ratio:
                crop = (1 - image_ratio / slide_ratio) / 2
                picture.crop_top = crop
                picture.crop_bottom = crop
    if name:
        picture.name = name
def build_deck(args: argparse.Namespace) -> dict[str, object]:
    input_dir = args.input_dir.resolve()
    output = args.output.resolve()
    spec = None
    if input_dir.is_file():
        spec, _ = load_page_spec(input_dir, require_images=True, strict=True)
        images = [(input_dir.parent / slide["image_file"]).resolve() for slide in spec["slides"]]
    elif input_dir.is_dir():
        images = collect_images(input_dir, tuple(args.extensions))
    else:
        raise FileNotFoundError(f"Input directory or Page Spec does not exist: {input_dir}")
    if not images:
        raise FileNotFoundError(f"No supported images found in: {input_dir}")
    width, height = args.width, args.height
    if spec:
        ratio = spec["canvas"]["width"] / spec["canvas"]["height"]
        if width is not None and height is not None and not math.isclose(width / height, ratio, rel_tol=1e-6):
            raise ValueError("Slide dimensions must preserve the Page Spec aspect ratio")
        if width is None:
            width = height * ratio if height is not None else 13.333333
        height = width / ratio
    else:
        width = width if width is not None else 13.333333
        height = height if height is not None else 7.5
    if not all(math.isfinite(value) and 1 <= value <= 56 for value in (width, height)):
        raise ValueError("Slide dimensions must be finite and between 1 and 56 inches")
    presentation = Presentation()
    presentation.slide_width = Inches(width)
    presentation.slide_height = Inches(height)
    presentation.core_properties.title = args.title or (spec["title"] if spec else output.stem)
    slide_width = presentation.slide_width
    slide_height = presentation.slide_height
    blank_layout = presentation.slide_layouts[6]
    process = args.max_width > 0 or args.jpeg_quality > 0
    with (tempfile.TemporaryDirectory() if process else nullcontext(None)) as tmp:
        for index, image_path in enumerate(images):
            embed_path = (
                prepare_image(
                    image_path, index, Path(tmp), args.max_width, args.jpeg_quality, args.background
                )
                if process
                else image_path
            )
            slide = presentation.slides.add_slide(blank_layout)
            add_background(slide, args.background)
            add_fitted_picture(
                slide, embed_path, slide_width, slide_height, args.fit, name=image_path.stem
            )
            if spec and "speaker_notes" in spec["slides"][index]:
                slide.notes_slide.notes_text_frame.text = canonical_notes(spec["slides"][index]["speaker_notes"])
        output.parent.mkdir(parents=True, exist_ok=True)
        fd, temp_output = tempfile.mkstemp(suffix=".pptx", dir=output.parent)
        os.close(fd)
        try:
            presentation.save(temp_output)
            verification = Presentation(temp_output)
            if len(verification.slides) != len(images):
                raise RuntimeError("Saved deck slide count does not match approved image count")
            if spec and (errors := check_notes(verification, spec["slides"])):
                raise ValueError("; ".join(errors))
            os.replace(temp_output, output)
        finally:
            Path(temp_output).unlink(missing_ok=True)
    return {
        "output": str(output),
        "slides": len(images),
        "slide_size_inches": [width, height],
        "fit": args.fit,
        "max_width": args.max_width,
        "jpeg_quality": args.jpeg_quality,
        "images": [path.name for path in images],
        "page_spec": str(input_dir) if spec else None,
        "content_version": spec["content_version"] if spec else None,
    }
def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_dir", type=Path, help="Approved Page Spec JSON, or a legacy directory of ordered images")
    parser.add_argument("output", type=Path, help="Output .pptx path")
    parser.add_argument("--fit", choices=("cover", "contain", "stretch"), default="cover")
    parser.add_argument("--width", type=float, help="Slide width in inches (default: 13.333333)")
    parser.add_argument("--height", type=float, help="Slide height (default: Page Spec ratio, or 7.5 for a directory)")
    parser.add_argument("--background", default="FFFFFF", help="Six-digit RGB background color")
    parser.add_argument("--title", default="", help="PowerPoint document title")
    parser.add_argument("--extensions", nargs="+", default=list(DEFAULT_EXTENSIONS))
    parser.add_argument(
        "--max-width",
        type=int,
        default=0,
        help="Downscale images wider than this many pixels before embedding (0 keeps original size)",
    )
    parser.add_argument(
        "--jpeg-quality",
        type=int,
        default=0,
        help="Re-encode images as JPEG at this quality (1-95) before embedding (0 keeps original format)",
    )
    args = parser.parse_args()
    if args.output.suffix.lower() != ".pptx":
        parser.error("output must use the .pptx extension")
    if any(value is not None and (not math.isfinite(value) or value <= 0) for value in (args.width, args.height)):
        parser.error("slide width and height must be finite and positive")
    if not re.fullmatch(r"[0-9A-Fa-f]{6}", args.background):
        parser.error("background must be a six-digit RGB value such as FFFFFF")
    args.background = args.background.upper()
    if args.max_width < 0:
        parser.error("max-width must be a non-negative pixel count")
    if args.jpeg_quality != 0 and not 1 <= args.jpeg_quality <= 95:
        parser.error("jpeg-quality must be between 1 and 95 (0 disables re-encoding)")
    return args
def main() -> None:
    args = parse_args()
    try:
        result = build_deck(args)
    except (ValueError, OSError) as error:
        raise SystemExit(f"build_image_ppt: {error}") from error
    print(json.dumps(result, ensure_ascii=True, indent=2))
if __name__ == "__main__":
    main()
