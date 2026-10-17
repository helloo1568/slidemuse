import copy
import json

import pytest
from audit_editability import audit
from build_editable_ppt import build_deck
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Pt
from scene import validate_scene
from table_layout import scaled_lengths


@pytest.fixture
def merged_scene():
    return {"version": "1.0", "canvas": {"width": 1600, "height": 900}, "slides": [
        {"id": "s01", "elements": [{
            "id": "results", "type": "table", "x": 100, "y": 100, "w": 1400, "h": 650,
            "font_size": 20, "font": "Microsoft YaHei", "valign": "middle",
            "rows": [["实验结果 / Results", "", ""], ["组别", "流量", "耗时"],
                     ["A", "100", "10 s"], ["", "200", "20 s"]],
            "column_widths": [2, 3, 3], "row_heights": [1.5, 1, 1, 1],
            "merges": [{"row": 0, "col": 0, "row_span": 1, "col_span": 3},
                       {"row": 2, "col": 0, "row_span": 2, "col_span": 1}],
        }]}]}


def write_scene(tmp_path, scene):
    path = tmp_path / "scene.json"
    path.write_text(json.dumps(scene, ensure_ascii=False), encoding="utf-8")
    return path


def test_native_merges_and_proportions_survive_save_reopen(tmp_path, merged_scene):
    path = write_scene(tmp_path, merged_scene)
    before = path.read_bytes()
    output = tmp_path / "merged.pptx"
    build_deck(path, output)
    table = Presentation(output).slides[0].shapes[0].table
    assert table.cell(0, 0).span_width == 3
    assert table.cell(2, 0).span_height == 2
    assert table.cell(0, 1).is_spanned
    assert table.cell(3, 0).is_spanned
    assert table.cell(0, 0).text == "实验结果 / Results"
    assert table.cell(2, 0).text == "A"
    assert table.cell(0, 0).vertical_anchor == MSO_ANCHOR.MIDDLE
    assert table.cell(0, 0).margin_top == Pt(5)
    assert table.rows[0].height / table.rows[1].height == pytest.approx(1.5)
    assert table.columns[1].width / table.columns[0].width == pytest.approx(1.5)
    assert not audit(output, path)["errors"]
    assert path.read_bytes() == before


def test_two_dimensional_merge_and_merged_text_edit(tmp_path, merged_scene):
    element = merged_scene["slides"][0]["elements"][0]
    element["rows"] = [["跨行跨列", "", "C"], ["", "", "D"]]
    element["row_heights"] = [1, 2]
    element["merges"] = [{"row": 0, "col": 0, "row_span": 2, "col_span": 2}]
    path = write_scene(tmp_path, merged_scene)
    output = tmp_path / "merged.pptx"
    build_deck(path, output)
    assert not audit(output, path)["errors"]
    prs = Presentation(output)
    prs.slides[0].shapes[0].table.cell(0, 0).text = "可以继续修改"
    prs.save(output)
    assert Presentation(output).slides[0].shapes[0].table.cell(0, 0).text == "可以继续修改"
    assert any("table content differs" in error for error in audit(output, path)["errors"])


@pytest.mark.parametrize("mutation, message", [
    ("covered_content", "Covered table cells"), ("overlap", "overlap"),
    ("out_of_bounds", "exceeds grid"), ("single", "at least two"),
    ("row_count", "row heights mismatch"), ("zero_weight", "schema"),
    ("negative_index", "schema"), ("fractional_span", "schema"),
    ("unknown_span_field", "schema"), ("bool_span", "schema"),
])
def test_bad_spans_rejected_before_replacing_output(tmp_path, merged_scene, mutation, message):
    element = merged_scene["slides"][0]["elements"][0]
    if mutation == "covered_content":
        element["rows"][0][1] = "不能丢失的内容"
    elif mutation == "overlap":
        element["merges"].append(copy.deepcopy(element["merges"][0]))
    elif mutation == "out_of_bounds":
        element["merges"][0]["col_span"] = 4
    elif mutation == "single":
        element["merges"][0]["col_span"] = 1
    elif mutation == "row_count":
        element["row_heights"] = [1]
    elif mutation == "zero_weight":
        element["row_heights"][0] = 0
    elif mutation == "negative_index":
        element["merges"][0]["row"] = -1
    elif mutation == "fractional_span":
        element["merges"][0]["row_span"] = 1.5
    elif mutation == "unknown_span_field":
        element["merges"][0]["colspan"] = 3
    else:
        element["merges"][0]["row_span"] = True
    path = write_scene(tmp_path, merged_scene)
    output = tmp_path / "existing.pptx"
    output.write_bytes(b"previous approved deck")
    with pytest.raises(ValueError, match=message):
        build_deck(path, output)
    assert output.read_bytes() == b"previous approved deck"


@pytest.mark.parametrize("mutation, message", [
    ("split", "merge differs"), ("columns", "column widths differ"),
    ("rows", "row heights differ"), ("covered_text", "content differs"),
])
def test_audit_detects_saved_table_mutations(tmp_path, merged_scene, mutation, message):
    path = write_scene(tmp_path, merged_scene)
    output = tmp_path / "merged.pptx"
    build_deck(path, output)
    prs = Presentation(output)
    table = prs.slides[0].shapes[0].table
    if mutation == "split":
        table.cell(0, 0).split()
    elif mutation == "columns":
        table.columns[0].width += 10000
        table.columns[1].width -= 10000
    elif mutation == "rows":
        table.rows[0].height += 10000
        table.rows[1].height -= 10000
    else:
        table.cell(0, 1).text = "hidden duplicate"
    prs.save(output)
    assert any(message in error for error in audit(output, path)["errors"])


def test_plain_tables_remain_valid(tmp_path, merged_scene):
    element = merged_scene["slides"][0]["elements"][0]
    element.pop("merges")
    element.pop("row_heights")
    assert not validate_scene(merged_scene, tmp_path)
    path = write_scene(tmp_path, merged_scene)
    output = tmp_path / "plain.pptx"
    build_deck(path, output)
    assert not audit(output, path)["errors"]


def test_extreme_weights_are_finite_and_unrepresentable_cells_are_rejected():
    assert scaled_lengths([1e308, 1e308], 100) == [50, 50]
    with pytest.raises(ValueError, match="zero-sized"):
        scaled_lengths([1e-300, 1], 100)
