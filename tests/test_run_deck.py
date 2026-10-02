import copy
import json
import subprocess
import sys
import time
from pathlib import Path

import pytest
import run_deck
from PIL import Image, ImageDraw
from pptx import Presentation


def save(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


@pytest.fixture
def pipeline_case(tmp_path, monkeypatch):
    root = Path(__file__).resolve().parents[1]
    spec = json.loads((root / "examples/page-spec.example.json").read_text(encoding="utf-8"))
    base = spec["slides"][0]
    spec["style"] = {"decision": "skipped", "authorization": "test fixture", "visual_spec": "test fixture"}
    spec["slides"] = []
    scene = {"version": "1.0", "canvas": spec["canvas"], "slides": []}
    (tmp_path / "slides").mkdir()
    for index, text in enumerate(("Alpha", "Beta"), 1):
        sid = f"s0{index}"
        slide = copy.deepcopy(base)
        slide.update(id=sid, page_number=index, title=text, core_message=text, source_refs=["test fixture"],
                     image_file=f"slides/{index}.png", image_status="approved", speaker_notes="test notes")
        slide["elements"] = [{"id": "text", "kind": "text", "role": "finding", "native_intent": "native",
                              "source_ref": "test fixture", "confirmation_status": "confirmed", "text": text}]
        spec["slides"].append(slide)
        Image.new("RGB", (160, 90), "white").save(tmp_path / slide["image_file"])
        scene["slides"].append({"id": sid, "speaker_notes": "test notes", "elements": [
            {"id": "text", "type": "text", "text": text, "x": 100, "y": 100, "w": 1000, "h": 150, "font_size": 24}]})
    spec_path, scene_path = tmp_path / "page-spec.json", tmp_path / "scene.json"
    save(spec_path, spec)
    save(scene_path, scene)
    config = {"version": "1.0", "mode": "editable", "page_spec": spec_path.name, "scene": scene_path.name,
              "backend": "powerpoint", "width": 320}
    job = tmp_path / "job.json"
    save(job, config)
    calls = []

    def renderer(deck, output, width, backend, slides):
        calls.append(list(slides))
        output.mkdir(parents=True, exist_ok=True)
        pptx = Presentation(deck)
        for n in slides:
            image = Image.new("RGB", (width, round(width * pptx.slide_height / pptx.slide_width)), "white")
            text = " ".join(shape.text for shape in pptx.slides[n - 1].shapes if shape.has_text_frame)
            ImageDraw.Draw(image).text((10, 10), text, fill="black")
            image.save(output / f"{n:03d}.png")
        return "powerpoint"

    monkeypatch.setattr(run_deck, "render_selected", renderer)
    return job, tmp_path / "workspace", spec, scene, calls


def review_fixture(case):
    _, work, spec, _, _ = case
    review = json.loads((work / "visual-review.json").read_text(encoding="utf-8"))
    for slide in review["slides"]:
        slide.update(status="pass", notes="Explicit test fixture review")
        slide["data_visual_inventory"].update(status="pass", notes="No data visuals in fixture", observed_element_ids=[])
    save(work / "visual-review.json", review)
    observations = json.loads((work / "content-observations.json").read_text(encoding="utf-8"))
    for slide, source in zip(observations["slides"], spec["slides"]):
        slide.update(status="complete", observed_text=source["elements"][0]["text"])
    save(work / "content-observations.json", observations)


def test_first_run_waits_for_review_and_resume_reuses_hash_bound_artifacts(pipeline_case):
    job, work, _, _, calls = pipeline_case
    first = run_deck.run(job, work)
    assert first["status"] == "awaiting_review"
    assert first["rendered_slides"] == ["s01", "s02"]
    assert first["actions"]
    assert read_status(work) == "awaiting_review"
    review_fixture(pipeline_case)
    resumed = run_deck.run(job, work)
    assert resumed["status"] == "complete"
    assert resumed["stages"]["build"] == "reused"
    assert resumed["reused_slides"] == ["s01", "s02"]
    assert calls == [[1, 2]]
    assert (work / "scorecard.json").is_file()


def read_status(work):
    return json.loads((work / "summary.json").read_text(encoding="utf-8"))["status"]


def test_changed_page_only_renders_that_page_and_invalidates_its_review(pipeline_case):
    job, work, spec, scene, calls = pipeline_case
    run_deck.run(job, work)
    review_fixture(pipeline_case)
    assert run_deck.run(job, work)["status"] == "complete"
    spec["slides"][0]["elements"][0]["text"] = "Changed Alpha"
    scene["slides"][0]["elements"][0]["text"] = "Changed Alpha"
    save(job.parent / "page-spec.json", spec)
    save(job.parent / "scene.json", scene)
    changed = run_deck.run(job, work)
    assert changed["status"] == "awaiting_review"
    assert changed["rendered_slides"] == ["s01"]
    assert changed["reused_slides"] == ["s02"]
    assert calls == [[1, 2], [1]]
    reviews = json.loads((work / "visual-review.json").read_text(encoding="utf-8"))["slides"]
    assert [s["status"] for s in reviews] == ["pending", "pass"]
    observations = json.loads((work / "content-observations.json").read_text(encoding="utf-8"))["slides"]
    assert [s["status"] for s in observations] == ["partial", "complete"]


def test_notes_only_change_rebuilds_export_and_keeps_visual_review(pipeline_case):
    job, work, spec, scene, calls = pipeline_case
    run_deck.run(job, work)
    review_fixture(pipeline_case)
    run_deck.run(job, work)
    spec["slides"][0]["speaker_notes"] = "Updated notes"
    scene["slides"][0]["speaker_notes"] = "Updated notes"
    save(job.parent / "page-spec.json", spec)
    save(job.parent / "scene.json", scene)
    updated = run_deck.run(job, work)
    assert updated["status"] == "complete"
    assert updated["stages"]["build"] == "built"
    assert updated["rendered_slides"] == []
    assert calls == [[1, 2]]
    assert Presentation(updated["deck"]).slides[0].notes_slide.notes_text_frame.text == "Updated notes"


def test_interrupted_renderer_keeps_build_and_recovers_next_run(pipeline_case, monkeypatch):
    job, work, _, _, calls = pipeline_case
    renderer = run_deck.render_selected

    def fail(*_args):
        raise RuntimeError("renderer unavailable")

    monkeypatch.setattr(run_deck, "render_selected", fail)
    failed = run_deck.run(job, work)
    assert failed["failed_stage"] == "render"
    assert run_deck.cached_state(work)["build"]
    monkeypatch.setattr(run_deck, "render_selected", renderer)
    resumed = run_deck.run(job, work)
    assert resumed["stages"]["build"] == "reused"
    assert resumed["status"] == "awaiting_review"
    assert calls == [[1, 2]]


@pytest.mark.parametrize("change", ["png", "deck", "width", "canvas", "refresh"])
def test_corrupt_output_and_changed_render_parameters_are_not_reused(pipeline_case, change):
    job, work, _, scene, calls = pipeline_case
    run_deck.run(job, work)
    state = run_deck.cached_state(work)
    if change == "png":
        Path(state["pages"]["s01"]["path"]).write_bytes(b"corrupt")
    elif change == "deck":
        Path(state["build"]["path"]).write_bytes(b"corrupt")
    elif change == "canvas":
        scene["canvas"]["width_inches"] = 12
        save(job.parent / "scene.json", scene)
    elif change == "width":
        config = json.loads(job.read_text(encoding="utf-8"))
        config["width"] = 640
        save(job, config)
    result = run_deck.run(job, work, refresh_render=change == "refresh")
    assert result["status"] == "awaiting_review"
    if change == "png":
        assert result["rendered_slides"] == ["s01"]
    elif change == "deck":
        assert result["stages"]["build"] == "built"
        assert len(calls) == 1
    else:
        assert result["rendered_slides"] == ["s01", "s02"]


def test_status_is_read_only_and_foreign_workspaces_are_rejected(pipeline_case):
    job, work, _, _, calls = pipeline_case
    assert run_deck.status(job, work)["status"] == "ready"
    assert not work.exists()
    assert not calls
    work.mkdir()
    (work / "precious.txt").write_text("preserve", encoding="utf-8")
    with pytest.raises(ValueError, match="new or previously owned"):
        run_deck.run(job, work)
    assert (work / "precious.txt").read_text(encoding="utf-8") == "preserve"


def test_stopped_build_continues_and_corrupt_state_cannot_bypass_checks(pipeline_case):
    job, work, _, _, calls = pipeline_case
    assert run_deck.run(job, work, stop_after="build")["status"] == "stopped"
    assert not calls
    assert run_deck.run(job, work)["stages"]["build"] == "reused"
    state = json.loads((work / "state.json").read_text(encoding="utf-8"))
    state["backend"] = "invented"
    save(work / "state.json", state)
    with pytest.raises(ValueError, match="state changed"):
        run_deck.run(job, work)


def test_old_pass_scorecard_is_removed_from_current_outputs_on_failure(pipeline_case, monkeypatch):
    job, work, _, _, _ = pipeline_case
    run_deck.run(job, work)
    review_fixture(pipeline_case)
    run_deck.run(job, work)
    assert (work / "scorecard.json").is_file()

    def broken(*_args):
        raise RuntimeError("renderer unavailable")

    monkeypatch.setattr(run_deck, "render_selected", broken)
    assert run_deck.run(job, work, refresh_render=True)["status"] == "failed"
    assert not (work / "scorecard.json").exists()
    assert list((work / "review-history").glob("scorecard-*.json"))


def test_image_mode_uses_approved_images_and_no_native_scene(pipeline_case):
    job, work, _, _, calls = pipeline_case
    save(job, {"version": "1.0", "mode": "image", "page_spec": "page-spec.json", "backend": "powerpoint", "width": 320})
    result = run_deck.run(job, work)
    assert result["status"] == "awaiting_review"
    assert len(Presentation(result["deck"]).slides) == 2
    assert calls == [[1, 2]]


def test_terminated_process_releases_workspace_lock_and_resumes_saved_build(pipeline_case):
    job, work, _, _, _ = pipeline_case
    scripts = Path(run_deck.__file__).parent
    marker = job.parent / "render-started"
    code = (f"import sys,time; from pathlib import Path; sys.path.insert(0,{str(scripts)!r}); import run_deck\n"
            f"def blocked(*args):\n Path({str(marker)!r}).write_text('started'); time.sleep(30)\n"
            f"run_deck.render_selected=blocked\nrun_deck.run(Path({str(job)!r}),Path({str(work)!r}))\n")
    process = subprocess.Popen([sys.executable, "-c", code], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        deadline = time.monotonic() + 10
        while not marker.exists() and process.poll() is None and time.monotonic() < deadline:
            time.sleep(0.05)
        assert marker.exists(), "Child must reach rendering after saving the build checkpoint"
        process.terminate()
        process.wait(timeout=5)
    finally:
        if process.poll() is None:
            process.terminate()
            process.wait(timeout=5)
    resumed = run_deck.run(job, work)
    assert resumed["stages"]["build"] == "reused"
    assert resumed["status"] == "awaiting_review"
