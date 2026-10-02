import copy
import json
from pathlib import Path

import pytest
from PIL import Image
from update_data_bindings import export_draft, inspect, plan_update, values_for
from validate_page_spec import load_page_spec


def dump(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


@pytest.fixture
def bound_case(tmp_path):
    spec = json.loads((Path(__file__).resolve().parents[1] / "examples/page-spec.example.json").read_text(encoding="utf-8"))
    slide = spec["slides"][0]
    slide["elements"] = [{"id": "finding", "kind": "text", "role": "evidence", "native_intent": "native",
                          "source_ref": "Synthetic fixture", "confirmation_status": "confirmed",
                          "text": "Total 150; share 80.0%"}]
    slide["speaker_notes"] = "120 / 150 * 100 = 80.0%; synthetic"
    slide["image_status"] = "approved"
    image = tmp_path / slide["image_file"]
    image.parent.mkdir()
    Image.new("RGB", (160, 90), "white").save(image)
    spec_path = tmp_path / "page-spec.json"
    dump(spec_path, spec)
    contract = {"version": "1.0", "documents": {"page_spec": "page-spec.json"},
                "inputs": {"current": {"value": 120, "unit": "vehicles", "source_ref": "Synthetic fixture"},
                           "other": {"value": 30, "unit": "vehicles", "source_ref": "Synthetic fixture"}},
                "formulas": {"share": "current / total * 100", "total": "current + other"},
                "bindings": [{"document": "page_spec", "pointer": "/slides/0/elements/0/text",
                              "template": "Total {total:.0f}; share {share:.1f}%"},
                             {"document": "page_spec", "pointer": "/slides/0/speaker_notes",
                              "template": "{current:.0f} / {total:.0f} * 100 = {share:.1f}%; synthetic"}]}
    path = tmp_path / "bindings.json"
    dump(path, contract)
    return path, spec_path, contract, spec


def test_preview_recalculates_transitive_prose_notes_and_leaves_source_unchanged(bound_case):
    path, spec_path, _, _ = bound_case
    original = (path.read_bytes(), spec_path.read_bytes())
    assert inspect(path)["status"] == "pass"
    plan, _ = plan_update(path, {"current": 90})
    assert plan["values"]["total"] == 120
    assert plan["values"]["share"] == 75
    assert len(plan["changes"]) == 2
    assert plan["documents"]["page_spec"]["slides"][0]["elements"][0]["text"] == "Total 120; share 75.0%"
    assert plan["affected_slides"] == ["s01"]
    assert (path.read_bytes(), spec_path.read_bytes()) == original


def test_draft_invalidates_visible_approval_and_copies_assets_without_old_qa(bound_case):
    path, spec_path, _, _ = bound_case
    (path.parent / "old-pass.json").write_text('"pass"', encoding="utf-8")
    plan, paths = plan_update(path, {"current": 90})
    output = path.parent / "draft"
    export_draft(path, output, plan, paths)
    with pytest.raises(ValueError, match="content_approved"):
        load_page_spec(output / "page-spec.json")
    draft = json.loads((output / "page-spec.json").read_text(encoding="utf-8"))
    assert draft["slides"][0]["image_status"] == "pending"
    assert (output / draft["slides"][0]["image_file"]).is_file()
    assert not (output / "old-pass.json").exists()
    assert json.loads(spec_path.read_text(encoding="utf-8"))["content_approved"] is True
    with pytest.raises(ValueError, match="new directory"):
        export_draft(path, output, plan, paths)


def test_stale_notes_block_update_even_if_numbers_are_correct(bound_case):
    path, spec_path, _, spec = bound_case
    spec["slides"][0]["speaker_notes"] = "Old percentage 79%"
    dump(spec_path, spec)
    assert inspect(path)["status"] == "fail"
    with pytest.raises(ValueError, match="stale content"):
        plan_update(path, {"current": 90})


@pytest.mark.parametrize("formula", ["unknown + 1", "total", "current / 0", "__import__('os')", "current ** 100", "current.real", "1e309"])
def test_unsafe_unknown_cyclic_and_nonfinite_formulas_fail(bound_case, formula):
    _, _, contract, _ = bound_case
    contract["formulas"]["total"] = formula
    with pytest.raises((ValueError, TypeError)):
        values_for(contract)


@pytest.mark.parametrize("change", ["unknown", "derived", "boolean", "nan", "asset", "metadata", "duplicate", "type-change", "format-access", "escape"])
def test_invalid_updates_and_bindings_fail_without_writes(bound_case, change):
    path, spec_path, contract, _ = bound_case
    updates = {"current": 90}
    if change in ("unknown", "derived"):
        updates = {"missing" if change == "unknown" else "share": 90}
    elif change in ("boolean", "nan"):
        updates = {"current": True if change == "boolean" else float("nan")}
    elif change == "asset":
        contract["bindings"][0]["pointer"] = "/slides/0/image_file"
    elif change == "metadata":
        contract["bindings"][0]["pointer"] = "/content_approved"
    elif change == "duplicate":
        contract["bindings"].append(copy.deepcopy(contract["bindings"][0]))
    elif change == "type-change":
        contract["bindings"][0].pop("template")
        contract["bindings"][0]["value"] = "current"
    elif change == "format-access":
        contract["bindings"][0]["template"] = "{current.__class__}"
    else:
        contract["documents"]["page_spec"] = "../escape.json"
    dump(path, contract)
    original = (path.read_bytes(), spec_path.read_bytes())
    with pytest.raises((ValueError, TypeError)):
        plan_update(path, updates)
    assert (path.read_bytes(), spec_path.read_bytes()) == original


def test_source_changed_after_preview_cannot_export_stale_plan(bound_case):
    path, spec_path, _, spec = bound_case
    plan, paths = plan_update(path, {"current": 90})
    spec["title"] = "Changed after preview"
    dump(spec_path, spec)
    with pytest.raises(ValueError, match="changed after preview"):
        export_draft(path, path.parent / "draft", plan, paths)
    assert not (path.parent / "draft").exists()


def test_example_updates_real_native_workbook_prose_and_notes(tmp_path):
    from audit_editability import audit
    from build_editable_ppt import build_deck
    from pptx import Presentation

    example = Path(__file__).resolve().parents[1] / "examples/data-update"
    for name in ("data-bindings.json", "page-spec.json", "scene.json"):
        (tmp_path / name).write_bytes((example / name).read_bytes())
    contract_path = tmp_path / "data-bindings.json"
    plan, _ = plan_update(contract_path, {"north_after": 110})
    assert len(plan["changes"]) == 14
    assert plan["values"]["after_total"] == 210
    assert plan["values"]["change_percent"] == 5
    assert round(plan["values"]["north_share"], 1) == 52.4
    for kind, document in plan["documents"].items():
        dump(tmp_path / ("page-spec.json" if kind == "page_spec" else "scene.json"), document)
    dump(contract_path, plan["contract"])
    assert inspect(contract_path)["status"] == "pass"
    output = tmp_path / "edited.pptx"
    build_deck(tmp_path / "scene.json", output, tmp_path / "page-spec.json")
    assert audit(output, tmp_path / "scene.json")["errors"] == []
    pptx = Presentation(output)
    for slide in pptx.slides:
        chart = next(shape.chart for shape in slide.shapes if shape.has_chart)
        assert list(chart.series[1].values) == [110, 100]
        assert "after 210" in slide.notes_slide.notes_text_frame.text
        assert "52.4%" in slide.notes_slide.notes_text_frame.text
        assert any("total 210" in shape.text for shape in slide.shapes if shape.has_text_frame)
    scene = plan["documents"]["scene"]
    scene["slides"][0]["elements"][1]["text"] = "Old total 190"
    dump(tmp_path / "scene.json", scene)
    assert inspect(contract_path)["status"] == "fail"
