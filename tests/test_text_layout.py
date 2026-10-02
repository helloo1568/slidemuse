import json

import pytest
import run_deck
from audit_editability import audit
from build_editable_ppt import build_deck, text_frame
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt
from text_layout import estimate_frame, inspect_deck, inspect_shape

pytest_plugins = ['test_run_deck']


def box(text='Normal text', width=3, height=1):
    prs = Presentation()
    shape = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0, 0, Inches(width), Inches(height))
    shape.name = 'body | text'
    text_frame(shape.text_frame, text, {'font_size': 20})
    return prs, shape


@pytest.mark.parametrize('text', ['Normal text', '中文标题与正文', 'Latin 与中文混排 2026', 'Line one\nLine two'])
def test_spacious_text_has_no_strong_risk_without_claiming_verified_fit(text):
    _, shape = box(text)
    result = estimate_frame(shape.text_frame, shape.width, shape.height)
    assert result['status'] == 'no_obvious_risk'
    assert 'unverified' in result['reason']


@pytest.mark.parametrize('text', ['Dense English paragraph ' * 50, '这是需要折行的中文正文' * 30, 'One\nTwo\nThree\nFour\nFive', 'One\vTwo\vThree\vFour\vFive'])
def test_dense_and_explicit_break_overflow_is_flagged(text):
    _, shape = box(text, width=2, height=.5)
    result = inspect_shape(shape)[0]
    assert result['status'] == 'risk'
    assert result['element_id'] == 'body'


def test_margins_run_override_and_paragraph_spacing_affect_estimate():
    _, shape = box('Normal', width=2, height=.5)
    assert inspect_shape(shape)[0]['status'] == 'no_obvious_risk'
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(80)
    assert inspect_shape(shape)[0]['status'] == 'risk'
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    shape.text_frame.paragraphs[0].space_after = Pt(50)
    assert inspect_shape(shape)[0]['status'] == 'risk'
    shape.text_frame.margin_left = Inches(2)
    assert 'margins' in inspect_shape(shape)[0]['reason']


@pytest.mark.parametrize('mutation', ['inherited_size', 'inherited_spacing', 'inherited_wrap', 'autofit', 'rotation', 'tab', 'complex', 'vertical', 'columns', 'internal_rotation'])
def test_unsupported_layout_is_unverified_instead_of_fabricated_fit(mutation):
    _, shape = box()
    frame = shape.text_frame
    if mutation == 'inherited_size':
        frame.paragraphs[0].font.size = None
    elif mutation == 'inherited_spacing':
        frame.paragraphs[0].line_spacing = None
    elif mutation == 'inherited_wrap':
        frame.word_wrap = None
    elif mutation == 'autofit':
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    elif mutation == 'rotation':
        shape.rotation = 90
    elif mutation in ('vertical', 'columns', 'internal_rotation'):
        attribute, value = {'vertical': ('vert', 'vert'), 'columns': ('numCol', '2'), 'internal_rotation': ('rot', '5400000')}[mutation]
        frame._txBody.bodyPr.set(attribute, value)
    else:
        text_frame(frame, 'Text\tTab' if mutation == 'tab' else 'مرحبا', {'font_size': 20})
    assert inspect_shape(shape)[0]['status'] == 'unverified'


def test_table_cells_and_transformed_groups_are_identified():
    prs, normal = box()
    table = prs.slides[0].shapes.add_table(1, 2, 0, 0, Inches(4), Inches(.4))
    table.name = 'matrix | table'
    for cell in table.table.rows[0].cells:
        text_frame(cell.text_frame, 'Long text ' * 40, {'font_size': 20})
    findings = inspect_shape(table)
    assert [f['cell'] for f in findings] == [[1, 1], [1, 2]]
    assert all(f['status'] == 'risk' for f in findings)
    group = prs.slides[0].shapes.add_group_shape([normal])
    assert next(f for f in inspect_deck(prs) if f['element_id'] == 'body')['status'] == 'no_obvious_risk'
    group.width *= 2
    assert next(f for f in inspect_deck(prs) if f['element_id'] == 'body')['status'] == 'unverified'


def test_export_audit_and_pipeline_preserve_text_and_flag_current_page(pipeline_case):
    job, work, _, scene, _ = pipeline_case
    source = job.parent / 'scene.json'
    scene['slides'][0]['elements'][0].update(h=10, font_size=80)
    run_deck.write(source, scene)
    target = job.parent / 'direct.pptx'
    before = source.read_bytes()
    built = build_deck(source, target)
    assert built['text_layout'][0]['status'] == 'risk'
    assert any('possible text overflow' in x for x in built['warnings'])
    actual = audit(target, source)
    assert not actual['errors']
    assert actual['text_layout'][0]['status'] == 'risk'
    assert source.read_bytes() == before
    assert Presentation(target).slides[0].shapes[0].text == 'Alpha'
    result = run_deck.run(job, work)
    tasks = [a for a in result['actions'] if a['kind'] == 'text_overflow']
    assert len(tasks) == 1 and tasks[0]['slide_id'] == 's01' and tasks[0]['element_id'] == 'text'
    assert tasks[0]['rendered'] and tasks[0]['reference'] and tasks[0]['page_number'] == 1
    assert '可能文字溢出' in (work / 'summary.md').read_text(encoding='utf-8')
    review = json.loads((work / 'visual-review.json').read_text(encoding='utf-8'))
    assert review['slides'][0]['status'] == 'pending'
    stopped = run_deck.run(job, work, stop_after='build')
    assert any(a['kind'] == 'text_overflow' for a in stopped['actions'])
