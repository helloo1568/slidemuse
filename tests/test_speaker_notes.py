import copy
import json
from pathlib import Path

import audit_editability
import audit_speaker_notes
import build_editable_ppt
import build_image_ppt
import evaluate_delivery
import export_speaker_notes
import plan_deck_update
import pytest
import render_deck
from PIL import Image
from pptx import Presentation
from speaker_notes import canonical_notes, read_notes
from test_evaluate_delivery import make_case, mark_visual_pass

ROOT = Path(__file__).resolve().parents[1]
NOTES = "Q 与 K 决定权重，V 提供被汇总的信息。\r\n\r\nAttention(Q,K,V) = softmax(QKᵀ/√dₖ)V\r\n来源：§3.2.1；不是所有任务都更优。\n  保留缩进与空行。  "


def case(tmp_path):
    spec = json.loads((ROOT / "examples/page-spec.example.json").read_text(encoding="utf-8"))
    first = spec["slides"][0]
    first.update(image_status="approved", speaker_notes=NOTES, notes="WORK-RECORD: do not export")
    second = copy.deepcopy(first)
    second.update(id="s02", page_number=2, image_file="slides/02.png", speaker_notes="政策目标，不是实测效果。\n2027 / 2030 / 2035")
    spec["slides"].append(second)
    scene = {"version": "1.0", "canvas": spec["canvas"], "slides": []}
    for page in spec["slides"]:
        image = tmp_path / page["image_file"]
        image.parent.mkdir(exist_ok=True)
        Image.new("RGB", (160, 90), "white").save(image)
        scene["slides"].append({"id": page["id"], "speaker_notes": page["speaker_notes"],
                                "notes": "WORK-RECORD: do not export",
                                "elements": [{"id": "title", "type": "text", "text": page["title"],
                                              "x": 10, "y": 10, "w": 100, "h": 50}]})
    spec_path, scene_path = tmp_path / "page-spec.json", tmp_path / "scene.json"
    save(spec_path, spec)
    save(scene_path, scene)
    return spec_path, scene_path, spec, scene


def save(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")


def image_build(spec_path, output):
    from argparse import Namespace
    return build_image_ppt.build_deck(Namespace(input_dir=spec_path, output=output,
        width=None, height=None, title="", max_width=0, jpeg_quality=0,
        background="FFFFFF", fit="cover", extensions=[".png"]))


def test_both_exports_keep_exact_multilingual_notes_without_work_records(tmp_path):
    spec_path, scene_path, spec, _ = case(tmp_path)
    image, editable = tmp_path / "image.pptx", tmp_path / "editable.pptx"
    image_build(spec_path, image)
    build_editable_ppt.build_deck(scene_path, editable, spec_path)
    for deck in (image, editable):
        actual = Presentation(deck)
        assert [read_notes(s) for s in actual.slides] == [canonical_notes(s["speaker_notes"]) for s in spec["slides"]]
        assert all("WORK-RECORD" not in read_notes(s) for s in actual.slides)
        assert audit_speaker_notes.audit(spec_path, deck)["status"] == "pass"
    assert not audit_editability.audit(editable, scene_path)["errors"]


@pytest.mark.parametrize("mutation", ["missing", "wrong", "reordered", "count", "empty"])
def test_notes_audit_detects_real_pptx_changes(tmp_path, mutation):
    spec_path, _, _, _ = case(tmp_path)
    deck = tmp_path / "image.pptx"
    image_build(spec_path, deck)
    presentation = Presentation(deck)
    if mutation in ("missing", "wrong", "empty"):
        presentation.slides[0].notes_slide.notes_text_frame.text = "wrong" if mutation == "wrong" else ""
    elif mutation == "reordered":
        ids = presentation.slides._sldIdLst
        ids.insert(0, ids[-1])
    else:
        presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(deck)
    report = audit_speaker_notes.audit(spec_path, deck)
    assert report["status"] == "fail" and report["errors"]


@pytest.mark.parametrize("mutation", ["missing", "changed", "reordered"])
def test_upstream_notes_mismatch_rejects_build_without_replacing_output(tmp_path, mutation):
    spec_path, scene_path, _, scene = case(tmp_path)
    if mutation == "missing":
        scene["slides"][0].pop("speaker_notes")
    elif mutation == "changed":
        scene["slides"][0]["speaker_notes"] = "Silently changed claim"
    else:
        scene["slides"].reverse()
    save(scene_path, scene)
    output = tmp_path / "existing.pptx"
    output.write_bytes(b"keep this accepted deliverable")
    with pytest.raises(ValueError, match="speaker_notes|IDs/order"):
        build_editable_ppt.build_deck(scene_path, output, spec_path)
    assert output.read_bytes() == b"keep this accepted deliverable"


def test_scene_only_audit_detects_tampered_legacy_notes(tmp_path):
    _, scene_path, _, scene = case(tmp_path)
    for page in scene["slides"]:
        page.pop("speaker_notes")
        page["notes"] = "Legacy presenter notes"
    save(scene_path, scene)
    deck = tmp_path / "editable.pptx"
    build_editable_ppt.build_deck(scene_path, deck)
    presentation = Presentation(deck)
    assert read_notes(presentation.slides[0]) == "Legacy presenter notes"
    presentation.slides[0].notes_slide.notes_text_frame.text = "tampered"
    presentation.save(deck)
    assert any("speaker notes" in e for e in audit_editability.audit(deck, scene_path)["errors"])


def test_notes_upstream_check_preserves_unresolved_reconstruction_workflow(tmp_path):
    spec_path, scene_path, spec, _ = case(tmp_path)
    spec["slides"][0]["elements"][0].update(confirmation_status="unresolved", confidence=0.7)
    save(spec_path, spec)
    output = tmp_path / "partial-restoration.pptx"
    build_editable_ppt.build_deck(scene_path, output, spec_path)
    assert audit_speaker_notes.audit(spec_path, output, scene_path)["status"] == "pass"


def test_undeclared_and_explicit_empty_notes_are_different_contracts(tmp_path):
    spec_path, _, spec, _ = case(tmp_path)
    for page in spec["slides"]:
        page.pop("speaker_notes")
    save(spec_path, spec)
    deck = tmp_path / "legacy.pptx"
    image_build(spec_path, deck)
    presentation = Presentation(deck)
    assert not presentation.slides[0].has_notes_slide
    presentation.slides[0].notes_slide.notes_text_frame.text = "Existing legacy remarks"
    presentation.save(deck)
    assert audit_speaker_notes.audit(spec_path, deck)["status"] == "not_declared"
    spec["slides"][0]["speaker_notes"] = ""
    save(spec_path, spec)
    assert audit_speaker_notes.audit(spec_path, deck)["status"] == "fail"
    image_build(spec_path, deck)
    assert audit_speaker_notes.audit(spec_path, deck)["status"] == "pass"


def test_notes_only_changes_rebuild_exports_without_regenerating_images(tmp_path):
    spec_path, _, spec, _ = case(tmp_path)
    baseline = plan_deck_update.snapshot(spec_path)
    spec["slides"][0]["speaker_notes"] += "\n补充解释。"
    plan = plan_deck_update.plan(baseline, spec, tmp_path)
    assert [s["action"] for s in plan["slides"]] == ["reuse", "reuse"]
    assert not any(s["stale_approval"] for s in plan["slides"])
    assert plan["slides"][0]["export_reasons"] == ["speaker_notes_changed"]
    assert plan["rebuild_image_deck"] and plan["rebuild_editable_deck"]
    spec["slides"][0]["elements"][0]["text"] += "new visible text"
    assert plan_deck_update.plan(baseline, spec, tmp_path)["slides"][0]["action"] == "regenerate"


def test_old_snapshot_without_notes_hash_remains_usable(tmp_path):
    spec_path, _, spec, _ = case(tmp_path)
    for page in spec["slides"]:
        page.pop("speaker_notes")
    save(spec_path, spec)
    baseline = plan_deck_update.snapshot(spec_path)
    for page in baseline["slides"]:
        page.pop("speaker_notes_hash")
    assert not plan_deck_update.plan(baseline, spec, tmp_path)["rebuild_image_deck"]
    spec["slides"][0]["speaker_notes"] = ""
    assert plan_deck_update.plan(baseline, spec, tmp_path)["rebuild_image_deck"]


def test_handout_preserves_text_and_sources_without_inferring_work_records(tmp_path):
    spec_path, _, spec, _ = case(tmp_path)
    spec["slides"][0]["speaker_notes"] += "\n```\n<script>literal reference</script>"
    save(spec_path, spec)
    output = tmp_path / "speaker-notes.md"
    report = export_speaker_notes.export(spec_path, output)
    text = output.read_text(encoding="utf-8")
    assert canonical_notes(spec["slides"][0]["speaker_notes"]) in text
    assert "````text" in text and "WORK-RECORD" not in text
    assert all(e["source_ref"] in text for s in spec["slides"] for e in s["elements"])
    assert report["page_spec_sha256"] == audit_speaker_notes.sha256(spec_path)
    for page in spec["slides"]:
        page.pop("speaker_notes")
    save(spec_path, spec)
    before = output.read_bytes()
    with pytest.raises(ValueError, match="No speaker_notes"):
        export_speaker_notes.export(spec_path, output)
    assert output.read_bytes() == before


def test_scorecard_rejects_missing_notes_despite_visual_and_content_pass(tmp_path):
    spec_path, render, review, observations, _ = make_case(tmp_path)
    mark_visual_pass(review)
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    spec["slides"][0]["speaker_notes"] = NOTES
    save(spec_path, spec)
    report = json.loads(render.read_text(encoding="utf-8"))
    report["page_spec_sha256"] = render_deck.sha256(spec_path)
    save(render, report)
    checked = json.loads(review.read_text(encoding="utf-8"))
    checked["page_spec_sha256"] = report["page_spec_sha256"]
    save(review, checked)
    result = evaluate_delivery.evaluate(spec_path, render, review, [observations])
    assert result["content"]["issues"] == 0 and result["visual"]["passed"] == 1
    assert result["status"] == "fail" and result["speaker_notes"]["errors"]
