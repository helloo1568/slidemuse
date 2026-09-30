import json
from pathlib import Path

import audit_page_content
import evaluate_delivery
import pytest
import render_deck
from PIL import Image
from pptx import Presentation


def make_case(tmp_path, data_kind=None):
    spec = json.loads((Path(__file__).resolve().parents[1] / "examples/page-spec.example.json").read_text(encoding="utf-8"))
    slide = spec["slides"][0]
    slide["image_status"] = "approved"
    if data_kind:
        slide["elements"].append({"id": "results", "kind": data_kind, "role": "results",
                                  "native_intent": "native", "source_ref": "report/table-1",
                                  "confirmation_status": "source-verified", "data": {"values": [12, 15]}})
    image_path = tmp_path / slide["image_file"]
    image_path.parent.mkdir()
    Image.new("RGB", (160, 90), "white").save(image_path)
    spec_path = tmp_path / "page-spec.json"
    spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
    deck = tmp_path / "deck.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(deck)
    rendered = tmp_path / "001.png"
    Image.new("RGB", (160, 90), "white").save(rendered)
    report = {"deck": str(deck), "deck_sha256": render_deck.sha256(deck),
              "page_spec_sha256": render_deck.sha256(spec_path), "backend": "test",
              "slides": [{"slide": 1, "rendered": str(rendered), "rendered_sha256": render_deck.sha256(rendered),
                          "reference": str(image_path), "reference_sha256": render_deck.sha256(image_path)}]}
    render_path = tmp_path / "render-report.json"
    render_path.write_text(json.dumps(report), encoding="utf-8")
    review = evaluate_delivery.review_template(spec, report)
    review_path = tmp_path / "visual-review.json"
    review_path.write_text(json.dumps(review), encoding="utf-8")
    observations = audit_page_content.template(spec_path, spec["slides"])
    observations["slides"][0].update(status="complete", observed_text="让数字服务真正走进田间\n以可信、易用的数字服务连接农户与产业资源")
    observation_path = tmp_path / "observations.json"
    observation_path.write_text(json.dumps(observations, ensure_ascii=False), encoding="utf-8")
    return spec_path, render_path, review_path, observation_path, rendered


def test_scorecard_needs_visual_review_and_current_artifacts(tmp_path):
    spec, render, review, observation, rendered = make_case(tmp_path)
    result = evaluate_delivery.evaluate(spec, render, review, [observation])
    assert result["status"] == "incomplete"
    assert result["content"]["issues"] == 0
    reviewed = json.loads(review.read_text(encoding="utf-8"))
    reviewed["slides"][0]["status"] = "pass"
    reviewed["slides"][0]["data_visual_inventory"].update(status="pass", notes="No data visuals observed", observed_element_ids=[])
    review.write_text(json.dumps(reviewed), encoding="utf-8")
    assert evaluate_delivery.evaluate(spec, render, review, [observation])["status"] == "pass"
    Image.new("RGB", (160, 90), "black").save(rendered)
    with pytest.raises(ValueError, match="Rendered slide 1 changed"):
        evaluate_delivery.evaluate(spec, render, review, [observation])


def test_scorecard_rejects_render_without_current_page_spec(tmp_path):
    spec, render, review, observation, _ = make_case(tmp_path)
    report = json.loads(render.read_text(encoding="utf-8"))
    report.pop("page_spec_sha256")
    render.write_text(json.dumps(report), encoding="utf-8")
    with pytest.raises(ValueError, match="lacks the current Page Spec hash"):
        evaluate_delivery.evaluate(spec, render, review, [observation])


def mark_visual_pass(review_path):
    review = json.loads(review_path.read_text(encoding="utf-8"))
    review["slides"][0]["status"] = "pass"
    review["slides"][0]["data_visual_inventory"].update(
        status="pass", notes="Observed only declared data visuals",
        observed_element_ids=[item["element_id"] for item in review["slides"][0]["data_reviews"]])
    review_path.write_text(json.dumps(review), encoding="utf-8")
    return review


def test_changed_reference_rejects_old_render_even_with_new_observations(tmp_path):
    spec, render, review, observation, _ = make_case(tmp_path)
    mark_visual_pass(review)
    approved = json.loads(spec.read_text(encoding="utf-8"))
    reference = tmp_path / approved["slides"][0]["image_file"]
    Image.new("RGB", (160, 90), "red").save(reference)
    observations = json.loads(observation.read_text(encoding="utf-8"))
    observations["slides"][0]["image_sha256"] = render_deck.sha256(reference)
    observation.write_text(json.dumps(observations), encoding="utf-8")
    with pytest.raises(ValueError, match="Reference slide 1 changed"):
        evaluate_delivery.evaluate(spec, render, review, [observation])
    # Even if a new render has identical pixels, its old visual approval is stale.
    report = json.loads(render.read_text(encoding="utf-8"))
    report["slides"][0]["reference_sha256"] = render_deck.sha256(reference)
    render.write_text(json.dumps(report), encoding="utf-8")
    with pytest.raises(ValueError, match="Visual review is stale"):
        evaluate_delivery.evaluate(spec, render, review, [observation])


@pytest.mark.parametrize("field,value", [("reference_sha256", None), ("reference", "wrong.png")])
def test_missing_or_mismatched_reference_evidence_is_rejected(tmp_path, field, value):
    spec, render, review, observation, _ = make_case(tmp_path)
    report = json.loads(render.read_text(encoding="utf-8"))
    report["slides"][0][field] = value
    render.write_text(json.dumps(report), encoding="utf-8")
    with pytest.raises(ValueError, match="Reference slide 1 changed or lacks a hash"):
        evaluate_delivery.evaluate(spec, render, review, [observation])


@pytest.mark.parametrize("change", ["spec", "backend", "legacy-review"])
def test_old_visual_review_cannot_approve_changed_review_context(tmp_path, change):
    spec, render, review, observation, _ = make_case(tmp_path)
    reviewed = mark_visual_pass(review)
    report = json.loads(render.read_text(encoding="utf-8"))
    if change == "spec":
        approved = json.loads(spec.read_text(encoding="utf-8"))
        approved["content_version"] = "outline-v2"
        spec.write_text(json.dumps(approved), encoding="utf-8")
        report["page_spec_sha256"] = render_deck.sha256(spec)
    elif change == "backend":
        report["backend"] = "different-renderer"
    else:
        reviewed["version"] = "1.0"
        review.write_text(json.dumps(reviewed), encoding="utf-8")
    render.write_text(json.dumps(report), encoding="utf-8")
    with pytest.raises(ValueError, match="create a new --init-review template"):
        evaluate_delivery.evaluate(spec, render, review, [observation])


@pytest.mark.parametrize("kind", ["chart", "table"])
def test_data_needs_independent_review_and_records_findings(tmp_path, kind):
    spec, render, review, observation, _ = make_case(tmp_path, kind)
    reviewed = mark_visual_pass(review)
    result = evaluate_delivery.evaluate(spec, render, review, [observation])
    assert result["status"] == "incomplete"
    assert result["data"]["pending"][0]["element_id"] == "results"
    reviewed["slides"][0]["data_reviews"][0].update(status="fail", notes="Rendered value is 13, expected 12")
    review.write_text(json.dumps(reviewed), encoding="utf-8")
    assert evaluate_delivery.evaluate(spec, render, review, [observation])["status"] == "fail"
    reviewed["slides"][0]["data_reviews"][0].update(status="pass", notes="Checked values 12 and 15, units and source report/table-1")
    entry = reviewed["slides"][0]["data_reviews"][0]
    entry["checks"] = dict.fromkeys(entry["checks"], "pass")
    review.write_text(json.dumps(reviewed), encoding="utf-8")
    result = evaluate_delivery.evaluate(spec, render, review, [observation])
    assert result["status"] == "pass"
    assert result["data"]["passed"] == 1
    assert result["data"]["reviews"][0]["notes"].startswith("Checked values")
    reviewed["slides"][0].pop("data_reviews")
    review.write_text(json.dumps(reviewed), encoding="utf-8")
    assert evaluate_delivery.evaluate(spec, render, review, [observation])["status"] == "incomplete"


@pytest.mark.parametrize("change", ["duplicate", "unknown", "empty-notes", "invalid-status"])
def test_invalid_data_review_cannot_pass(tmp_path, change):
    spec, render, review, observation, _ = make_case(tmp_path, "chart")
    reviewed = mark_visual_pass(review)
    entries = reviewed["slides"][0]["data_reviews"]
    if change == "duplicate":
        entries.append(dict(entries[0]))
    elif change == "unknown":
        entries[0]["element_id"] = "missing"
    elif change == "empty-notes":
        entries[0].update(status="pass", notes="  ")
    else:
        entries[0]["status"] = "complete"
    review.write_text(json.dumps(reviewed), encoding="utf-8")
    with pytest.raises(ValueError, match="data review|Data review"):
        evaluate_delivery.evaluate(spec, render, review, [observation])


@pytest.mark.parametrize("check_status,expected", [("pending", "incomplete"), ("fail", "fail"), ("pass", "pass")])
def test_correct_labels_do_not_bypass_geometry_review(tmp_path, check_status, expected):
    spec, render, review, observation, _ = make_case(tmp_path, "chart")
    reviewed = mark_visual_pass(review)
    item = reviewed["slides"][0]["data_reviews"][0]
    item.update(status="pass", notes="Source and text labels correct; geometric scale checked separately")
    item["checks"].update(source_values="pass", labels_units="pass", scale_geometry=check_status)
    review.write_text(json.dumps(reviewed), encoding="utf-8")
    assert evaluate_delivery.evaluate(spec, render, review, [observation])["status"] == expected


def test_unlisted_chart_on_a_text_page_fails_inventory_review(tmp_path):
    spec, render, review, observation, _ = make_case(tmp_path)
    reviewed = mark_visual_pass(review)
    reviewed["slides"][0]["data_visual_inventory"].update(
        observed_element_ids=["invented-trend"], notes="An extra unapproved trend chart was observed")
    review.write_text(json.dumps(reviewed), encoding="utf-8")
    result = evaluate_delivery.evaluate(spec, render, review, [observation])
    assert result["status"] == "fail"
    assert result["data_visual_inventory"]["failed"]


def test_missing_inventory_cannot_pass_even_when_text_and_layout_pass(tmp_path):
    spec, render, review, observation, _ = make_case(tmp_path)
    reviewed = mark_visual_pass(review)
    reviewed["slides"][0].pop("data_visual_inventory")
    review.write_text(json.dumps(reviewed), encoding="utf-8")
    assert evaluate_delivery.evaluate(spec, render, review, [observation])["status"] == "incomplete"
