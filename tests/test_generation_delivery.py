"""Behavioral regressions for content loss and unsafe generation delivery.

All images/transcriptions in these tests are synthetic protocol fixtures, not
evidence of model image quality or authenticated real-world sources.
"""
import copy
import json
from pathlib import Path

import pytest
from content_density import validate_content
from generation_delivery import CHECKS, CONTRACT, REVIEW, review
from PIL import Image
from pptx import Presentation
from test_generate_backend import invoke, run_pipeline, write_slide_image


def record_pages(tmp_path):
    deck = run_pipeline(tmp_path)
    for number in (1, 2):
        image = tmp_path / f"source-{number}.png"
        write_slide_image(image)
        assert invoke("--offline", "dispatch", deck, "--slide", number, "--agent-id", "main").returncode == 0
        result = invoke("--offline", "result", deck, "--slide", number, "--agent-id", "main",
                        "--backend-used", "built-in image tool", "--selected-source", image,
                        "--qa-note", "Synthetic protocol fixture")
        assert result.returncode == 0, result.stderr
    return deck


def fill_review(deck):
    review(deck, deck / REVIEW, initialize=True)
    data = json.loads((deck / REVIEW).read_text(encoding="utf-8"))
    contract = json.loads((deck / CONTRACT).read_text(encoding="utf-8"))
    for page, item in zip(contract["pages"], data["slides"]):
        item.update(observed_text="\n".join(page["expected_text"]), notes="Synthetic protocol observation",
                    checks=dict.fromkeys(CHECKS, "pass"))
    (deck / REVIEW).write_text(json.dumps(data), encoding="utf-8")


@pytest.fixture
def content():
    return {"title": "Synthetic result", "role": "data-evidence", "key_points": ["Source fact", "Comparison", "Limitation"],
            "core_message": "A supported conclusion", "source_refs": ["Synthetic fixture table 1"],
            "explanation": "Comparison uses the same time period."}


@pytest.mark.parametrize("field", ["role", "core_message", "key_points", "source_refs", "explanation"])
def test_required_content_cannot_be_silently_omitted(content, field):
    content.pop(field)
    with pytest.raises(ValueError):
        validate_content(content, required=True)


@pytest.mark.parametrize("mode,points", [("brief", ["One focused point"]),
                                         ("evidence-led", ["Graph carries the comparison"]),
                                         ("detailed", [str(i) for i in range(6)])])
def test_density_exceptions_need_explicit_reasons(content, mode, points):
    content.update(key_points=points, density={"mode": mode})
    with pytest.raises(ValueError, match="reason"):
        validate_content(content, required=True)
    content["density"]["reason"] = "Synthetic example: a full graph makes extra bullets redundant"
    validate_content(content, required=True)


def test_normalized_duplicate_points_are_rejected(content):
    content["key_points"] = ["Same point", "Same POINT！", "Different"]
    with pytest.raises(ValueError, match="duplicate"):
        validate_content(content, required=True)


def test_invalid_prepare_is_read_only_before_creating_jobs(tmp_path):
    path = tmp_path / "spec.json"
    path.write_text(json.dumps({"slides": [{"role": "concept", "key_points": ["alone"]}]}), encoding="utf-8")
    deck = tmp_path / "bad"
    result = invoke("--offline", "prepare", "--spec", path, "--out-dir", deck)
    assert result.returncode != 0 and not deck.exists()


def test_contract_and_prompt_retain_all_approved_content(tmp_path):
    deck = run_pipeline(tmp_path)
    contract = json.loads((deck / CONTRACT).read_text(encoding="utf-8"))
    prompt = json.loads((deck / "prompts/slide_02.json").read_text(encoding="utf-8"))["prompt"]
    assert contract["pages"][1]["content"]["source_refs"] == ["合成测试 fixture"]
    for text in contract["pages"][1]["expected_text"]:
        assert text in prompt


@pytest.mark.parametrize("mutation", ["pending", "missing", "changed", "extra", "ratio", "review", "prompt", "contract"])
def test_assembly_rejects_invalid_input_and_keeps_previous_deck(tmp_path, mutation):
    deck = record_pages(tmp_path)
    if mutation != "review":
        fill_review(deck)
    if mutation == "pending":
        path = deck / "slide_jobs.json"
        jobs = json.loads(path.read_text(encoding="utf-8"))
        jobs["slides"][1]["status"] = "pending"
        path.write_text(json.dumps(jobs), encoding="utf-8")
    elif mutation == "missing":
        (deck / "origin_image/slide_02.png").unlink()
    elif mutation == "changed":
        Image.new("RGB", (1600, 900), "red").save(deck / "origin_image/slide_02.png")
    elif mutation == "extra":
        write_slide_image(deck / "origin_image/slide_03.png")
    elif mutation == "ratio":
        # A recorded, reviewed landscape deck cannot be stretched to 4:3.
        pass
    elif mutation == "prompt":
        path = deck / "prompts/slide_02.json"
        path.write_text(path.read_text(encoding="utf-8") + "\n", encoding="utf-8")
    elif mutation == "contract":
        (deck / CONTRACT).unlink()
    target = deck / f"{deck.name}.pptx"
    target.write_bytes(b"previous approved delivery")
    args = ("--aspect-ratio", "4:3") if mutation == "ratio" else ()
    result = invoke("--offline", "assemble", tmp_path, target.name, *args)
    assert result.returncode != 0, result.stdout
    assert target.read_bytes() == b"previous approved delivery"


def test_missing_text_pending_checks_and_stale_reviews_never_pass(tmp_path):
    deck = record_pages(tmp_path)
    record = deck / REVIEW
    assert review(deck, record, initialize=True)["status"] == "awaiting_review"
    assert review(deck, record)["status"] != "pass"
    record.unlink()
    fill_review(deck)
    data = json.loads(record.read_text(encoding="utf-8"))
    changed = copy.deepcopy(data)
    changed["slides"][1]["observed_text"] = "证据"
    record.write_text(json.dumps(changed), encoding="utf-8")
    assert review(deck, record)["findings"][0]["missing_text"]
    record.write_text(json.dumps(data), encoding="utf-8")
    assert review(deck, record)["status"] == "pass"
    contract = deck / CONTRACT
    contract.write_text(contract.read_text(encoding="utf-8") + "\n", encoding="utf-8")
    with pytest.raises(ValueError, match="Stale"):
        review(deck, record)


def test_lossless_assembly_keeps_exact_image_bytes(tmp_path):
    deck = record_pages(tmp_path)
    fill_review(deck)
    result = invoke("--offline", "assemble", tmp_path, f"{deck.name}.pptx")
    assert result.returncode == 0, result.stderr
    prs = Presentation(deck / f"{deck.name}.pptx")
    for index, slide in enumerate(prs.slides, 1):
        assert slide.shapes[0].image.blob == (deck / f"origin_image/slide_{index:02d}.png").read_bytes()


def test_help_and_content_template_are_usable(tmp_path):
    assert invoke("--help").returncode == 0
    path = tmp_path / "template.json"
    assert invoke("prepare", "--write-template", path).returncode == 0
    for slide in json.loads(path.read_text(encoding="utf-8"))["slides"]:
        validate_content(slide, required=True)


def approved_page_spec(tmp_path):
    root = Path(__file__).resolve().parents[1]
    data = json.loads((root / "examples/page-spec.example.json").read_text(encoding="utf-8"))
    data["content_policy"] = "1.0"
    slide = data["slides"][0]
    slide.update(role="data-evidence", key_points=["合成要点一", "合成要点二", "合成要点三"],
                 explanation="合成示例检查来源和解释的传递", boundary="仅用于管道验证，不代表真实实验",
                 speaker_notes="独立讲稿，不应出现在生图可见文本中", image_status="approved")
    path = tmp_path / "page-spec.json"
    image = tmp_path / slide["image_file"]
    image.parent.mkdir(parents=True)
    write_slide_image(image)
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    return path, data


def test_page_spec_to_generation_preserves_semantic_ids_and_notes(tmp_path):
    path, source = approved_page_spec(tmp_path)
    deck = tmp_path / "converted"
    before = path.read_bytes()
    result = invoke("prepare", "--page-spec", path, "--out-dir", deck)
    assert result.returncode == 0, result.stderr
    contract = json.loads((deck / CONTRACT).read_text(encoding="utf-8"))
    assert contract["pages"][0]["content"]["elements"] == source["slides"][0]["elements"]
    assert source["slides"][0]["speaker_notes"] in (deck / "speech.md").read_text(encoding="utf-8")
    assert source["slides"][0]["speaker_notes"] not in contract["pages"][0]["expected_text"]
    assert path.read_bytes() == before


def test_page_spec_to_rebuilding_prompt_keeps_approved_semantics(tmp_path):
    import subprocess
    import sys
    path, source = approved_page_spec(tmp_path)
    root = Path(__file__).resolve().parents[1]
    run = tmp_path / "rebuilt"
    def rebuild(*args):
        return subprocess.run([sys.executable, str(root / "slidemuse.py"), "rebuild", *map(str, args)],
                              capture_output=True, encoding="utf-8", errors="replace", check=False)
    result = rebuild("--offline", "prepare", "--page-spec", path, "--job-dir", run)
    assert result.returncode == 0, result.stderr
    result = rebuild("page-prompt", run, "--page", "1", "--out", "worker-prompt.md")
    assert result.returncode == 0, result.stderr
    prompt = (run / "pages/page_001/worker-prompt.md").read_text(encoding="utf-8")
    assert source["slides"][0]["boundary"] in prompt
    assert source["slides"][0]["explanation"] in prompt
    for element in source["slides"][0]["elements"]:
        assert element["id"] in prompt
