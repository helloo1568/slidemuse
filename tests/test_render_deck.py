import json
import sys

import pytest
import render_deck
from PIL import Image
from pptx import Presentation


def test_render_deck_writes_every_slide_and_comparison(tmp_path, monkeypatch):
    deck = tmp_path / "deck.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(deck)
    for number, color in enumerate(("white", "black"), 1):
        path = tmp_path / f"source-{number}.png"
        Image.new("RGB", (160, 90), color).save(path)
    spec = tmp_path / "page-spec.json"
    spec.write_text(json.dumps({"slides": [
        {"id": "s01", "image_status": "approved", "image_file": "source-1.png"},
        {"id": "s02", "image_status": "approved", "image_file": "source-2.png"},
    ]}), encoding="utf-8")

    def fake_renderer(_deck, output, width, height):
        for index in range(1, 3):
            Image.new("RGB", (width, height), "white").save(output / f"{index:03d}.png")

    monkeypatch.setattr(render_deck, "render_powerpoint", fake_renderer)
    monkeypatch.setattr(sys, "argv", ["render_deck.py", str(deck), str(tmp_path / "review"),
                                      "--backend", "powerpoint", "--page-spec", str(spec)])
    render_deck.main()
    report = json.loads((tmp_path / "review/render-report.json").read_text(encoding="utf-8"))
    assert len(report["slides"]) == 2
    assert report["slides"][0]["mean_absolute_difference"] == 0
    assert report["slides"][1]["mean_absolute_difference"] == 255
    assert report["slides"][0]["reference_sha256"] == render_deck.sha256(tmp_path / "source-1.png")
    assert report["slides"][1]["reference_sha256"] == render_deck.sha256(tmp_path / "source-2.png")
    assert (tmp_path / "review/review.png").is_file()
    assert all((tmp_path / "review" / f"{index:03d}.png").is_file() for index in (1, 2))


def test_selected_render_keeps_original_index_and_checks_dimensions(tmp_path, monkeypatch):
    deck = tmp_path / "deck.pptx"
    presentation = Presentation()
    for _ in range(3):
        presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(deck)
    calls = []

    def selected(_deck, output, width, height, slides):
        calls.append(slides)
        for index in slides:
            Image.new("RGB", (width, height), "white").save(output / f"{index:03d}.png")

    monkeypatch.setattr(render_deck, "render_powerpoint", selected)
    output = tmp_path / "selected"
    assert render_deck.render_selected(deck, output, 320, "powerpoint", [3]) == "powerpoint"
    assert calls == [[3]]
    assert sorted(p.name for p in output.iterdir()) == ["003.png"]
    for invalid in ([1, 1], [4], [True], []):
        with pytest.raises(ValueError, match="indices"):
            render_deck.render_selected(deck, tmp_path / "bad", 320, "powerpoint", invalid)

    def incorrect(_deck, output, width, height, slides):
        Image.new("RGB", (1, 1)).save(output / "001.png")

    monkeypatch.setattr(render_deck, "render_powerpoint", incorrect)
    with pytest.raises(RuntimeError, match="incorrect dimensions"):
        render_deck.render_selected(deck, tmp_path / "wrong", 320, "powerpoint", [1])
    assert not (tmp_path / "wrong/001.png").exists()
