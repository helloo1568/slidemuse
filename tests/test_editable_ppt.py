import copy
import io
import json
import subprocess
import sys
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from audit_editability import audit
from build_editable_ppt import build_deck
from extract_assets import extract
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from scene import load_scene, validate_scene


@pytest.mark.parametrize("mutation", ["cache_only", "workbook_only", "formula"])
def test_audit_checks_editable_chart_workbook_not_only_display_cache(scene_file, mutation):
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    with ZipFile(output) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    if mutation == "cache_only":
        # Display and Scene agree, but Edit Data would restore the old value.
        scene = json.loads(scene_file.read_text(encoding="utf-8"))
        scene["slides"][0]["elements"][-1]["series"][0]["values"][0] = 99
        scene_file.write_text(json.dumps(scene), encoding="utf-8")
        xml = ET.fromstring(entries["ppt/charts/chart1.xml"])
        ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
        xml.find(".//c:val/c:numRef/c:numCache/c:pt/c:v", ns).text = "99"
        entries["ppt/charts/chart1.xml"] = ET.tostring(xml)
    else:
        name = next(name for name in entries if name.startswith("ppt/embeddings/"))
        with ZipFile(io.BytesIO(entries[name])) as archive:
            workbook = {part: archive.read(part) for part in archive.namelist()}
        xml = ET.fromstring(workbook["xl/worksheets/sheet1.xml"])
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        cell = xml.find(".//s:c[@r='B2']", ns)
        if mutation == "workbook_only":
            cell.find("s:v", ns).text = "99"
        else:
            # A cached formula result is not the literal Scene data contract.
            ET.SubElement(cell, "{" + ns["s"] + "}f").text = "99"
        workbook["xl/worksheets/sheet1.xml"] = ET.tostring(xml)
        buffer = io.BytesIO()
        with ZipFile(buffer, "w") as archive:
            for part, data in workbook.items():
                archive.writestr(part, data)
        entries[name] = buffer.getvalue()
    with ZipFile(output, "w") as archive:
        for part, data in entries.items():
            archive.writestr(part, data)
    errors = audit(output, scene_file)["errors"]
    assert any("embedded chart workbook data differs" in error for error in errors)


def test_audit_checks_all_series_and_unicode_workbook_labels(scene_file):
    scene = json.loads(scene_file.read_text(encoding="utf-8"))
    chart = scene["slides"][0]["elements"][-1]
    chart["categories"] = ["第一期", "第二期"]
    chart["series"].append({"name": "另一组", "values": [-1.5, 0]})
    scene_file.write_text(json.dumps(scene), encoding="utf-8")
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    assert audit(output, scene_file)["errors"] == []


@pytest.fixture
def scene_file(tmp_path):
    Image.new("RGBA", (200, 100), (10, 100, 200, 120)).save(tmp_path / "asset.png")
    scene = {
        "version": "1.0",
        "canvas": {"width": 1600, "height": 900},
        "slides": [
            {
                "id": "s1",
                "notes": "Source: verified test data",
                "elements": [
                    {
                        "id": "title",
                        "type": "text",
                        "x": 100,
                        "y": 40,
                        "w": 1400,
                        "h": 100,
                        "text": "中文与 English\n增长 25%",
                        "font": "Microsoft YaHei",
                        "font_size": 30,
                    },
                    {
                        "id": "photo",
                        "type": "image",
                        "x": 900,
                        "y": 500,
                        "w": 400,
                        "h": 300,
                        "path": "asset.png",
                    },
                    {
                        "id": "flow",
                        "type": "group",
                        "children": [
                            {
                                "id": "node",
                                "type": "shape",
                                "x": 100,
                                "y": 200,
                                "w": 250,
                                "h": 100,
                                "shape": "roundRect",
                                "fill": "#008877",
                            },
                            {
                                "id": "nested",
                                "type": "group",
                                "children": [
                                    {
                                        "id": "label",
                                        "type": "text",
                                        "x": 110,
                                        "y": 220,
                                        "w": 200,
                                        "h": 60,
                                        "text": "独立节点",
                                        "font_size": 20,
                                    }
                                ],
                            },
                            {
                                "id": "arrow",
                                "type": "line",
                                "x1": 350,
                                "y1": 250,
                                "x2": 550,
                                "y2": 250,
                                "arrow": "end",
                            },
                        ],
                    },
                    {
                        "id": "table",
                        "type": "table",
                        "x": 100,
                        "y": 400,
                        "w": 650,
                        "h": 230,
                        "rows": [["阶段", "数值"], ["第一期", "10"]],
                        "column_widths": [2, 1],
                    },
                    {
                        "id": "chart",
                        "type": "chart",
                        "x": 900,
                        "y": 180,
                        "w": 600,
                        "h": 300,
                        "chart_type": "column",
                        "categories": ["A", "B"],
                        "series": [
                            {"name": "增长", "values": [10, 20], "color": "#008877"}
                        ],
                    },
                ],
            }
        ],
    }
    path = tmp_path / "scene.json"
    path.write_text(json.dumps(scene, ensure_ascii=False), encoding="utf-8")
    return path


def test_native_roundtrip(scene_file):
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    report = audit(output, scene_file)
    assert report["errors"] == []
    assert report["warnings"] == []
    assert report["slides"][0]["native_objects"] == 6
    assert report["slides"][0]["movable_raster_objects"] == 1
    prs = Presentation(output)
    assert prs.slides[0].shapes[0].text == "中文与 English\n增长 25%"
    assert (
        prs.slides[0].notes_slide.notes_text_frame.text == "Source: verified test data"
    )
    assert prs.slide_width == Inches(13.333333)
    with ZipFile(output) as package:
        assert any(
            p.startswith("ppt/embeddings/") and p.endswith(".xlsx")
            for p in package.namelist()
        )
        assert b"tailEnd" in package.read("ppt/slides/slide1.xml")


@pytest.mark.parametrize("chart_type", ["bar", "column", "line", "pie", "doughnut"])
def test_chart_types(scene_file, chart_type):
    scene, _ = load_scene(scene_file)
    chart = scene["slides"][0]["elements"][-1]
    chart.update(chart_type=chart_type, data_labels=True, legend=True)
    scene_file.write_text(json.dumps(scene), encoding="utf-8")
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    assert audit(output, scene_file)["errors"] == []


def test_chart_style_roundtrip_and_audit(scene_file):
    scene, _ = load_scene(scene_file)
    chart = scene["slides"][0]["elements"][-1]
    chart.update(
        data_labels=True,
        value_axis_min=0,
        value_axis_max=50,
        major_gridlines=True,
        major_gridline_color="#AABBCC",
        tick_label_color="#123456",
        data_label_color="#654321",
    )
    scene_file.write_text(json.dumps(scene), encoding="utf-8")
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    assert audit(output, scene_file)["errors"] == []

    prs = Presentation(output)
    prs.slides[0].shapes[-1].chart.value_axis.maximum_scale = 60
    prs.save(output)
    assert any(
        "chart style differs" in error for error in audit(output, scene_file)["errors"]
    )


@pytest.mark.parametrize(
    "change",
    [
        {"chart_type": "pie", "value_axis_min": 0},
        {"value_axis_min": 50, "value_axis_max": 0},
        {"major_gridlines": False, "major_gridline_color": "#AABBCC"},
        {"data_label_color": "#AABBCC"},
    ],
)
def test_invalid_chart_style_rejected(scene_file, change):
    scene, _ = load_scene(scene_file)
    scene["slides"][0]["elements"][-1].update(change)
    with pytest.raises(ValueError):
        validate_scene(scene, scene_file.parent)


@pytest.mark.parametrize("fit", ["contain", "cover", "stretch"])
def test_image_fit_preserves_alpha(scene_file, fit):
    scene, _ = load_scene(scene_file)
    scene["slides"][0]["elements"][1]["fit"] = fit
    scene_file.write_text(json.dumps(scene), encoding="utf-8")
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    picture = Presentation(output).slides[0].shapes[1]
    assert audit(output, scene_file)["errors"] == []
    import io

    with Image.open(io.BytesIO(picture.image.blob)) as image:
        assert image.getchannel("A").getextrema() == (120, 120)
    if fit == "cover":
        assert picture.crop_left > 0
    elif fit == "contain":
        assert picture.width / picture.height == pytest.approx(2)


@pytest.mark.parametrize("fit", ["contain", "cover", "stretch"])
@pytest.mark.parametrize("attribute", ["left", "top", "width", "height", "crop_left", "crop_top", "crop_right", "crop_bottom"])
def test_audit_rejects_changed_picture_geometry_and_crop(scene_file, fit, attribute):
    scene, _ = load_scene(scene_file)
    scene["slides"][0]["elements"][1]["fit"] = fit
    scene_file.write_text(json.dumps(scene), encoding="utf-8")
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    prs = Presentation(output)
    picture = prs.slides[0].shapes[1]
    delta = 0.1 if attribute.startswith("crop_") else Inches(0.4)
    setattr(picture, attribute, getattr(picture, attribute) + delta)
    prs.save(output)
    error = "crop differs" if attribute.startswith("crop_") else "geometry differs"
    assert any(error in e for e in audit(output, scene_file)["errors"])


@pytest.mark.parametrize("nested", [False, True])
@pytest.mark.parametrize("mutation", ["move", "resize", "rotate", "child_offset", "child_extent", "flip"])
def test_audit_rejects_group_transform_edits(scene_file, nested, mutation):
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    prs = Presentation(output)
    group = prs.slides[0].shapes[2]
    if nested:
        group = group.shapes[1]
    transform = group._element.grpSpPr.xfrm
    if mutation == "move":
        group.left += Inches(1)
    elif mutation == "resize":
        group.width += Inches(1)
    elif mutation == "rotate":
        group.rotation = 90
    elif mutation == "child_offset":
        transform.chOff.x += Inches(1)
    elif mutation == "child_extent":
        transform.chExt.cy += Inches(1)
    elif mutation == "flip":
        transform.flipH = True
    prs.save(output)
    error = "rotation differs" if mutation == "rotate" else "group transform differs"
    assert any(error in e for e in audit(output, scene_file)["errors"])


@pytest.mark.parametrize("fit", ["contain", "cover", "stretch"])
def test_audit_accepts_image_inside_nested_group(scene_file, fit):
    scene, _ = load_scene(scene_file)
    elements = scene["slides"][0]["elements"]
    photo = elements.pop(1)
    photo["fit"] = fit
    elements[1]["children"][1]["children"].append(photo)
    scene_file.write_text(json.dumps(scene), encoding="utf-8")
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    assert audit(output, scene_file)["errors"] == []


@pytest.mark.parametrize("axis", ["flipH", "flipV"])
def test_audit_rejects_picture_flip(scene_file, axis):
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    prs = Presentation(output)
    setattr(prs.slides[0].shapes[1]._element.spPr.xfrm, axis, True)
    prs.save(output)
    assert any("image flip differs" in e for e in audit(output, scene_file)["errors"])


@pytest.mark.parametrize(
    "mutation",
    [
        "duplicate",
        "outside",
        "unknown",
        "nan",
        "ragged",
        "values",
        "traversal",
        "missing",
        "zero_line",
        "negative_pie",
    ],
)
def test_invalid_scene_rejected(scene_file, mutation):
    scene, _ = load_scene(scene_file)
    els = scene["slides"][0]["elements"]
    if mutation == "duplicate":
        els.append(copy.deepcopy(els[0]))
    elif mutation == "outside":
        els[0]["x"] = 1500
    elif mutation == "unknown":
        els[0]["typo"] = True
    elif mutation == "nan":
        els[0]["w"] = float("nan")
    elif mutation == "ragged":
        els[-2]["rows"].append(["bad"])
    elif mutation == "values":
        els[-1]["series"][0]["values"] = [10]
    elif mutation == "traversal":
        els[1]["path"] = "../asset.png"
    elif mutation == "missing":
        els[1]["path"] = "missing.png"
    elif mutation == "zero_line":
        els[2]["children"][-1]["x2"] = 350
    elif mutation == "negative_pie":
        els[-1]["chart_type"] = "pie"
        els[-1]["series"][0]["values"] = [-1, 2]
    with pytest.raises(ValueError):
        validate_scene(scene, scene_file.parent)


def test_failed_build_keeps_existing_output(scene_file):
    output = scene_file.with_suffix(".pptx")
    output.write_bytes(b"existing deck")
    scene, _ = load_scene(scene_file)
    scene["slides"][0]["elements"][1]["path"] = "missing.png"
    scene_file.write_text(json.dumps(scene), encoding="utf-8")
    with pytest.raises(ValueError):
        build_deck(scene_file, output)
    assert output.read_bytes() == b"existing deck"


def test_audit_detects_edits_and_missing_objects(scene_file):
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    prs = Presentation(output)
    prs.slides[0].shapes[0].text = "Changed"
    node = prs.slides[0].shapes[1]._element
    node.getparent().remove(node)
    prs.save(output)
    errors = audit(output, scene_file)["errors"]
    assert any("text differs" in e for e in errors)
    assert any("missing object" in e for e in errors)


def test_reference_image_rejected_even_after_reencoding(scene_file):
    scene, _ = load_scene(scene_file)
    scene["slides"][0]["source_image"] = "asset.png"
    scene_file.write_text(json.dumps(scene), encoding="utf-8")
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    assert any("reference pixels" in e for e in audit(output, scene_file)["errors"])


def test_extract_masks_and_bounds(tmp_path):
    Image.new("RGBA", (100, 60), (250, 0, 0, 128)).save(tmp_path / "source.png")
    Image.new("L", (100, 60), 128).save(tmp_path / "mask.png")
    path = tmp_path / "extract.json"
    spec = {
        "source": "source.png",
        "regions": [{"id": "item", "box": [10.2, 5.3, 30, 20], "mask": "mask.png"}],
    }
    path.write_text(json.dumps(spec), encoding="utf-8")
    result = extract(path, tmp_path / "assets")
    assert result["assets"][0]["source_box"] == [10, 5, 31, 21]
    with Image.open(tmp_path / "assets/item.png") as im:
        assert im.getchannel("A").getextrema() == (64, 64)
    with pytest.raises(ValueError, match="exists"):
        extract(path, tmp_path / "assets")
    spec["regions"][0]["box"] = [90, 50, 30, 20]
    path.write_text(json.dumps(spec), encoding="utf-8")
    with pytest.raises(ValueError, match="outside"):
        extract(path, tmp_path / "bad")
    assert not (tmp_path / "bad").exists()


def test_cli_from_other_working_directory(scene_file, tmp_path):
    script = Path(__file__).resolve().parents[1] / "scripts/build_editable_ppt.py"
    output = tmp_path / "space folder" / "result.pptx"
    result = subprocess.run(
        [sys.executable, str(script), str(scene_file), str(output)],
        cwd=tmp_path,
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    assert audit(output, scene_file)["errors"] == []


def test_native_data_can_be_changed_after_export(scene_file):
    from pptx.chart.data import CategoryChartData

    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    prs = Presentation(output)
    slide = prs.slides[0]
    slide.shapes[0].left += Inches(0.2)
    slide.shapes[-2].table.cell(1, 1).text = "99"
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("增长", [99, 100])
    slide.shapes[-1].chart.replace_data(data)
    prs.save(output)
    reopened = Presentation(output).slides[0]
    assert list(reopened.shapes[-1].chart.series[0].values) == [99, 100]
    assert reopened.shapes[-2].table.cell(1, 1).text == "99"
    errors = audit(output, scene_file)["errors"]
    assert any("chart data differs" in e for e in errors)
    assert any("table content differs" in e for e in errors)
    assert any("geometry differs" in e for e in errors)


def test_negative_point_colors_and_plot_layout_survive_export_and_are_audited(scene_file):
    from chart_style import NS
    from pptx.enum.chart import XL_LABEL_POSITION
    scene, _ = load_scene(scene_file)
    chart = scene["slides"][0]["elements"][-1]
    chart.update(chart_type="bar", category_axis_visible=False, value_axis_visible=False,
                 category_reverse_order=True, gap_width=78, data_labels=True,
                 data_label_position="inside_end", data_label_font_size=18,
                 plot_layout={"x": 0, "y": 0, "w": 0.9, "h": 1})
    chart["series"][0].update(values=[3.2, -5.5], point_colors=["#9582DE", "#B8A8E2"], invert_if_negative=False)
    scene["slides"][0]["elements"][2]["children"][0]["corner_radius"] = 0.04
    scene_file.write_text(json.dumps(scene), encoding="utf-8")
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    assert audit(output, scene_file)["errors"] == []
    prs = Presentation(output)
    native = prs.slides[0].shapes[-1].chart
    assert native.plots[0].data_labels.position == XL_LABEL_POSITION.INSIDE_END
    point = native.series[0].points[1]
    assert str(point.format.fill.fore_color.rgb) == "B8A8E2"
    assert point.format._element.find(NS + "invertIfNegative").get("val") == "0"
    point.format._element.find(NS + "invertIfNegative").set("val", "1")
    prs.save(output)
    assert any("point 1 negative fill" in error for error in audit(output, scene_file)["errors"])


def test_doughnut_hole_angle_and_colors_are_scene_settings(scene_file):
    from chart_style import NS
    scene, _ = load_scene(scene_file)
    chart = scene["slides"][0]["elements"][-1]
    chart.update(chart_type="doughnut", hole_size=70, first_slice_angle=20)
    chart["series"][0]["point_colors"] = ["#9582DE", "#DCD9EA"]
    scene_file.write_text(json.dumps(scene), encoding="utf-8")
    output = scene_file.with_suffix(".pptx")
    build_deck(scene_file, output)
    assert audit(output, scene_file)["errors"] == []
    prs = Presentation(output)
    plot = prs.slides[0].shapes[-1].chart.plots[0]
    assert plot._element.find(NS + "holeSize").get("val") == "70"
    plot._element.find(NS + "holeSize").set("val", "30")
    prs.save(output)
    assert any("hole_size" in error for error in audit(output, scene_file)["errors"])


@pytest.mark.parametrize("change", [
    {"hole_size": 70}, {"first_slice_angle": 10}, {"gap_width": 600},
    {"plot_layout": {"x": 0.5, "y": 0, "w": 0.8, "h": 1}},
    {"data_labels": False, "data_label_position": "center"},
    {"series": [{"name": "bad", "values": [1, 2], "point_colors": ["#123456"]}]},
])
def test_invalid_native_chart_settings_are_rejected(scene_file, change):
    scene, _ = load_scene(scene_file)
    scene["slides"][0]["elements"][-1].update(change)
    with pytest.raises(ValueError):
        validate_scene(scene, scene_file.parent)
