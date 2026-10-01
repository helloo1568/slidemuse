"""Adversarial evidence tests, not benchmark cases or visual acceptance.

PPTX objects are really built and audited. PNGs are deliberately synthetic;
the renderer marker only exercises report validation without Office installed.
"""

import copy
import json
import subprocess
import sys
from argparse import Namespace
from pathlib import Path

import build_editable_ppt
import build_image_ppt
import evaluate_benchmark as benchmark
import evaluate_delivery
import pytest
from PIL import Image
from pptx import Presentation
from test_evaluate_delivery import make_case, mark_visual_pass

COMMIT = "a" * 40


def save(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def load(path):
    return json.loads(path.read_text(encoding="utf-8"))


@pytest.mark.parametrize("scope", ["formal", "exercise"])
@pytest.mark.parametrize("mutation", ["changed", "removed"])
def test_source_image_dependencies_cannot_change_after_sealing(tmp_path, scope, mutation):
    batch_path, frozen, _ = frozen_batch(tmp_path, count=1)
    attempt = delivery_fixture(tmp_path)
    source = tmp_path / "reference-source.png"
    Image.new("RGB", (160, 90), "red").save(source)
    if scope == "formal":
        paths = [Path(attempt["deliveries"]["editable"]["scene"])]
    else:
        checks = load(Path(attempt["records"]["edit_check"]))
        paths = [Path(operation["after_scene"]) for operation in checks["operations"]]
    for path in paths:
        scene = load(path)
        scene["slides"][0]["source_image"] = source.name
        save(path, scene)
    approve(frozen, attempt, tmp_path)
    attempt_path = tmp_path / "attempt.json"
    save(attempt_path, attempt)
    record = benchmark.record_attempt(batch_path, attempt_path, tmp_path / "history")
    assert record["assessment"]["status"] == "pass"
    assert str(source.resolve()) in record["assessment"]["files"]
    if mutation == "changed":
        Image.new("RGB", (160, 90), "blue").save(source)
    else:
        source.unlink()
    report = benchmark.summarize(batch_path, tmp_path / "history")
    assert report["cases"][0]["first_recorded_status"] == "pass"
    assert report["cases"][0]["first"]["status"] == "incomplete"


def test_review_contract_excludes_work_judgments_and_nested_unknown_data(tmp_path):
    spec_path, *_ = make_case(tmp_path, "chart")
    spec = load(spec_path)
    page = spec["slides"][0]
    page.update(generation_prompt="SECRET_GENERATION", notes="SECRET_REPAIR",
                speaker_notes="Visible declared presenter content")
    page["elements"][0].update(prompt_record="SECRET_PROMPT_RECORD", notes="SECRET_AUTHOR_PASS")
    page["elements"][-1]["data"] = {
        "categories": ["A", "B"], "series": [{"name": "系列", "values": [12, 15],
                                             "author_review": {"notes": "SECRET_NESTED"}}],
        "unit": "km", "author_review": {"verdict": "SECRET_DATA_PASS"}}
    save(spec_path, spec)
    output = tmp_path / "contract.json"
    record = benchmark.create_review_contract(spec_path, output)
    assert "SECRET" not in output.read_text(encoding="utf-8")
    assert record["contract"]["omitted_data_fields"] == 2
    data = record["contract"]["slides"][0]["elements"][-1]["data"]
    assert data == {"categories": ["A", "B"], "series": [{"name": "系列", "values": [12, 15]}], "unit": "km"}
    assert record["contract"]["slides"][0]["speaker_notes"] == "Visible declared presenter content"
    assert benchmark.verify_review_contract(output, spec_path) == record
    with pytest.raises(FileExistsError):
        benchmark.create_review_contract(spec_path, output)
    with pytest.raises(ValueError, match="must not replace"):
        benchmark.create_review_contract(spec_path, spec_path)


@pytest.mark.parametrize("mutation", ["source", "projection", "extra_field"])
def test_review_contract_cannot_be_resealed_with_changed_or_extra_content(tmp_path, mutation):
    spec_path, *_ = make_case(tmp_path)
    output = tmp_path / "contract.json"
    benchmark.create_review_contract(spec_path, output)
    if mutation == "source":
        spec = load(spec_path)
        spec["slides"][0]["elements"][0]["text"] = "Different declared content"
        save(spec_path, spec)
    else:
        record = load(output)
        record.pop("record_sha256")
        if mutation == "extra_field":
            record["generation_prompt"] = "Unexpected private instruction"
        else:
            record["contract"]["slides"][0]["elements"][0]["text"] = "Different content"
            record["projection_sha256"] = benchmark.object_digest(record["contract"])
        record["record_sha256"] = benchmark.object_digest(record)
        save(output, record)
    with pytest.raises(ValueError, match="differs"):
        benchmark.verify_review_contract(output, spec_path)


def test_review_contract_rejects_opaque_data_and_objects_hidden_in_cells(tmp_path):
    spec_path, *_ = make_case(tmp_path, "table")
    spec = load(spec_path)
    spec["slides"][0]["elements"][-1]["data"] = {"private_result": "SECRET"}
    save(spec_path, spec)
    with pytest.raises(ValueError, match="no supported"):
        benchmark.create_review_contract(spec_path, tmp_path / "contract.json")
    spec["slides"][0]["elements"][-1]["data"] = {"rows": [[{"notes": "SECRET"}]]}
    save(spec_path, spec)
    with pytest.raises(ValueError, match="scalar cells"):
        benchmark.create_review_contract(spec_path, tmp_path / "contract.json")
    assert not (tmp_path / "contract.json").exists()


def registry(tmp_path, count=2, scope="reconstruction"):
    rubric = tmp_path / "rubric.md"
    rubric.write_text("Actual reviewer must check all five dimensions.", encoding="utf-8")
    value = {"version": "1.0", "candidate_commit": COMMIT, "rubric": str(rubric),
             "cases": [{"id": f"fixture-{i}", "category": "reconstruction",
                        "split": "development", "source_id": f"fixture:source-{i}",
                        "delivery_scope": scope} for i in range(count)]}
    path = tmp_path / "registry.json"
    save(path, value)
    return path, value


def frozen_batch(tmp_path, **kwargs):
    path, value = registry(tmp_path, **kwargs)
    output = tmp_path / "frozen.json"
    return output, benchmark.freeze(path, output), value


def render_fixture(deck, output, reference, spec=None):
    rendered = output.with_suffix(".png")
    Image.new("RGB", (160, 90), "white").save(rendered)
    report = {"deck": str(deck), "deck_sha256": benchmark.digest(deck),
              "backend": "powerpoint",  # Test marker; no actual rendering claim.
              "slides": [{"slide": 1, "rendered": str(rendered),
                          "rendered_sha256": benchmark.digest(rendered),
                          "reference": str(reference), "reference_sha256": benchmark.digest(reference)}]}
    if spec:
        report["page_spec_sha256"] = benchmark.digest(spec)
    save(output, report)
    return report


def delivery_fixture(tmp_path):
    spec_path, render_path, review_path, observations, _ = make_case(tmp_path)
    spec = load(spec_path)
    reference = tmp_path / spec["slides"][0]["image_file"]
    scene = {"version": "1.0", "canvas": spec["canvas"], "slides": [{"id": "s01", "elements": [
        {"id": "title", "type": "text", "text": spec["slides"][0]["elements"][0]["text"],
         "x": 100, "y": 100, "w": 700, "h": 100},
        {"id": "tagline", "type": "text", "text": spec["slides"][0]["elements"][1]["text"],
         "x": 100, "y": 300, "w": 700, "h": 100},
        {"id": "agronomist", "type": "image", "path": spec["slides"][0]["image_file"],
         "x": 1000, "y": 100, "w": 300, "h": 300}]}]}
    scene_path = tmp_path / "scene.json"
    save(scene_path, scene)
    deck = tmp_path / "deck.pptx"
    build_editable_ppt.build_deck(scene_path, deck, spec_path)
    report = render_fixture(deck, render_path, reference, spec_path)
    save(review_path, evaluate_delivery.review_template(spec, report))
    mark_visual_pass(review_path)
    operations = []
    for index, kind in enumerate(("text_edit", "object_move")):
        changed = copy.deepcopy(scene)
        target = changed["slides"][0]["elements"][0]
        target["text" if kind == "text_edit" else "x"] = "Actual changed text" if kind == "text_edit" else 200
        after_scene = tmp_path / f"edited-{index}.json"
        after_deck = tmp_path / f"edited-{index}.pptx"
        after_render = tmp_path / f"edited-{index}-render.json"
        save(after_scene, changed)
        build_editable_ppt.build_deck(after_scene, after_deck)
        render_fixture(after_deck, after_render, reference)
        operations.append({"kind": kind, "status": "pass", "notes": "Synthetic exercise, actual OOXML change",
                           "slide_id": "s01", "element_id": "title", "before_scene": str(scene_path),
                           "after_scene": str(after_scene), "after_deck": str(after_deck),
                           "render_report": str(after_render)})
    source = tmp_path / "source.txt"
    source.write_text("Synthetic source for evidence tests; not a real benchmark case.", encoding="utf-8")
    draft = tmp_path / "content.md"
    draft.write_text("Synthetic content draft", encoding="utf-8")
    authorizations = tmp_path / "authorization.json"
    save(authorizations, {"user_quote": "Fixture authorization, not a human consent record",
                          "stages": {"content": True, "style": True, "editable": True}})
    run_log = tmp_path / "run-log.json"
    save(run_log, {"complete": True, "elapsed_seconds": 1,
                   "events": [{"stage": "render", "outcome": "pass", "evidence": str(render_path)}]})
    edit_check = tmp_path / "edit-check.json"
    save(edit_check, {"operations": operations})
    attempt = {"case_id": "fixture-0", "candidate_commit": COMMIT, "execution_status": "delivered",
               "producer_id": "fixture-producer", "sources": [str(source)],
               "records": {"content_draft": str(draft), "authorizations": str(authorizations),
                           "run_log": str(run_log), "edit_check": str(edit_check)},
               "deliveries": {"editable": {"page_spec": str(spec_path), "render_report": str(render_path),
                                           "visual_review": str(review_path), "observations": [str(observations)],
                                           "scene": str(scene_path)}}}
    return attempt


def approve(frozen, attempt, base, *, dimension_status="pass", errors=0):
    assessment = benchmark.assess(frozen, attempt, base)
    assert "review_input_sha256" in assessment, assessment["issues"]
    path = base / "independent-review.json"
    save(path, {"version": "1.0", "input_sha256": assessment["review_input_sha256"],
                "reviewer_id": "fixture-reviewer", "independent": True,
                "prior_conclusions_exposed": False,
                "major_factual_errors": errors,
                "dimensions": {name: {"status": dimension_status, "notes": "Synthetic judgment for test only"}
                               for name in benchmark.DIMENSIONS}})
    attempt["independent_review"] = str(path)
    return path


def test_freeze_cannot_replace_existing_batch(tmp_path):
    registry_path, _ = registry(tmp_path)
    frozen = tmp_path / "frozen.json"
    benchmark.freeze(registry_path, frozen)
    original = frozen.read_bytes()
    with pytest.raises(FileExistsError):
        benchmark.freeze(registry_path, frozen)
    assert frozen.read_bytes() == original


@pytest.mark.parametrize("identities", [
    ("https://arxiv.org/abs/1706.03762v1", "https://arxiv.org/html/1706.03762v7"),
    ("https://doi.org/10.1234/Test", "doi:10.1234/test"),
    ("https://example.org/report#one", "https://example.org/report#two"),
])
def test_versions_and_locations_of_one_source_are_not_independent_cases(tmp_path, identities):
    path, value = registry(tmp_path)
    for case, source_id in zip(value["cases"], identities):
        case["source_id"] = source_id
    save(path, value)
    with pytest.raises(ValueError, match="Duplicate material"):
        benchmark.freeze(path, tmp_path / "frozen.json")


def test_explicit_source_alias_overlap_is_rejected(tmp_path):
    path, value = registry(tmp_path)
    value["cases"][1]["source_aliases"] = [value["cases"][0]["source_id"]]
    save(path, value)
    with pytest.raises(ValueError, match="Duplicate material"):
        benchmark.freeze(path, tmp_path / "frozen.json")


def test_unrun_and_aborted_cases_stay_in_denominator(tmp_path):
    batch_path, _, _ = frozen_batch(tmp_path)
    attempt_path = tmp_path / "attempt.json"
    save(attempt_path, {"case_id": "fixture-0", "candidate_commit": COMMIT, "execution_status": "aborted"})
    benchmark.record_attempt(batch_path, attempt_path, tmp_path / "history")
    report = benchmark.summarize(batch_path, tmp_path / "history")
    assert report["overall"] == {"total": 2, "first_passed": 0, "first_usable_rate": 0, "final_passed": 0}
    assert report["cases"][1]["first_recorded_status"] == "not_run"
    assert not report["ultimate_goal_checks_passed"]


@pytest.mark.parametrize("tamper", ["rubric", "registry"])
def test_frozen_inputs_cannot_be_silently_rewritten(tmp_path, tamper):
    batch_path, _, value = frozen_batch(tmp_path)
    if tamper == "rubric":
        Path(value["rubric"]).write_text("Changed acceptance after execution", encoding="utf-8")
    else:
        changed = load(batch_path)
        changed["registry"]["cases"].pop()
        save(batch_path, changed)
    with pytest.raises(ValueError, match="rubric changed|changed or is unsealed"):
        benchmark.summarize(batch_path, tmp_path / "history")


def test_holdout_requires_explicit_no_tuning_and_clean_execution(tmp_path):
    path, value = registry(tmp_path, count=1)
    value["cases"][0]["split"] = "holdout"
    save(path, value)
    with pytest.raises(ValueError, match="unused-for-tuning"):
        benchmark.freeze(path, tmp_path / "frozen.json")
    value["cases"][0]["used_for_tuning"] = False
    save(path, value)
    frozen = benchmark.freeze(path, tmp_path / "frozen.json")
    assessed = benchmark.assess(frozen, {"case_id": "fixture-0", "candidate_commit": COMMIT}, tmp_path)
    assert assessed["status"] == "incomplete"
    assert "Holdout" in assessed["issues"][0]


def test_first_failure_cannot_be_replaced_by_later_success(tmp_path):
    batch_path, frozen, _ = frozen_batch(tmp_path)
    attempt_path = tmp_path / "attempt.json"
    save(attempt_path, {"case_id": "fixture-0", "candidate_commit": COMMIT, "execution_status": "aborted"})
    benchmark.record_attempt(batch_path, attempt_path, tmp_path / "history")
    attempt = delivery_fixture(tmp_path)
    approve(frozen, attempt, tmp_path)
    save(attempt_path, attempt)
    benchmark.record_attempt(batch_path, attempt_path, tmp_path / "history")
    report = benchmark.summarize(batch_path, tmp_path / "history")
    assert report["overall"]["first_passed"] == 0
    assert report["overall"]["final_passed"] == 1
    assert report["cases"][0]["first"]["status"] == "incomplete"
    assert report["cases"][0]["final"]["status"] == "pass"


def test_structurally_valid_delivery_requires_actual_semantic_review_record(tmp_path):
    _, frozen, _ = frozen_batch(tmp_path)
    attempt = delivery_fixture(tmp_path)
    assessment = benchmark.assess(frozen, attempt, tmp_path)
    assert assessment["deliveries"]["editable"]["status"] == "pass"
    assert assessment["status"] == "incomplete"
    approve(frozen, attempt, tmp_path)
    assert benchmark.assess(frozen, attempt, tmp_path)["status"] == "pass"


@pytest.mark.parametrize("change", ["same-reviewer", "missing-dimension", "wrong-binding", "empty-explanation"])
def test_invalid_semantic_review_never_counts_as_pass(tmp_path, change):
    _, frozen, _ = frozen_batch(tmp_path)
    attempt = delivery_fixture(tmp_path)
    path = approve(frozen, attempt, tmp_path)
    review = load(path)
    if change == "same-reviewer":
        review["reviewer_id"] = " FIXTURE-PRODUCER "
    elif change == "missing-dimension":
        review["dimensions"].pop("content")
    elif change == "wrong-binding":
        review["input_sha256"] = "0" * 64
    else:
        review["dimensions"]["content"]["notes"] = "  "
    save(path, review)
    assert benchmark.assess(frozen, attempt, tmp_path)["status"] == "incomplete"


@pytest.mark.parametrize("dimension_status,errors,expected", [("fail", 0, "fail"), ("pass", 1, "fail"),
                                                            ("incomplete", 0, "incomplete")])
def test_independent_failures_and_factual_errors_override_machine_pass(tmp_path, dimension_status, errors, expected):
    _, frozen, _ = frozen_batch(tmp_path)
    attempt = delivery_fixture(tmp_path)
    approve(frozen, attempt, tmp_path, dimension_status=dimension_status, errors=errors)
    assert benchmark.assess(frozen, attempt, tmp_path)["status"] == expected


@pytest.mark.parametrize("exposure", ["missing", True])
@pytest.mark.parametrize("errors", [0, 1])
def test_review_exposed_to_author_conclusions_cannot_pass_but_preserves_factual_failure(tmp_path, exposure, errors):
    _, frozen, _ = frozen_batch(tmp_path)
    attempt = delivery_fixture(tmp_path)
    path = approve(frozen, attempt, tmp_path, errors=errors)
    review = load(path)
    if exposure == "missing":
        review.pop("prior_conclusions_exposed")
    else:
        review["prior_conclusions_exposed"] = True
    save(path, review)
    assessed = benchmark.assess(frozen, attempt, tmp_path)
    assert assessed["status"] == ("fail" if errors else "incomplete")
    assert assessed["major_factual_errors"] == errors
    assert any("prior conclusions" in issue for issue in assessed["issues"])


def test_changed_additional_reproduction_evidence_revokes_pass(tmp_path):
    batch_path, frozen, _ = frozen_batch(tmp_path)
    attempt = delivery_fixture(tmp_path)
    reproduction = tmp_path / "reproduction.py"
    reproduction.write_text("# Synthetic reproduction evidence, not an actual installation report.\n", encoding="utf-8")
    attempt["records"]["additional_evidence"] = [str(reproduction)]
    approve(frozen, attempt, tmp_path)
    attempt_path = tmp_path / "attempt.json"
    save(attempt_path, attempt)
    history = tmp_path / "history"
    recorded = benchmark.record_attempt(batch_path, attempt_path, history)
    assert recorded["assessment"]["status"] == "pass"
    assert recorded["assessment"]["files"][str(reproduction)] == benchmark.digest(reproduction)
    reproduction.write_text("# Changed the reproduction script after approval.\n", encoding="utf-8")
    report = benchmark.summarize(batch_path, history)
    row = report["cases"][0]
    assert row["first_recorded_status"] == "pass"
    assert row["first"]["status"] == row["final"]["status"] == "incomplete"
    assert str(reproduction) in row["first"]["changed_files"]
    assert report["overall"]["first_passed"] == report["overall"]["final_passed"] == 0


@pytest.mark.parametrize("changed_record", ["source", "rendered", "review"])
def test_changed_recorded_evidence_revokes_current_pass_but_preserves_original(tmp_path, changed_record):
    batch_path, frozen, _ = frozen_batch(tmp_path)
    attempt = delivery_fixture(tmp_path)
    approve(frozen, attempt, tmp_path)
    attempt_path = tmp_path / "attempt.json"
    save(attempt_path, attempt)
    history = tmp_path / "history"
    recorded = benchmark.record_attempt(batch_path, attempt_path, history)
    assert recorded["assessment"]["status"] == "pass"
    if changed_record == "source":
        Path(attempt["sources"][0]).write_text("Different source snapshot", encoding="utf-8")
    elif changed_record == "rendered":
        Image.new("RGB", (160, 90), "black").save(tmp_path / "render-report.png")
    else:
        Path(attempt["independent_review"]).write_text("{}", encoding="utf-8")
    report = benchmark.summarize(batch_path, history)
    assert report["overall"]["first_passed"] == 0
    assert report["cases"][0]["first_recorded_status"] == "pass"
    assert report["cases"][0]["first"]["status"] == "incomplete"
    assert benchmark.sealed_read(history / "attempt-000001.json")["assessment"]["status"] == "pass"


@pytest.mark.parametrize("change", ["unchanged-scene", "unchanged-pptx", "missing-move", "stale-render"])
def test_edit_exercises_require_declared_scene_change_and_corresponding_artifacts(tmp_path, change):
    _, frozen, _ = frozen_batch(tmp_path)
    attempt = delivery_fixture(tmp_path)
    path = Path(attempt["records"]["edit_check"])
    checks = load(path)
    operation = checks["operations"][0]
    if change == "unchanged-scene":
        operation["after_scene"] = operation["before_scene"]
    elif change == "unchanged-pptx":
        operation["after_deck"] = str(tmp_path / "deck.pptx")
    elif change == "missing-move":
        checks["operations"].pop()
    else:
        report_path = Path(operation["render_report"])
        report = load(report_path)
        report["deck_sha256"] = "0" * 64
        save(report_path, report)
    save(path, checks)
    assessed = benchmark.assess(frozen, attempt, tmp_path)
    assert assessed["status"] == "incomplete"
    assert not assessed["deliveries"]


def test_deleted_attempt_in_middle_of_chain_is_rejected(tmp_path):
    batch_path, _, _ = frozen_batch(tmp_path)
    path = tmp_path / "attempt.json"
    save(path, {"case_id": "fixture-0", "candidate_commit": COMMIT, "execution_status": "aborted"})
    history = tmp_path / "history"
    for _ in range(3):
        benchmark.record_attempt(batch_path, path, history)
    (history / "attempt-000002.json").unlink()
    with pytest.raises(ValueError, match="history is missing"):
        benchmark.summarize(batch_path, history)


def test_deleted_attempt_at_tail_of_chain_is_rejected(tmp_path):
    batch_path, _, _ = frozen_batch(tmp_path)
    path = tmp_path / "attempt.json"
    save(path, {"case_id": "fixture-0", "candidate_commit": COMMIT, "execution_status": "aborted"})
    history = tmp_path / "history"
    for _ in range(2):
        benchmark.record_attempt(batch_path, path, history)
    (history / "attempt-000002.json").unlink()
    with pytest.raises(ValueError, match="history tail is missing"):
        benchmark.summarize(batch_path, history)


def test_late_semantic_review_requires_new_attempt_and_cannot_promote_first(tmp_path):
    batch_path, frozen, _ = frozen_batch(tmp_path)
    attempt = delivery_fixture(tmp_path)
    review_path = tmp_path / "independent-review.json"
    attempt["independent_review"] = str(review_path)
    path = tmp_path / "attempt.json"
    save(path, attempt)
    history = tmp_path / "history"
    benchmark.record_attempt(batch_path, path, history)
    approve(frozen, attempt, tmp_path)
    report = benchmark.summarize(batch_path, history)
    assert report["overall"]["first_passed"] == report["overall"]["final_passed"] == 0
    benchmark.record_attempt(batch_path, path, history)
    report = benchmark.summarize(batch_path, history)
    assert report["overall"]["first_passed"] == 0
    assert report["overall"]["final_passed"] == 1


def test_edit_exercises_cannot_use_an_unrelated_starting_scene(tmp_path):
    _, frozen, _ = frozen_batch(tmp_path)
    attempt = delivery_fixture(tmp_path)
    checks_path = Path(attempt["records"]["edit_check"])
    checks = load(checks_path)
    unrelated = load(tmp_path / "scene.json")
    unrelated["slides"][0]["elements"][1]["text"] = "Different delivery"
    unrelated_path = tmp_path / "unrelated.json"
    save(unrelated_path, unrelated)
    checks["operations"][0]["before_scene"] = str(unrelated_path)
    save(checks_path, checks)
    assessment = benchmark.assess(frozen, attempt, tmp_path)
    assert assessment["status"] == "incomplete"
    assert "delivered Scene" in assessment["issues"][0]


def test_run_ledger_cannot_use_a_nonexistent_evidence_locator(tmp_path):
    _, frozen, _ = frozen_batch(tmp_path)
    attempt = delivery_fixture(tmp_path)
    path = Path(attempt["records"]["run_log"])
    log = load(path)
    log["events"][0]["evidence"] = "made-up-invocation.json"
    save(path, log)
    assessment = benchmark.assess(frozen, attempt, tmp_path)
    assert assessment["status"] == "incomplete"
    assert "made-up-invocation.json" in assessment["issues"][0]


def test_image_workflow_cannot_skip_four_style_trials(tmp_path):
    _, frozen, _ = frozen_batch(tmp_path, scope="image_and_editable")
    attempt = delivery_fixture(tmp_path)
    attempt["deliveries"]["image"] = {key: value for key, value in attempt["deliveries"]["editable"].items()
                                       if key != "scene"}
    assessment = benchmark.assess(frozen, attempt, tmp_path)
    assert assessment["status"] == "incomplete"
    assert "four style calls" in assessment["issues"][0]


def test_attempt_from_different_candidate_is_not_scored_as_frozen_candidate(tmp_path):
    _, frozen, _ = frozen_batch(tmp_path)
    assessed = benchmark.assess(frozen, {"case_id": "fixture-0", "candidate_commit": "b" * 40,
                                         "execution_status": "delivered", "producer_id": "fixture"}, tmp_path)
    assert assessed["status"] == "incomplete"
    assert assessed["issues"] == ["Candidate differs from frozen commit"]


def test_eligible_case_count_does_not_establish_ultimate_goal_success(tmp_path):
    path, value = registry(tmp_path, count=30)
    for index, case in enumerate(value["cases"]):
        case["category"] = benchmark.CATEGORIES[index // 5]
        if index < 10:
            case.update(split="holdout", used_for_tuning=False)
    save(path, value)
    batch_path = tmp_path / "frozen.json"
    benchmark.freeze(path, batch_path)
    report = benchmark.summarize(batch_path, tmp_path / "history")
    assert report["ultimate_goal_eligible_batch"]
    assert report["overall"]["total"] == 30
    assert report["overall"]["first_usable_rate"] == 0
    assert not report["ultimate_goal_checks_passed"]


@pytest.mark.parametrize("kind", ["chart", "table"])
def test_declared_data_visuals_require_data_edit_exercise(tmp_path, kind):
    _, frozen, _ = frozen_batch(tmp_path)
    attempt = delivery_fixture(tmp_path)
    item = attempt["deliveries"]["editable"]
    spec_path = Path(item["page_spec"])
    spec = load(spec_path)
    spec["slides"][0]["elements"].append({"id": "results", "kind": kind, "role": "test data",
                                        "native_intent": "native", "source_ref": "fixture/source",
                                        "confirmation_status": "source-verified", "data": {"values": [10, 20]}})
    save(spec_path, spec)
    report_path = Path(item["render_report"])
    report = load(report_path)
    report["page_spec_sha256"] = benchmark.digest(spec_path)
    save(report_path, report)
    review_path = Path(item["visual_review"])
    save(review_path, evaluate_delivery.review_template(spec, report))
    mark_visual_pass(review_path)
    assessment = benchmark.assess(frozen, attempt, tmp_path)
    assert assessment["status"] == "incomplete"
    assert "chart/table data edits" in assessment["issues"][0]


@pytest.mark.parametrize("kind", ["chart", "table"])
@pytest.mark.parametrize("change_deck", [True, False])
def test_data_edit_requires_changed_native_embedded_data(tmp_path, kind, change_deck):
    _, _, cases = frozen_batch(tmp_path)
    attempt = delivery_fixture(tmp_path)
    paths = {key: Path(value) for key, value in attempt["records"].items()}
    scene_path = Path(attempt["deliveries"]["editable"]["scene"])
    scene = load(scene_path)
    element = {"id": "results", "type": kind, "x": 100, "y": 500, "w": 600, "h": 250}
    if kind == "chart":
        element.update(chart_type="column", categories=["A", "B"], series=[{"name": "Series", "values": [10, 20]}])
    else:
        element["rows"] = [["Period", "Value"], ["A", "10"]]
    scene["slides"][0]["elements"].append(element)
    save(scene_path, scene)
    checks = load(paths["edit_check"])
    reference = tmp_path / "slides/01-cover.png"
    for operation in checks["operations"]:
        after = Path(operation["after_scene"])
        changed = load(after)
        changed["slides"][0]["elements"].append(copy.deepcopy(element))
        save(after, changed)
        deck = Path(operation["after_deck"])
        build_editable_ppt.build_deck(after, deck)
        render_fixture(deck, Path(operation["render_report"]), reference)
    before_deck = tmp_path / "before-data.pptx"
    build_editable_ppt.build_deck(scene_path, before_deck)
    changed = copy.deepcopy(scene)
    target = changed["slides"][0]["elements"][-1]
    if kind == "chart":
        target["series"][0]["values"][0] = 99
    else:
        target["rows"][1][1] = "99"
    after = tmp_path / "edited-data.json"
    save(after, changed)
    deck = tmp_path / "edited-data.pptx"
    build_editable_ppt.build_deck(after if change_deck else scene_path, deck)
    report_path = tmp_path / "edited-data-render.json"
    render_fixture(deck, report_path, reference)
    checks["operations"].append({"kind": f"{kind}_data_edit", "status": "pass", "notes": "Actual native data exercise",
                                 "slide_id": "s01", "element_id": "results", "before_scene": str(scene_path),
                                 "after_scene": str(after), "after_deck": str(deck), "render_report": str(report_path)})
    save(paths["edit_check"], checks)
    if change_deck:
        exercised = benchmark.validate_work_records(paths, cases["cases"][0], lambda path: path, scene_path)
        assert f"{kind}_data_edit" in exercised
    else:
        with pytest.raises(ValueError, match="Modified deck differs"):
            benchmark.validate_work_records(paths, cases["cases"][0], lambda path: path, scene_path)


@pytest.mark.parametrize("mutation", [None, "nonuniform", "content-change", "no-movement", "missing-child"])
def test_group_move_translates_all_absolute_children_without_changing_content(tmp_path, mutation):
    _, _, value = frozen_batch(tmp_path)
    attempt = delivery_fixture(tmp_path)
    paths = {key: Path(path) for key, path in attempt["records"].items()}
    scene_path = Path(attempt["deliveries"]["editable"]["scene"])
    scene = load(scene_path)
    group = {"id": "mechanism", "type": "group", "children": [
        {"id": "node", "type": "shape", "shape": "rect", "x": 900, "y": 500, "w": 100, "h": 80},
        {"id": "labels", "type": "group", "children": [
            {"id": "node-text", "type": "text", "text": "Mechanism", "x": 910, "y": 510, "w": 80, "h": 30}]},
        {"id": "edge", "type": "line", "x1": 1000, "y1": 535, "x2": 1100, "y2": 535}]}
    scene["slides"][0]["elements"].append(group)
    save(scene_path, scene)
    build_editable_ppt.build_deck(scene_path, tmp_path / "deck.pptx")
    checks = load(paths["edit_check"])

    def translate(element):
        if element["type"] == "group":
            for child in element["children"]:
                translate(child)
        else:
            for field in ("x", "x1", "x2"):
                if field in element:
                    element[field] += 20
            for field in ("y", "y1", "y2"):
                if field in element:
                    element[field] += 30

    for operation in checks["operations"]:
        after_scene = Path(operation["after_scene"])
        after = load(after_scene)
        actual_group = copy.deepcopy(group)
        after["slides"][0]["elements"].append(actual_group)
        if operation["kind"] == "object_move":
            after["slides"][0]["elements"][0]["x"] = scene["slides"][0]["elements"][0]["x"]
            operation["element_id"] = "mechanism"
            if mutation != "no-movement":
                translate(actual_group)
            if mutation == "nonuniform":
                actual_group["children"][0]["x"] += 10
            elif mutation == "content-change":
                actual_group["children"][1]["children"][0]["text"] = "Different mechanism"
            elif mutation == "missing-child":
                actual_group["children"].pop()
        save(after_scene, after)
        after_deck = Path(operation["after_deck"])
        build_editable_ppt.build_deck(after_scene, after_deck)
        render_fixture(after_deck, Path(operation["render_report"]), tmp_path / "slides/01-cover.png")
    save(paths["edit_check"], checks)
    if mutation is None:
        assert benchmark.validate_work_records(paths, value["cases"][0], lambda path: path, scene_path) == {
            "text_edit", "object_move"}
    else:
        with pytest.raises(ValueError, match="translate the whole object without altering its content"):
            benchmark.validate_work_records(paths, value["cases"][0], lambda path: path, scene_path)


def test_chart_series_rename_is_not_a_data_edit(tmp_path):
    _, _, value = frozen_batch(tmp_path)
    attempt = delivery_fixture(tmp_path)
    paths = {key: Path(path) for key, path in attempt["records"].items()}
    scene_path = Path(attempt["deliveries"]["editable"]["scene"])
    scene = load(scene_path)
    chart = {"id": "results", "type": "chart", "x": 100, "y": 500, "w": 600, "h": 250,
             "chart_type": "column", "categories": ["A", "B"], "series": [{"name": "Before", "values": [10, 20]}]}
    scene["slides"][0]["elements"].append(chart)
    save(scene_path, scene)
    changed = copy.deepcopy(scene)
    changed["slides"][0]["elements"][-1]["series"][0]["name"] = "Renamed only"
    after = tmp_path / "series-name.json"
    deck = tmp_path / "series-name.pptx"
    report = tmp_path / "series-name-render.json"
    save(after, changed)
    build_editable_ppt.build_deck(after, deck)
    render_fixture(deck, report, tmp_path / "slides/01-cover.png")
    # Put the actual data operation first so unrelated fixture operations cannot
    # mask rejection of a real renamed chart with unchanged embedded values.
    save(paths["edit_check"], {"operations": [{"kind": "chart_data_edit", "status": "pass", "notes": "Rename only",
                                              "slide_id": "s01", "element_id": "results", "before_scene": str(scene_path),
                                              "after_scene": str(after), "after_deck": str(deck), "render_report": str(report)}]})
    with pytest.raises(ValueError, match="change actual series values"):
        benchmark.validate_work_records(paths, value["cases"][0], lambda path: path, scene_path)


def test_inspect_cli_prepares_binding_without_recording_attempt_or_changing_inputs(tmp_path):
    batch_path, frozen, _ = frozen_batch(tmp_path)
    attempt = delivery_fixture(tmp_path)
    attempt_path = tmp_path / "attempt.json"
    output = tmp_path / "inspection.json"
    save(attempt_path, attempt)
    expected = benchmark.assess(frozen, attempt, tmp_path)["review_input_sha256"]
    original = {str(path): benchmark.digest(path) for path in tmp_path.rglob("*") if path.is_file()}
    script = Path(benchmark.__file__)
    result = subprocess.run([sys.executable, "-X", "utf8", str(script), "inspect", str(batch_path),
                             str(attempt_path), str(output)], text=True, encoding="utf-8", capture_output=True, check=True)
    stdout = json.loads(result.stdout)
    assert stdout["review_input_sha256"] == expected
    assert stdout["status"] == "incomplete"  # No semantic review provided yet.
    assert benchmark.sealed_read(output)["review_input_sha256"] == expected
    assert not list(tmp_path.rglob("attempt-*.json"))
    assert not list(tmp_path.rglob("head.json"))
    assert all(benchmark.digest(path) == digest for path, digest in original.items())


def image_deck_fixture(tmp_path, fit="cover", image_size=(160, 90)):
    spec_path, _, _, _, _ = make_case(tmp_path)
    spec = load(spec_path)
    reference = tmp_path / spec["slides"][0]["image_file"]
    Image.new("RGB", image_size, "white").save(reference)
    output = tmp_path / "image-deck.pptx"
    # Strict PageSpec correctly rejects mismatched ratios. Legacy directory
    # input exercises the builder's supported nontrivial cover/contain geometry.
    input_path = spec_path if image_size == (160, 90) else reference.parent
    build_image_ppt.build_deck(Namespace(input_dir=input_path, output=output, width=None, height=None,
                                       title="", max_width=0, jpeg_quality=0, background="FFFFFF",
                                       fit=fit, extensions=[".png"]))
    return output, spec, {"slides": [{"reference": str(reference)}]}


@pytest.mark.parametrize("fit", ["cover", "contain"])
@pytest.mark.parametrize("image_size", [(160, 90), (160, 160), (300, 90)])
def test_image_branch_accepts_real_image_builder_fit_geometry(tmp_path, fit, image_size):
    deck, spec, report = image_deck_fixture(tmp_path, fit, image_size)
    benchmark.verify_image_deck(deck, spec, report)


@pytest.mark.parametrize("mutation", ["wrong-pixels", "small-picture", "extra-text", "wrong-crop"])
def test_image_branch_rejects_false_reference_or_partial_page(tmp_path, mutation):
    deck, spec, report = image_deck_fixture(tmp_path)
    if mutation == "wrong-pixels":
        Image.new("RGB", (160, 90), "black").save(report["slides"][0]["reference"])
    else:
        presentation = Presentation(deck)
        slide = presentation.slides[0]
        if mutation == "small-picture":
            slide.shapes[0].width //= 2
        elif mutation == "extra-text":
            slide.shapes.add_textbox(0, 0, 100000, 100000).text = "Overlay"
        else:
            slide.shapes[0].crop_left = .2
        presentation.save(deck)
    with pytest.raises(ValueError, match="reference pixels|contain/cover geometry|exactly one approved"):
        benchmark.verify_image_deck(deck, spec, report)


def test_editable_native_deck_cannot_masquerade_as_image_branch(tmp_path):
    _, frozen, _ = frozen_batch(tmp_path, scope="image_and_editable")
    attempt = delivery_fixture(tmp_path)
    editable = attempt["deliveries"]["editable"]
    attempt["deliveries"]["image"] = {key: value for key, value in editable.items() if key != "scene"}
    log_path = Path(attempt["records"]["run_log"])
    log = load(log_path)
    log["events"].extend({"stage": stage, "outcome": "pass", "evidence": editable["render_report"]}
                         for stage in ["style"] * 4 + ["initial_page"])
    save(log_path, log)
    assessed = benchmark.assess(frozen, attempt, tmp_path)
    assert assessed["status"] == "incomplete"
    assert "exactly one approved page image" in assessed["issues"][0]


@pytest.mark.parametrize("mutation", [None, "wrong-page-id", "duplicate-rendered-path"])
def test_edited_render_requires_contiguous_page_ids_and_distinct_files(tmp_path, mutation):
    _, _, value = frozen_batch(tmp_path)
    attempt = delivery_fixture(tmp_path)
    paths = {key: Path(path) for key, path in attempt["records"].items()}
    scene_path = Path(attempt["deliveries"]["editable"]["scene"])
    scene = load(scene_path)
    second = copy.deepcopy(scene["slides"][0])
    second["id"] = "s02"
    scene["slides"].append(second)
    save(scene_path, scene)
    # Two real PPTX pages; only report PNGs are synthetic test evidence.
    checks = load(paths["edit_check"])
    for operation in checks["operations"]:
        after_scene = Path(operation["after_scene"])
        after = load(after_scene)
        after["slides"].append(copy.deepcopy(second))
        save(after_scene, after)
        after_deck = Path(operation["after_deck"])
        build_editable_ppt.build_deck(after_scene, after_deck)
        report_path = Path(operation["render_report"])
        report = render_fixture(after_deck, report_path, tmp_path / "slides/01-cover.png")
        second_render = report_path.with_name(report_path.stem + "-002.png")
        Image.new("RGB", (160, 90), "white").save(second_render)
        second_page = {**report["slides"][0], "slide": 2, "rendered": str(second_render),
                       "rendered_sha256": benchmark.digest(second_render)}
        report["slides"].append(second_page)
        if mutation == "wrong-page-id":
            second_page["slide"] = 1
        elif mutation == "duplicate-rendered-path":
            second_page["rendered"] = report["slides"][0]["rendered"]
        save(report_path, report)
    if mutation is None:
        assert benchmark.validate_work_records(paths, value["cases"][0], lambda path: path, scene_path) == {
            "text_edit", "object_move"}
    else:
        with pytest.raises(ValueError, match="unique and in slide order"):
            benchmark.validate_work_records(paths, value["cases"][0], lambda path: path, scene_path)
