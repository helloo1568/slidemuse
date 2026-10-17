import json
import os
import subprocess
import sys
from pathlib import Path

import generate_backend
import pytest
from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]


def invoke(*arguments):
    return subprocess.run([sys.executable, str(ROOT / "slidemuse.py"), "generate", *map(str, arguments)],
                          capture_output=True, encoding="utf-8", check=False)


def deck_spec(name="示例课件"):
    """A deck spec matching the vendored backend's documented job input."""
    return {
        "deck_name": name,
        "language": "Chinese",
        "goal": "解释核心论点并用两页展示证据。",
        "selected_image_backend": "built-in image tool",
        "style": {"name": "清爽专业风", "visual_direction": "calm professional explainer"},
        "slides": [
            {"number": 1, "title": "封面", "role": "cover", "intent": "开始汇报",
             "key_points": ["主题", "场合", "署名"]},
            {"number": 2, "title": "证据", "role": "data-evidence", "intent": "展示来源数据",
             "key_points": ["来源与口径", "图形表达的关系", "结论边界"]},
        ],
    }


def write_slide_image(path, size=(1600, 900)):
    Image.new("RGB", size, "#F8F5EB").save(path)


def run_pipeline(tmp_path, name="示例课件"):
    deck = tmp_path / name
    (tmp_path / "deck_spec.json").write_text(json.dumps(deck_spec(name), ensure_ascii=False), encoding="utf-8")
    prepared = invoke("--offline", "prepare", "--spec", tmp_path / "deck_spec.json", "--out-dir", deck)
    assert prepared.returncode == 0, prepared.stdout + prepared.stderr
    return deck


def test_complete_sources_are_verified():
    result = generate_backend.capabilities(verify=True)
    assert result["source_integrity"]["checked"] >= 41
    assert not result["source_integrity"]["mismatches"]
    assert result["commit"] == "941dcafe2f23ea4ae61de99e9b22fb0ef3d9e06c"
    assert result["dependencies_ready"]
    assert "outline-driven-slide-jobs" in result["features"]


def test_prepare_writes_per_slide_jobs_and_state(tmp_path):
    deck = run_pipeline(tmp_path)
    prompt = json.loads((deck / "prompts/slide_02.json").read_text(encoding="utf-8"))
    assert prompt["out"] == "slide_02.png"
    assert prompt["expected_backend"] == "built-in image tool"
    assert "图形表达的关系" in prompt["prompt"]
    jobs = json.loads((deck / "slide_jobs.json").read_text(encoding="utf-8"))
    assert [entry["status"] for entry in jobs["slides"]] == ["pending", "pending"]
    assert (deck / "origin_image").is_dir()
    status = invoke("--offline", "status", deck, "--json")
    assert status.returncode == 0, status.stderr
    summary = json.loads(status.stdout)
    assert summary["dispatchable_slides"] == ["slide_01", "slide_02"]
    assert summary["run_status"] == "jobs_prepared"


def test_dispatch_result_and_assemble_produce_deck_with_notes(tmp_path):
    deck = run_pipeline(tmp_path)
    (deck / "speech.md").write_text("## Slide 1: 封面\n开场说明。\n\n## Slide 2: 证据\n数据来自报告表 2。\n",
                                   encoding="utf-8")
    for number in (1, 2):
        slide_id = f"slide_{number:02d}"
        generated = tmp_path / f"{slide_id}-generated.png"
        write_slide_image(generated)
        dispatched = invoke("--offline", "dispatch", deck, "--slide", number, "--agent-id", "main")
        assert dispatched.returncode == 0, dispatched.stdout + dispatched.stderr
        recorded = invoke("--offline", "result", deck, "--slide", slide_id, "--agent-id", "main",
                          "--backend-used", "built-in image tool", "--selected-source", generated,
                          "--qa-note", "文字与要点齐全")
        assert recorded.returncode == 0, recorded.stdout + recorded.stderr
    jobs = json.loads((deck / "slide_jobs.json").read_text(encoding="utf-8"))
    assert jobs["run_status"] == "slides_recorded"
    assert (deck / "origin_image/slide_02.png").is_file()
    assembled = invoke("--offline", "assemble", tmp_path, f"{deck.name}.pptx")
    assert assembled.returncode == 0, assembled.stdout + assembled.stderr
    presentation = Presentation(tmp_path / deck.name / f"{deck.name}.pptx")
    assert len(presentation.slides) == 2
    assert "数据来自报告表 2" in presentation.slides[1].notes_slide.notes_text_frame.text


def test_result_rejects_local_rendering_as_final_image(tmp_path):
    deck = run_pipeline(tmp_path)
    write_slide_image(tmp_path / "local.png")
    assert invoke("--offline", "dispatch", deck, "--slide", 1, "--agent-id", "main").returncode == 0
    rejected = invoke("--offline", "result", deck, "--slide", 1, "--agent-id", "main",
                      "--backend-used", "Pillow script render", "--selected-source", tmp_path / "local.png",
                      "--qa-note", "本地绘制")
    assert rejected.returncode != 0
    assert "Forbidden backend" in rejected.stderr + rejected.stdout


def test_offline_blocks_networked_runtime_commands(tmp_path):
    blocked = invoke("--offline", "runtime", "doctor")
    assert blocked.returncode == 2
    assert "--offline does not support" in blocked.stderr


def test_offline_isolates_credentials_and_keeps_explicit_home(monkeypatch, tmp_path):
    calls = []

    def capture(command, **kwargs):
        calls.append(kwargs["env"])
        return subprocess.CompletedProcess(command, 0)

    monkeypatch.setattr(generate_backend.subprocess, "run", capture)
    monkeypatch.setenv("CODEX_PPT_HOME", str(tmp_path / "user-home"))
    monkeypatch.setenv("OPENAI_API_KEY", "must-not-leak")

    assert generate_backend.run_backend(["status", "deck"]) == 0
    assert calls[0]["CODEX_PPT_HOME"] == str(tmp_path / "user-home")
    assert calls[0]["OPENAI_API_KEY"] == "must-not-leak"

    assert generate_backend.run_backend(["status", "deck"], offline=True) == 0
    assert calls[1]["CODEX_PPT_HOME"] != str(tmp_path / "user-home")
    assert "OPENAI_API_KEY" not in calls[1]
    assert calls[1]["PYTHONUTF8"] == "1"


def test_unknown_command_reports_available_commands():
    result = invoke("not-a-command")
    assert result.returncode == 2
    assert "assemble" in result.stderr


@pytest.mark.parametrize("name", ["普通 deck", "带 空格 的 deck"])
def test_deck_paths_with_spaces_and_unicode(tmp_path, name):
    deck = run_pipeline(tmp_path, name)
    assert (deck / "prompts/slide_01.json").is_file()
    assert len(os.listdir(deck / "origin_image")) == 0
