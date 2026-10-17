import json
import os
import runpy
import subprocess
import sys
from pathlib import Path

import editable_backend
import pytest
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

import install as installer

ROOT = Path(__file__).resolve().parents[1]
RUNTIME = ROOT / editable_backend.VENDOR / editable_backend.SKILL / "cli/editppt/runtime"


def invoke(*arguments):
    return subprocess.run([sys.executable, str(ROOT / "slidemuse.py"), "rebuild", *map(str, arguments)],
                          capture_output=True, encoding="utf-8", check=False)


def synthetic_image(path):
    image = Image.new("RGB", (640, 360), "white")
    try:
        font = ImageFont.truetype("arial.ttf" if os.name == "nt" else "DejaVuSans.ttf", 26)
    except OSError:
        font = ImageFont.load_default(size=26)
    ImageDraw.Draw(image).text((50, 80), "Source 123", font=font, fill="black")
    image.save(path)


def test_complete_sources_and_local_patches_are_verified():
    result = editable_backend.capabilities(verify=True)
    assert result["source_integrity"]["checked"] >= 50
    assert not result["source_integrity"]["mismatches"]
    assert result["dependencies_ready"]
    assert "native-custom-paths-dashes-arrows" in result["features"]


def test_offline_unicode_prepare_measures_text_without_using_configured_tokens(tmp_path, monkeypatch):
    source = tmp_path / "中文输入.png"
    synthetic_image(source)
    monkeypatch.setenv("PADDLE_OCR_TOKEN", "not-a-real-token-must-not-be-used")
    monkeypatch.setenv("CODEX_AUTH_FILE", str(tmp_path / "private-auth.json"))
    job = tmp_path / "中文任务"
    result = invoke("--offline", "prepare", source, "--job-dir", job)
    assert result.returncode == 0, result.stdout + result.stderr
    deck = json.loads((job / "deck_manifest.json").read_text(encoding="utf-8"))
    assert deck["page_count"] == 1
    assert deck["image_backend"]["backend_id"] == "builtin-imagegen"
    page = job / "pages/page_001"
    hints = json.loads((page / "text_hints.json").read_text(encoding="utf-8"))
    assert hints["backend"] == "builtin-ink"
    assert hints["lines"]
    assert Image.open(page / "source.png").size == (640, 360)
    prompt = invoke("page-prompt", job, "--page", "1", "--out", "worker-prompt.md")
    assert prompt.returncode == 0, prompt.stderr
    text = (page / "worker-prompt.md").read_text(encoding="utf-8")
    assert "SlideMuse integration" in text and "slidemuse.py" in text
    assert invoke("page-prompt", job, "--page", "1", "--out", tmp_path / "escape.md").returncode != 0
    assert not (tmp_path / "escape.md").exists()


@pytest.mark.parametrize("kind", ["pdf", "pptx"])
def test_normalization_keeps_pages_and_pptx_notes(tmp_path, kind):
    if kind == "pdf":
        import fitz
        source = tmp_path / "input.pdf"
        with fitz.open() as document:
            document.new_page(width=320, height=180)
            document.new_page(width=180, height=320)
            document.save(source)
        count = 2
    else:
        image = tmp_path / "slide.png"
        synthetic_image(image)
        source = tmp_path / "input.pptx"
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(8), Inches(4.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(image), 0, 0, width=prs.slide_width, height=prs.slide_height)
        slide.notes_slide.notes_text_frame.text = "Original speaker notes 123"
        prs.save(source)
        count = 1
    job = tmp_path / "run"
    result = invoke("--offline", "prepare", source, "--job-dir", job)
    assert result.returncode == 0, result.stdout + result.stderr
    assert json.loads((job / "deck_manifest.json").read_text(encoding="utf-8"))["page_count"] == count
    if kind == "pptx":
        notes = json.loads((job / "notes_manifest.json").read_text(encoding="utf-8"))
        assert "Original speaker notes 123" in notes["notes"][0]["text"]


def test_native_manifest_build_preserves_curve_table_styles_and_text(tmp_path):
    manifest = {"source": {"width_px": 800, "height_px": 450},
                "slide": {"width": 8, "height": 4.5}, "preview_scale": 80,
                "content_box": {"left": 0, "top": 0, "width": 8, "height": 4.5},
                "text_boxes": [{"text": "Backend smoke test", "box_px": [40, 20, 600, 60],
                                "font_size": 24, "fit_text": False}],
                "shapes": [{"type": "path", "box_px": [40, 100, 300, 80], "fill": "none",
                            "stroke": "#2463EB", "stroke_width": 2, "dash": "dash", "end_arrow": "triangle",
                            "path_px": [{"op": "moveTo", "points": [[40, 180]]},
                                        {"op": "cubicBezTo", "points": [[80, 100], [260, 100], [340, 180]]}]}],
                "tables": [{"box_px": [40, 210, 700, 180], "row_heights": [1.5, 1],
                            "style": {"font": "Arial", "font_size": 16, "fit_text": False},
                            "cells": [[{"text": "Merged", "col_span": 2, "style": {"fill": "#DDEEFF"}}, ""],
                                      ["A", {"text": "10", "style": {"bold": True}}]]}],
                "images": [], "visual_inventory": [],
                "background_strategy": {"mode": "native-or-script", "comparison_note": "Synthetic white canvas"},
                "quality_checks": dict.fromkeys(["font_size_calibrated", "visual_inventory_matched",
                                                  "background_strategy_checked", "shape_corner_geometry_checked"], True)}
    (tmp_path / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    built = invoke("--offline", "page", "build", tmp_path)
    assert built.returncode == 0, built.stderr
    validated = invoke("--offline", "page", "validate", tmp_path, "--report", "validation.json")
    assert validated.returncode == 0, validated.stdout + validated.stderr
    assert json.loads((tmp_path / "validation.json").read_text(encoding="utf-8"))["passed"]
    prs = Presentation(tmp_path / "page.pptx")
    table = next(shape.table for shape in prs.slides[0].shapes if shape.has_table)
    assert table.cell(0, 0).span_width == 2
    assert table.cell(1, 1).text == "10"
    assert (tmp_path / "preview.png").is_file()


def test_offline_blocks_live_image_requests_before_reading_sources():
    result = invoke("--offline", "image", "generate", "--prompt", "unused")
    assert result.returncode == 2
    assert "--dry-run" in result.stderr


def test_extra_dependency_failure_preserves_existing_install(tmp_path, monkeypatch):
    target = tmp_path / "slidemuse"
    target.mkdir()
    (target / ".skill-install.json").write_text("{}")
    (target / "previous.txt").write_text("approved")
    monkeypatch.setattr(installer, "install_dependencies", lambda stage, final: Path(sys.executable))

    def fail(*args, **kwargs):
        raise subprocess.CalledProcessError(1, args[0])

    monkeypatch.setattr(installer.subprocess, "run", fail)
    with pytest.raises(subprocess.CalledProcessError):
        installer.install_skill(ROOT, target, "codex", "slidemuse", False, False, True)
    assert (target / "previous.txt").read_text() == "approved"


def test_formula_png_uses_real_pdf_renderer_without_windows_convert(tmp_path):
    import fitz
    pdf = tmp_path / "formula.pdf"
    with fitz.open() as document:
        document.new_page(width=100, height=40).insert_text((10, 25), "x = 1")
        document.save(pdf)
    helpers = runpy.run_path(str(RUNTIME / "formula_renderer.py"))
    output = tmp_path / "formula.png"
    assert helpers["convert_pdf"](pdf, output, "png", dpi=144, timeout=10) == "PyMuPDF"
    assert Image.open(output).size == (200, 80)
